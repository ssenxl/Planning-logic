"""
Replan — เอางาน booking ที่ยังไม่ล็อก มาวางแผนใหม่พร้อมกับ SC ในรอบเดียว
=========================================================================

เดิม: ดีมานด์ = SC ใหม่อย่างเดียว ; booking = "ของที่จองเครื่องไปแล้ว" ถูกหักออกจากเครื่องว่าง
ใหม่: ดีมานด์ = SC ใหม่ + ยอด booking ตั้งแต่สัปดาห์ที่ปลดล็อก แล้ววางพร้อมกันทั้งก้อน

เส้นแบ่งล็อก/ปลดล็อก = TODAY_IDX + 2 (พื้นเดียวกับที่ Planning ใช้เป็นสัปดาห์แรกที่วางงานได้)
  - WEEK ที่ idx <  TODAY_IDX+2 = ล็อก  → คงแผนเดิมของ booking ไว้ (กำลังทอ/ตั้งเครื่องแล้ว)
  - WEEK ที่ idx >= TODAY_IDX+2 = ปลดล็อก → ยอดถูกยกออกมาเป็นดีมานด์ + คืนเครื่องให้ว่าง

ล็อกที่ผู้ใช้กดเอง (ชีท "Lock_Plan" ใน MasterMC.xlsx — กดจากการ์ดบล็อกบน Gantt หรือแก้มือ):
  แถวที่ตรงคีย์ (WEEK, ITEM_CODE, MC_GROUP, GUAGE) ถือว่า "ล็อก" แม้อยู่หลังเส้น TODAY+2
  → ไม่ถูกยกเป็นดีมานด์ และเครื่องของแถวนั้นไม่ถูกคืนให้ว่าง (ดู write_replan_capacity)

ใบจองที่ออกเลข KP_NO แล้ว (21 ก.ย. 2569): ถือว่า "ล็อก" เหมือนกัน แม้อยู่หลังเส้น TODAY+2
  → คงสัปดาห์/เครื่องเดิมของ booking ไว้ ไม่ยกมาวางใหม่ (ดู kp_locked_mask)

กันนับเครื่องซ้ำ (สำคัญ): ถ้าปล่อยให้ booking ทั้งกินเครื่องอยู่ (TOTAL_MC_REMAIN หักไปแล้ว)
และขอเครื่องใหม่พร้อมกัน เครื่องจะหายไปสองเท่า → จึงต้องสร้างไฟล์ capacity คู่ขนานที่
    - DETAIL  : เหลือเฉพาะสัปดาห์ที่ล็อก
    - SUMMARY : สัปดาห์ที่ปลดล็อก MC_USE_CEIL = 0 และ TOTAL_MC_REMAIN = TOTAL_MC

ผลลัพธ์ (data_plan/):
    booking_demand.xlsx        – ยอด booking ที่ถูกยกมาเป็นดีมานด์ + ชีท audit การกันซ้ำกับ SC
    order_replan.xlsx          – ดีมานด์รวม (order_ready + booking) ที่ Planning รอบ replan ใช้
    booking_final_replan.xlsx  – capacity คู่ขนาน (ล็อก/ปลดล็อกตามด้านบน)
    replan_append.xlsx         – แถวที่ต้องแปะเข้าแผนโดยไม่ถูกจัดใหม่ (งานที่ตรึง 🔐 + core ที่เบรกไว้)
                                 Planning รอบ replan อ่านไปต่อท้ายแผนแล้วเติม TARGET_KNIT / ทีมให้
    replan_result.xlsx         – แผนที่วางใหม่ (ชีท PLAN) + ชีท COMPARE (booking เดิม vs แผนใหม่)
    production_plan_*.xlsx     – ไฟล์แผนผลิตที่หน้าเว็บเปิด = สำเนาของ replan_result
                                 (ชีท COMPARE เปลี่ยนชื่อเป็น REPLAN_COMPARE)

ตั้งแต่ 2026-09-11 ไม่มี Planning รอบหลัก (แผน SC ล้วน) แล้ว — แผนผลิตมีรอบเดียวคือรอบนี้
ชีท PLAN ของไฟล์แผนจึงเป็นแผน replan (ไฟล์รุ่นเก่ามี PLAN = SC ล้วน + REPLAN แยก ; ดู plan_view)

Usage:
    python Replan.py
"""

import configparser as _cp
import os
import re
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

from Calendar import (_CALENDAR_LOCAL_PATH, calendar_week_map, load_calendar,
                      WeekTimeline, yw_week, yw_year, yw_label)
import ColorOpen  # so_text: ล้าง ".0"/"nan" ของเลขใบสั่งที่ pandas อ่านเป็นตัวเลข
import file_guard  # เขียนไฟล์แบบสลับชื่อ + เข้าคิวกับหน้าเว็บที่แก้ไฟล์เดียวกัน
import plan_config  # ค่าปรับแต่งที่แก้ได้จากหน้า "ตั้งค่าระบบ"

# ── คีย์สัปดาห์ภายในไฟล์นี้ = YW (ปี*100+สัปดาห์) ────────────────────────────────────
# ไฟล์ที่อ่านเข้ามา (booking DETAIL / ไฟล์แผน / Lock_Plan) แยกสัปดาห์กับปีคนละคอลัมน์
# → แปลงเป็น YW ตอนโหลด (_to_yw) ; ไฟล์ที่เขียนออกแยกกลับเป็นสัปดาห์ + ปี (_split_yw)
# เลขสัปดาห์ที่ไม่มีปี = ครั้งที่ใกล้วันนี้สุด (WeekTimeline.resolve_bare)
_TL: WeekTimeline | None = None


def _timeline() -> WeekTimeline:
    global _TL
    if _TL is None:
        cal = load_calendar(_CALENDAR_LOCAL_PATH, sheet_name="Sheet1")
        _TL = WeekTimeline(calendar_week_map(cal))
    return _TL


def _to_yw(df, week_col="WEEK", year_col="YEAR"):
    """คอลัมน์สัปดาห์ของตาราง → YW (แก้ในที่) ; ใช้ปีจาก year_col ถ้ามี ; W99 / YW เดิมคงไว้"""
    if df is None or df.empty or week_col not in df.columns:
        return df
    tl = _timeline()
    w = pd.to_numeric(df[week_col], errors="coerce")
    y = pd.to_numeric(df[year_col], errors="coerce") if year_col in df.columns else None
    df[week_col] = pd.to_numeric(pd.Series(
        [None if pd.isna(wv) else tl.key(int(wv), None if y is None or pd.isna(y.iat[i]) else int(y.iat[i]))
         for i, wv in enumerate(w)],
        index=df.index), errors="coerce")
    return df


def _split_yw(df, pairs):
    """ก่อนเขียนไฟล์: คอลัมน์ YW → เลขสัปดาห์ + คอลัมน์ปี (คืนสำเนา) ; ค่าที่ไม่ใช่ YW คงเดิม
    pairs = [(คอลัมน์สัปดาห์, คอลัมน์ปี หรือ None)]"""
    if df is None or df.empty:
        return df
    out = df.copy()
    for wc, yc in pairs:
        if wc not in out.columns:
            continue
        nums = pd.to_numeric(out[wc], errors="coerce")
        is_yw = nums >= 100000
        if not is_yw.any():
            continue
        if yc:
            old_y = pd.to_numeric(out[yc], errors="coerce") if yc in out.columns else pd.Series(pd.NA, index=out.index)
            out[yc] = pd.array([int(n) // 100 if f else (None if pd.isna(o) else int(o))
                                for n, f, o in zip(nums, is_yw, old_y)], dtype="Int64")
        out[wc] = [int(n) % 100 if f else v for n, f, v in zip(nums, is_yw, out[wc])]
    return out

# =========================
# CONFIG
# =========================
BASE_DIR = Path(__file__).parent
DATA_PLAN_DIR = BASE_DIR / "data_plan"

ORDER_FILE = DATA_PLAN_DIR / "order_ready.xlsx"
# ไฟล์ SC ดิบก่อน Order.py กรอง — ใช้หา DYE_END_DATE ของใบ YD ที่ถูกกรองออกจาก order_ready
# (Order.py ตัด CL-ORDERS/FQC/blocklist ทิ้ง แต่ booking ยังยกงานใบนั้นมาวางใหม่ได้)
ORDER_SRC_DIR = BASE_DIR / "Order"
DEMAND_FILE = DATA_PLAN_DIR / "booking_demand.xlsx"
ORDER_REPLAN_FILE = DATA_PLAN_DIR / "order_replan.xlsx"
BOOKING_REPLAN_FILE = DATA_PLAN_DIR / "booking_final_replan.xlsx"
RESULT_FILE = DATA_PLAN_DIR / "replan_result.xlsx"
# แถวที่แปะเข้าแผนตรง ๆ (งานที่ตรึง + core ที่เบรก + ใบที่ออกเลข KP) — Planning รอบ replan
# อ่านผ่าน env KP_REPLAN_APPEND_FILE
APPEND_FILE = DATA_PLAN_DIR / "replan_append.xlsx"
APPEND_PINNED_SHEET = "PINNED"
APPEND_CORE_SHEET = "CORE"
APPEND_KP_SHEET = "KP"
# ที่วางไฟล์ production_plan_*.xlsx — env KP_PLAN_OUTPUT_DIR ไว้ให้รันทดสอบลง scratchpad
# (ไฟล์ที่ใหม่สุดในโฟลเดอร์จริงคือแผนที่หน้าเว็บเปิด ห้ามให้รอบทดสอบไปทับ)
PLAN_OUTPUT_DIR = Path(os.environ.get("KP_PLAN_OUTPUT_DIR") or DATA_PLAN_DIR)
COMPARE_SHEET = "COMPARE"                  # ชื่อชีทใน replan_result
PUBLISHED_COMPARE_SHEET = "REPLAN_COMPARE"  # ชื่อในไฟล์แผน (replan_notify / หน้าเว็บอ่านชื่อนี้)

# สัปดาห์แรกที่วางงานได้ = สัปดาห์ปัจจุบัน + ค่านี้ (ค่าเดียวกับพื้น TODAY_IDX+FREEZE_WEEKS ใน
# Planning.py — มาจากหน้า "ตั้งค่าระบบ" ที่เดียว จึงเลื่อนพร้อมกันเสมอ)
# ทุกอย่างก่อนหน้านั้นถือว่าล็อก — ตั้งเครื่อง/ขึ้นด้ายไปแล้ว ขยับไม่ได้
FREEZE_WEEKS_AHEAD = plan_config.get("KP_FREEZE_WEEKS")

# ล็อกที่ผู้ใช้กดเอง — ชีทเดียวกับที่ webapp/backend/plan_lock.py เขียน (MasterMC.xlsx)
# แก้คอลัมน์ที่นี่ต้องแก้ที่ plan_lock.py ด้วย (คนละโปรเซส อ่านไฟล์เดียวกัน)
LOCK_PLAN_SHEET = "Lock_Plan"

# คู่ (SC, item) ที่อยู่ทั้งใน booking และ view_sc = งานเดียวกัน (ดู booking_view_field_map)
# ถ้า append ตรง ๆ จะได้งานซ้ำสองเท่า → ให้ยอดจาก booking ชนะ เพราะ booking แตกรายสัปดาห์แล้ว
# (รู้ว่าส่วนไหนทอไป/ล็อกไปแล้ว) ส่วน PENDING_PLAN ของ SC เป็นยอดค้างทั้งใบซึ่งรวมส่วนที่ล็อกอยู่ด้วย
BOOKING_WINS_ON_DUPLICATE = plan_config.get("KP_BOOKING_WINS_ON_DUPLICATE")

# ORDER_TYPE ที่ Planning มี routing รองรับจริง (ดู Planning.py — branch LAB-DIP / SC-ORDERS /
# ผ้าเติม / YD-ORDERS ; ค่าอื่นตกท้าย else แล้วถูกโยนเข้าชีท UNPLANNED "routing ไม่ครอบคลุม")
# ใช้ 2 ที่: เลือกชนิดตอนแตก FR_CODE ที่ยุบหลายใบ + เตือนก่อนส่งให้ Planning
# แก้รายการได้ที่หน้า "ตั้งค่าระบบ" › ชนิดใบสั่งที่ระบบวางแผนให้ (เทียบตรงตัว ไม่แปลงตัวพิมพ์)
PLANNABLE_ORDER_TYPES = tuple(
    str(_t).strip() for _t in plan_config.get_list("KP_PLANNABLE_ORDER_TYPES") if str(_t).strip()
)


def _latest_booking_final() -> Path:
    """ไฟล์ booking_final_ready ที่ใหม่สุด (เหมือน Planning._latest_booking_final)"""
    cands = sorted(
        DATA_PLAN_DIR.glob("booking_final_ready_*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if cands:
        return cands[0]
    return DATA_PLAN_DIR / "booking_final_ready25.xlsx"


# =========================
# LOCK_PLAN: ล็อกที่ผู้ใช้กดเอง
# =========================
def _master_mc_path() -> Path | None:
    """path ของ MasterMC.xlsx จาก config.ini (เหมือน Planning._MASTER_MC_PATH)"""
    cfg = _cp.ConfigParser(interpolation=None)
    if not cfg.read(BASE_DIR / "config.ini", encoding="utf-8"):
        return None
    try:
        return Path(os.path.expandvars(cfg["paths"]["master_mc"]))
    except KeyError:
        return None


def _norm_gauge(g) -> str:
    """เกจ → '22' (ต้องตรงกับ _gkey ใน plan_view.py / plan_lock.py)"""
    try:
        return str(int(float(g)))
    except (TypeError, ValueError):
        return str(g).strip()


def _so_tokens(v) -> set:
    """SO_NO ('720188, 720189') → {'720188','720189'}

    ต้องตัด '.0' ท้ายเลข: คอลัมน์ SO_NO ที่มีช่องว่างบางแถวถูก pandas อ่านเป็น float
    ('720188.0') ส่วนชีท Lock_Plan อ่านด้วย openpyxl ได้ int ('720188') — ไม่ normalize
    ให้ตรงกัน ล็อกที่ระบุ SO จะจับคู่กับ booking ไม่ติดเลย
    """
    out = set()
    for t in re.split(r"[,\s;]+", str(v or "")):
        t = t.strip().upper()
        if not t or t in ("NAN", "NONE", "NAT"):
            continue
        if re.fullmatch(r"\d+\.0+", t):
            t = t.split(".")[0]
        out.add(t)
    return out


def _read_lock_sheet() -> pd.DataFrame | None:
    """ชีท Lock_Plan → DataFrame (หัวคอลัมน์ normalize แล้ว) ; ไม่มีไฟล์/ชีท = None

    ไม่มีล็อกเลยไม่ใช่ error — replan ทำงานตามเส้นเวลาปกติได้
    """
    p = _master_mc_path()
    if p is None or not p.exists():
        return None
    try:
        df = pd.read_excel(p, sheet_name=LOCK_PLAN_SHEET)
    except Exception:
        return None
    df.columns = [str(c).strip().upper().replace(" ", "_") for c in df.columns]
    if not {"WEEK", "ITEM_CODE"} <= set(df.columns):
        print(f"⚠️ ชีท {LOCK_PLAN_SHEET} ไม่มีคอลัมน์ WEEK/ITEM_CODE → ข้ามล็อกของผู้ใช้")
        return None
    return df


def _lock_source(v) -> str:
    """SOURCE ของแถวล็อก — ว่าง = BOOKING (แถวเก่าที่บันทึกก่อนมีคอลัมน์นี้)"""
    s = str(v or "").strip().upper()
    return "PLAN" if s == "PLAN" else "BOOKING"


def load_plan_locks(source: str = "BOOKING") -> dict:
    """ชีท Lock_Plan → {(week, item, mc_group, gauge): set(SO ที่ล็อก)}

    source = BOOKING → ล็อกที่ชี้ไปแถวใน booking DETAIL (ของที่จองเครื่องไปแล้ว)
    set ว่าง = ล็อกทุก SO ของ item นั้นในช่องนั้น (ช่อง SO_NO ในชีทเว้นว่างไว้)
    """
    df = _read_lock_sheet()
    if df is None:
        return {}

    want = _lock_source(source)
    out: dict = {}
    for _, r in df.iterrows():
        if _lock_source(r.get("SOURCE")) != want:
            continue
        wk = pd.to_numeric(r.get("WEEK"), errors="coerce")
        item = str(r.get("ITEM_CODE") or "").strip().upper()
        if pd.isna(wk) or not item or item == "NAN":
            continue
        # คีย์สัปดาห์ = YW ; คอลัมน์ YEAR ว่าง (ล็อกที่บันทึกก่อนมีคอลัมน์นี้) = ครั้งที่ใกล้วันนี้สุด
        key = (_timeline().key(int(wk), r.get("YEAR")), item,
               str(r.get("MC_GROUP") or "").strip().upper(),
               _norm_gauge(r.get("GUAGE")))
        out.setdefault(key, set()).update(_so_tokens(r.get("SO_NO")))
    return out


def load_plan_row_locks() -> list:
    """ล็อกฝั่ง "แถวในแผน" (SOURCE=PLAN) → [{key, so, qty, mc, sheet}]

    ต่างจากล็อก booking ตรงที่ต้องเก็บ QTY/MC/SHEET ไว้ด้วย:
      - SHEET บอกว่าแถวอยู่ไฟล์ไหน (PLAN = production_plan_* / REPLAN = replan_result)
      - QTY/MC เป็นค่าสำรองไว้ใช้เมื่อย้อนไปหาแถวจริงในไฟล์แผนไม่เจอ
    """
    df = _read_lock_sheet()
    if df is None:
        return []

    out = []
    for _, r in df.iterrows():
        if _lock_source(r.get("SOURCE")) != "PLAN":
            continue
        wk = pd.to_numeric(r.get("WEEK"), errors="coerce")
        item = str(r.get("ITEM_CODE") or "").strip().upper()
        if pd.isna(wk) or not item or item == "NAN":
            continue
        sheet = str(r.get("SHEET") or "").strip().upper().replace(" ", "_") or "PLAN"
        out.append({
            "key": (_timeline().key(int(wk), r.get("YEAR")), item,
                    str(r.get("MC_GROUP") or "").strip().upper(),
                    _norm_gauge(r.get("GUAGE"))),
            "so": _so_tokens(r.get("SO_NO")),
            "qty": pd.to_numeric(r.get("QTY"), errors="coerce"),
            "mc": pd.to_numeric(r.get("MC"), errors="coerce"),
            "sheet": sheet,
        })
    return out


def user_locked_mask(detail: pd.DataFrame, locks: dict) -> pd.Series:
    """แถวไหนใน DETAIL ตรงกับล็อกของผู้ใช้ → True

    SO ในล็อกว่าง = ทุก SO ผ่าน ; มี SO = แถวต้องมี SO นั้นอยู่ (SO_NO 1 ช่องมีหลายใบได้)
    """
    if not locks or detail.empty:
        return pd.Series(False, index=detail.index)

    wk = pd.to_numeric(_s(detail, "WEEK"), errors="coerce")
    item = _str_col(detail, "ITEM_CODE").str.upper()
    mcg = _str_col(detail, "MC_GROUP").str.upper()
    gauge = _s(detail, "GUAGE").map(_norm_gauge)
    so_raw = _str_col(detail, "SO_NO").str.upper()

    hit = []
    for i in detail.index:
        w = wk.at[i]
        if pd.isna(w):
            hit.append(False)
            continue
        want = locks.get((int(w), item.at[i], mcg.at[i], gauge.at[i]))
        if want is None:
            hit.append(False)
        elif not want:
            hit.append(True)                      # ล็อกทั้งช่อง (ทุก SO)
        else:
            hit.append(bool(want & _so_tokens(so_raw.at[i])))
    return pd.Series(hit, index=detail.index)


# ค่าที่ยังนับว่า "ไม่มีเลขใบจอง" (เซลล์ว่าง / NaN ที่กลายเป็นข้อความ / ศูนย์ที่ค้างจาก float)
KP_NO_BLANK = {"", "NAN", "NONE", "NAT", "-", "0", "0.0"}


def kp_locked_mask(df: pd.DataFrame) -> pd.Series:
    """แถวไหนมีเลขใบจอง (KP_NO) แล้ว → True = ไม่เอาเข้ารอบวางใหม่

    ใบที่ออกเลข KP แล้วคือของที่ยืนยันกับฝ่ายผลิตไปแล้ว ย้ายสัปดาห์/เครื่องไม่ได้
    จึงถูกจัดฝั่ง "ล็อก" เหมือนล็อกที่ผู้ใช้กดเอง (ผู้ใช้สั่ง 21 ก.ย. 2569)

    ⚠️ ช่อง KP_NO ของชีท DETAIL ยุบหลายใบไว้ด้วยกันคั่น ' | ' (AVA_MC._uniq_values)
       → ต้องแตกก่อน แล้วถือว่า "มีเลข" ถ้ามีตัวใดตัวหนึ่งไม่ว่าง
    ไฟล์ booking ที่สร้างก่อน 16 ก.ย. 2569 ยังไม่มีคอลัมน์นี้ → คืน False ทั้งคอลัมน์
    (= พฤติกรรมเดิม ไม่ใช่ล็อกทั้งไฟล์)
    """
    if df.empty or "KP_NO" not in df.columns:
        return pd.Series(False, index=df.index)

    def _has_kp(v) -> bool:
        return any(
            t.strip().upper() not in KP_NO_BLANK
            for t in str(v if v is not None else "").split("|")
        )

    return df["KP_NO"].map(_has_kp)


# =========================
# CORE ITEM: เข้ารอบ replan ตามปกติ (ยกเว้นตัวที่ user สั่งกันไว้)
# =========================
# core item = ผ้าสต๊อกที่ผลิตเป็นรอบตาม trigger ของสต๊อก (ดู Planning.py — CORE ITEM / batch)
#
# ⚠️ 16 ก.ย. 2569 (กลับคำจาก 15 ก.ย.): **เลิกล็อก core อัตโนมัติแล้ว** — ล็อกทั้งก้อนทำให้แผน core
# ไม่เคยถูกคิดใหม่ตามสต๊อกจริง ซึ่งขัดกับหัวใจของ core เอง ; core จึงเข้ารอบวางใหม่เหมือนงานทั่วไป
# = ยอดถูกยกมาเป็นดีมานด์ เครื่องถูกคืนเข้าพูล แล้ว Planning วางให้ตาม trigger สต๊อกของมันเอง
# โดยตรึงสัปดาห์เดิมจาก booking ตราบใดที่สต๊อกยังไม่ต่ำกว่า MIN (ดู CORE HOLD WEEK)
# สิ่งที่เคยหวังจากล็อก (เครื่องไม่หลุดมือระหว่างรอรอบผลิตถัดไป) ทำด้วยแถว "Hold รอ target"
# ในแผนแทน — กันเครื่องไว้ทุกสัปดาห์ที่เว้นระหว่าง 2 รอบ (Planning.py: CORE_HOLD_SOURCE)
# แถว AUTO-CORE ที่ค้างใน Lock_Plan ถูก program_lock.py กวาดออกให้เอง
# (กติกา "งานสียืมเครื่อง core แล้วคืน" ถูกลบไปพร้อมด่านซ่อมจัดทรงเมื่อ 2 ก.ย. — ไม่มีแล้ว)
#
# ประตูของ user (พลิกความหมายจากเดิม): คอลัมน์ "Core Status" ในทะเบียน — ใส่ค่าอะไรก็ได้
# = เบรก item ตัวนั้นไว้ที่แผนเดิม ไม่ให้เข้ารอบ replan ; เว้นว่าง = เข้าแผนตามปกติ
CORE_CUSTOMERS = {str(_c).strip() for _c in plan_config.get_list("KP_CORE_CUSTOMERS") if str(_c).strip()}
ITEMCORE_FALLBACK = BASE_DIR / "data" / "Itemcore" / "Itemcore.xlsx"

# คอลัมน์ธงของ user + ค่าที่ยังนับว่า "ไม่ได้ใส่" (เซลล์ว่าง/ขีด/ศูนย์ที่ค้างจากการพิมพ์)
CORE_STATUS_COL = "Core Status"
CORE_STATUS_BLANK = {"", "NAN", "NONE", "NAT", "-", "0", "0.0"}


def _core_status_filled(v) -> bool:
    """ช่อง Core Status ถือว่า "user ใส่ค่าแล้ว" ไหม"""
    return str(v or "").strip().upper() not in CORE_STATUS_BLANK


def load_core_registry():
    """ทะเบียน core item → (DataFrame, ชื่อไฟล์) ; (None, "") เมื่ออ่านไม่ได้

    แหล่งเดียวกับ Planning.py: ไฟล์ Target_Stock (config.ini → paths.target_stock) ซึ่ง user
    แก้บน server ได้ ; อ่านไม่ได้ค่อย fallback ไฟล์ Itemcore.xlsx เดิม
    อ่านทะเบียนไม่ได้ = replan งาน core ทั้งหมดตามปกติ (ไม่ล้มทั้งรอบ) แต่ต้องเตือนให้เห็น
    อ่านครั้งเดียวแล้ว cache — ทั้งด่านเบรก (Core Status) และการเติมลูกค้าให้ดีมานด์ใช้ตัวเดียวกัน
    """
    if getattr(load_core_registry, "_cache", None) is not None:
        return load_core_registry._cache

    cfg = _cp.ConfigParser(interpolation=None)
    src = None
    if cfg.read(BASE_DIR / "config.ini", encoding="utf-8"):
        try:
            p = Path(os.path.expandvars(cfg["paths"]["target_stock"]))
            src = p if p.exists() else None
        except KeyError:
            src = None
    if src is None and ITEMCORE_FALLBACK.exists():
        src = ITEMCORE_FALLBACK
    if src is None:
        print("⚠️ ไม่พบทะเบียน core item (target_stock / Itemcore.xlsx) → งาน core เข้ารอบ replan ทั้งหมด")
        load_core_registry._cache = (None, "")
        return load_core_registry._cache
    try:
        df = pd.read_excel(src)
    except Exception as e:
        print(f"⚠️ อ่านทะเบียน core item ({src.name}) ไม่ได้: {e} → งาน core เข้ารอบ replan ทั้งหมด")
        load_core_registry._cache = (None, "")
        return load_core_registry._cache

    df.columns = df.columns.str.strip()
    if "Item code" not in df.columns:
        print(f"⚠️ {src.name} ไม่มีคอลัมน์ 'Item code' → งาน core เข้ารอบ replan ทั้งหมด")
        load_core_registry._cache = (None, "")
        return load_core_registry._cache
    load_core_registry._cache = (df, src.name)
    return load_core_registry._cache


def load_core_item_teams() -> dict:
    """ทะเบียน core → {(ITEM_CODE, TEAM): ชื่อลูกค้า} สำหรับเติมลูกค้าให้ดีมานด์ booking

    ทะเบียนเก็บ "ทีม" คู่กับ "ลูกค้า" อยู่แล้ว (RTS→ที คัลเจอร์, NYK1-1→CENTER DOMESTIC)
    ส่วน booking มี TEAM_NAME ทุกแถว → ใช้ทีมย้อนกลับไปหาลูกค้าได้เมื่อ KNIT_CUS_NAME ว่าง
    """
    df, _ = load_core_registry()
    if df is None or "Team Name" not in df.columns or "Customer" not in df.columns:
        return {}
    out = {}
    for it, tm, cu in zip(
        df["Item code"].astype(str).str.strip().str.upper(),
        df["Team Name"].astype(str).str.strip().str.upper(),
        df["Customer"].astype(str).str.strip(),
    ):
        if it and it not in ("NAN", "NONE") and tm and tm not in ("NAN", "NONE") and cu in CORE_CUSTOMERS:
            out[(it, tm)] = cu
    return out


def load_core_item_customers() -> dict:
    """ทะเบียน core → {ITEM_CODE: เซ็ตลูกค้า core ของผ้านั้น}

    ใช้เติมลูกค้าชั้นที่สอง เมื่อทีมของ booking ยุบหลายค่าจนย้อนหาลูกค้าไม่ได้:
    ผ้าที่ทะเบียนมีลูกค้า core คนเดียว → เติมได้โดยไม่ต้องเดา (ข้อมูลจริง 47/61 ผ้าเป็นแบบนี้)
    ผ้าที่มีสองลูกค้า (RTS + NYK1-1 คนละเป้า คนละก้อนสต๊อก) → ห้ามเดา ปล่อยว่างไว้
    """
    df, _ = load_core_registry()
    if df is None or "Customer" not in df.columns:
        return {}
    out: dict = {}
    for it, cu in zip(
        df["Item code"].astype(str).str.strip().str.upper(),
        df["Customer"].astype(str).str.strip(),
    ):
        if it and it not in ("NAN", "NONE") and cu in CORE_CUSTOMERS:
            out.setdefault(it, set()).add(cu)
    return out


def load_core_hold_items() -> set:
    """core item ที่ user สั่ง "กันไว้ที่แผนเดิม" → เซ็ต ITEM_CODE ตัวใหญ่ (ปกติว่าง)"""
    df, src_name = load_core_registry()
    if df is None:
        return set()
    # ต้องกรอง Customer ด้วย — กฎ CORE ITEM ใน Planning ใช้ item + customer คู่กัน
    # (item ที่อยู่ในทะเบียนแต่ลูกค้าไม่ตรง = งานสั่งปกติ ไม่ใช่ core)
    cust = (df["Customer"].astype(str).str.strip() if "Customer" in df.columns
            else pd.Series("", index=df.index))
    # ไม่มีคอลัมน์ Core Status ในไฟล์ = ไม่มีใครสั่งเบรก → core เข้ารอบ replan ทั้งหมด
    has_flag_col = CORE_STATUS_COL in df.columns
    flag = (df[CORE_STATUS_COL] if has_flag_col else pd.Series([None] * len(df), index=df.index))

    hold, n_core = {}, 0
    for it, cu, fl in zip(df["Item code"].astype(str).str.strip().str.upper(), cust, flag):
        if not it or it in ("NAN", "NONE") or cu not in CORE_CUSTOMERS:
            continue
        n_core += 1
        if _core_status_filled(fl):
            # item เดียวมีได้หลายแถว (คนละลูกค้า/ประเภท) — ใส่ธงแถวไหนก็ถือว่าเบรกทั้ง item
            # เพราะแผน booking แยกตาม ITEM_CODE ไม่ได้แยกตามลูกค้า
            hold[it] = str(fl).strip()

    msg = f"📘 ทะเบียน core item: {n_core} แถว จาก {src_name} → เข้ารอบ replan ตามปกติ"
    if hold:
        _sample = ", ".join(f"{k}={v}" for k, v in list(hold.items())[:10])
        msg += (f" | ยกเว้นที่ user เบรกไว้ด้วย {CORE_STATUS_COL} {len(hold)} รายการ: "
                + _sample + (" ..." if len(hold) > 10 else ""))
    print(msg)
    return set(hold)


def _pick_core_customer(v) -> str:
    """ชื่อลูกค้าที่ AVA_MC ยุบไว้ด้วย ' | ' → เลือกชื่อลูกค้า core ถ้ามีอยู่ในนั้น

    ⚠ ขาดตัวนี้ไม่ได้เมื่อ core เข้ารอบ replan: Planning ตัดสิน CORE ITEM จากคู่ (ผ้า, ลูกค้า)
    ที่ต้องอยู่แถวเดียวกันในทะเบียน (Planning.itemcore_pairs) — ค่าที่ยุบเป็น 'A | B' จะไม่ตรง
    ทะเบียนสักชื่อ งาน core จึงกลายเป็นงานสั่งปกติที่วางตาม RDD = รอบผลิตสต๊อกพัง
    ชีท DETAIL_SC คีย์ด้วย SO อยู่แล้วจึงมักมีชื่อเดียว — ตัวนี้กันกรณี fallback ไปใช้ DETAIL
    (ไฟล์ booking ที่ไม่มี DETAIL_SC) ที่ยุบทุกใบของบล็อกเข้าด้วยกัน
    ไม่มีลูกค้า core อยู่ในนั้น = คืนค่าเดิมทุกตัวอักษร (งานทั่วไปต้องไม่ถูกแตะ)
    """
    s = str(v or "").strip()
    if " | " not in s:
        return s
    for part in s.split(" | "):
        if part.strip() in CORE_CUSTOMERS:
            return part.strip()
    return s


def core_item_mask(detail: pd.DataFrame, core_items: set) -> pd.Series:
    """แถวไหนใน DETAIL เป็น core item ที่ถูกเบรกไว้ → True (จับด้วย ITEM_CODE อย่างเดียว —
    booking DETAIL ไม่มีคอลัมน์ลูกค้า ; การกรองลูกค้าทำตอนสร้างทะเบียนแล้ว)"""
    if not core_items or detail.empty:
        return pd.Series(False, index=detail.index)
    return _str_col(detail, "ITEM_CODE").str.upper().isin(core_items)


# ── core → แถวในชีทแผน (ให้แก้บนหน้าเว็บได้) ─────────────────────────────────
# ไม่จัดใหม่ ≠ แก้ไม่ได้ — planner ยังต้องลากเปลี่ยนสัปดาห์/แก้จำนวนเองได้ จึงต้องแปะ core
# กลับเป็น "แถวแผน" ไม่ใช่ปล่อยให้เป็นบล็อก booking อ่านอย่างเดียวบน Gantt
#
# เครื่องไม่ถูกนับซ้ำเพราะตั้ง MC_BOOKING = ACTUAL_MC → หน้าเว็บคิดเครื่องที่แผนกินพูลจริง
# = ACTUAL_MC − MC_BOOKING = 0 (📦 นั่งบนเครื่องที่วิ่งอยู่แล้ว ดู PlanGantt.planMcAgg)
# ส่วนเครื่องจริงถูกกันไว้แล้วใน SUMMARY_MC_REMAIN และโควตา setup ถูกกินผ่าน DETAIL
CORE_PIN_FLAG = "_CORE_PINNED"          # ธงใน DETAIL — plan_view ใช้ซ่อนบล็อกที่ซ้อนกับแถวแผน
CORE_PLAN_SOURCE = "CORE_KEEP"
CORE_REMARK = "📘 core item — คงแผนเดิม ไม่ถูกจัดใหม่ (ลาก/แก้บนหน้าแผนได้)"

# ใบที่ออกเลข KP แล้ว — ใช้ทางเดียวกับ core แต่ล็อกแน่นกว่า: ยืนยันกับฝ่ายผลิตไปแล้ว
# จึงห้ามลาก/แก้บนหน้าเว็บด้วย (หน้าเว็บดูที่ PLAN_SOURCE นี้ — PlanGantt: KEEP_LOCKED_SOURCE)
# ⚠️ ไม่ติด IS_CORE_ITEM ให้แถวพวกนี้ ต่อให้ item เป็น core จริง — ธงนั้นเปลี่ยนพฤติกรรมหน้าเว็บ
# (ปุ่มกรอง core, ยกเว้นคาดแดง "ทอไม่ทัน") ซึ่งไม่ใช่สิ่งที่การออกใบ KP หมายถึง
KP_PLAN_SOURCE = "KP_KEEP"
KP_REMARK = "🎫 ออกใบจอง KP แล้ว — คงแผนเดิม ไม่ถูกจัดใหม่ (ลาก/แก้ไม่ได้)"


def _color_part_of_orders(so_col: pd.Series, item_col: pd.Series) -> pd.Series:
    """ธง COLOR_PART ของแถวที่ยกมาจาก booking ตรง ๆ (ไม่ผ่านลูปวางของ Planning)

    เกณฑ์เดียวกับ _main_row_is_color ใน Planning.py: ใบสั่งนั้นมียอดเปิดสีไหม
    ("Yes" = มี, ว่าง = ไม่มี) — คิดต่อใบสั่ง ไม่ใช่ต่อผ้า เพราะใบที่ยังไม่เปิดสีเลย
    ของผ้าที่ใบอื่นเปิดสีไว้ ไม่ใช่งานสี (user 27 ส.ค. 2569)

    ช่อง SC ที่ยุบหลายใบด้วยคอมมาต้องแตกก่อนเทียบ (ดู ColorOpen.split_sc)
    ยังไม่มีไฟล์ view → quota ว่าง → ทุกแถวได้ธงว่าง เหมือนพฤติกรรมเดิมทุกประการ
    """
    _q = ColorOpen.quota_by_sc_item()
    if not _q:
        return pd.Series("", index=so_col.index)
    return pd.Series(
        [
            "Yes" if any(_q.get((_s, ColorOpen.norm_item(_it)), 0.0) > 0
                         for _s in ColorOpen.split_sc(_so)) else ""
            for _so, _it in zip(so_col, item_col)
        ],
        index=so_col.index,
    )


def core_plan_rows(core_detail: pd.DataFrame, plan_cols: list = None, year_of_week=None,
                   source: str = CORE_PLAN_SOURCE, remark: str = CORE_REMARK,
                   is_core: bool = True) -> pd.DataFrame:
    """booking DETAIL_SC ของ core → แถวชีท PLAN (1 แถว = 1 SC)

    plan_cols = None → คืนทุกคอลัมน์ที่สร้างได้ (ตอนเตรียมก่อนรัน Planning ยังไม่รู้คอลัมน์ของแผน
    — Planning จัดให้ตรงคอลัมน์แผนเองตอนแปะ)

    ใส่ครบ 3 กลุ่ม:
      1. เครื่อง/เวลา — NEW_MC / SETUP_DAYS / ACTUAL_WORKING_DAYS / DAILY_CAPACITY / MC_BOOKING
         (หน้าเว็บใช้คำนวณตอนลาก) ; ค่าเครื่อง/setup เป็นของ "ทั้งบล็อก" → ใบที่ยอดมากสุด
         ถือไว้ใบเดียว ใบอื่นในบล็อกเดียวกันเขียน 0 + MC_SHARED (ดูคอมเมนต์ตรงที่คิด _owner)
      2. ข้อมูลออร์เดอร์ของใบนั้น — ORDERS_QTY / FG_WEEK / RDD_WEEK / ORDER_TYPE / ลูกค้า / PO
         เพื่อให้การ์ดรายละเอียดอ่านได้เท่ากับงานแผนปกติ และเห็นกำหนดส่งของแต่ละใบจริง ๆ
      3. ธง core — IS_CORE_ITEM / PLAN_SOURCE=CORE_KEEP : คงแผนเดิม ไม่ถูกจัดใหม่ ไม่ยึด target

    source / remark / is_core = สลับให้ใช้กับงาน "คงแผนเดิม" ชนิดอื่นที่มาจาก booking DETAIL_SC
    เหมือนกัน — ตอนนี้คือใบที่ออกเลข KP แล้ว (KP_PLAN_SOURCE, ไม่ติดธง core)

    ที่ยังเว้นว่าง: PLAN_QTY (ยอดคงเหลือระหว่างวางแผน — แถว booking ไม่ได้ผ่านการวางแผน)
    และ TARGET_KNIT ที่ Planning เติมให้ทีหลัง (มีปฏิทิน + offset ตาม FOB_TYPE อยู่ที่นั่น)
    """
    if core_detail.empty:
        return pd.DataFrame(columns=plan_cols)

    d = core_detail
    ceil = pd.to_numeric(_s(d, "MC_USE_CEIL"), errors="coerce").fillna(0).astype(int)
    inc = pd.to_numeric(_s(d, "_MC_INCREASE"), errors="coerce").fillna(0).astype(int)
    is_new = _s(d, "_IS_NEW_SETUP").fillna(True).astype(bool)
    # สูตรเดียวกับที่ Planning ใช้นับ job จาก DETAIL (_is_new_setup=True → ทุกเครื่องตั้งใหม่)
    new_mc = ceil.where(is_new, inc).clip(lower=0)
    setup_per_mc = pd.to_numeric(_s(d, "_SETUP_DAYS"), errors="coerce").fillna(0)
    # ⚠️ SETUP_DAYS ในชีทแผนเป็น "วันรวมของทุกเครื่องที่ตั้งใหม่" (หน้าเว็บหารด้วย NEW_MC
    # กลับเป็นวัน/เครื่อง) ส่วน booking เก็บเป็นวัน/เครื่อง → ต้องคูณก่อน ไม่งั้น setup หด
    setup_days = (setup_per_mc * new_mc).where(new_mc > 0, 0)
    wd = pd.to_numeric(
        _s(d, "_EFFECTIVE_WORKING_DAYS") if "_EFFECTIVE_WORKING_DAYS" in d.columns else _s(d, "WORKING_DAY"),
        errors="coerce",
    ).fillna(0)

    week = pd.to_numeric(_s(d, "WEEK"), errors="coerce")
    # PLAN_YEAR ต้องมีเสมอ — แกน Gantt เก็บแค่เลขสัปดาห์เปล่า ถ้าไม่มีปี งาน W1 ของปีหน้า
    # จะถูกตัดสินเป็น "สัปดาห์อดีตของปีนี้" แล้วไปโผล่ในแถบ freeze (ดู PlanGantt yearOfWeek)
    # WEEK ใน main เป็น YW แล้ว → ปีอยู่ในตัว ; ค่าเลขเปล่า (ข้อมูลเก่า) ค่อยพึ่ง YEAR / year_of_week
    if "YEAR" in d.columns:
        year = pd.to_numeric(_s(d, "YEAR"), errors="coerce")
    else:
        year = pd.Series([pd.NA] * len(d), index=d.index, dtype="object")
    year = pd.Series(
        [yw_year(w) if pd.notna(w) and yw_year(w) is not None
         else (y if pd.notna(y) else (year_of_week(int(w)) if (year_of_week is not None and pd.notna(w)) else pd.NA))
         for y, w in zip(year, week)],
        index=d.index,
    )
    week = week.map(lambda w: yw_week(w) if pd.notna(w) else w)

    # ── ข้อมูลออร์เดอร์ของใบนั้น (DETAIL_SC แตกมาให้แล้ว 1 แถว = 1 SC) ──────────
    # ยอดของออร์เดอร์ = ผลรวมทุกสัปดาห์ของ (SO, item) ในชุด core นี้ ไม่ใช่ยอดรายสัปดาห์ —
    # การ์ดบนหน้าแผนเทียบ ORDERS_QTY กับผลรวม PRODUCE_QTY ต่อออร์เดอร์ ถ้าใส่ยอดรายสัปดาห์
    # ทุกแถวจะขึ้นว่า "สั่งเท่ากับที่วาง" ทั้งที่ใบนั้นกระจายอยู่หลายสัปดาห์
    _so = _str_col(d, "SO_NO")
    _item_c = _str_col(d, "ITEM_CODE")
    _qty = pd.to_numeric(_s(d, "KP_WEIGHT"), errors="coerce").fillna(0)
    _order_qty = _qty.groupby([_so, _item_c]).transform("sum")
    # FG week เก็บเป็น YYYYWW เหมือนแถวแผนปกติ (Planning ใช้ตัวนี้คิด TARGET_KNIT ให้ทีหลัง)
    # ไม่มี FG → ใช้ RDD แทน (user 24/9/2569 — กติกาเดียวกับ build_booking_demand) ไม่งั้นแถว
    # core/KP ไม่มี Target บนการ์ด ; ติดป้าย DUE_BASIS='RDD' ให้แยกจาก FG ต้นทาง ; RDD_WEEK ไม่แตะ
    _fg_y = pd.to_numeric(_s(d, "FG_YEAR"), errors="coerce")
    _fg_w = pd.to_numeric(_s(d, "FG_WEEK"), errors="coerce")
    _fg = (_fg_y * 100 + _fg_w).where(_fg_y.notna() & _fg_w.notna())
    _rdd = pd.to_numeric(_s(d, "RDD_WEEK"), errors="coerce")
    _fg_by_rdd = _fg.isna() & _rdd.between(100001, 999953)
    _fg = _fg.where(~_fg_by_rdd, _rdd)
    # ORDER_TYPE มีที่เดียวคือ FR_CODE ('S{SC}-{ORDER_TYPE}') — กติกาเดียวกับ build_booking_demand
    _otype = _str_col(d, "FR_CODE").map(lambda v: _fr_parse(v)[1])

    # ── หลายใบสั่งในก้อนงานเดียว = เครื่องชุดเดียวกัน ไม่ใช่ชุดละใบ ─────────────────
    # DETAIL_SC แตก 1 แถวต่อ 1 ใบสั่ง แต่ MC_USE / MC_USE_CEIL / _SETUP_DAYS เป็นค่าของ
    # "ทั้งก้อน" ที่ถูกคัดลอกซ้ำทุกใบ (W41 FD3BASKD21/03A0: 354.95 + 5,495.05 กก. = 4 เครื่อง
    # แต่ DETAIL_SC เขียน CEIL 4 ทั้งสองแถว) ทุกที่ที่นับเครื่องใช้สูตร "max ภายในใบสั่ง →
    # บวกข้ามใบสั่ง" (_pool_blocks_from_plan/Planning.py, pinned_mc_use, planMcAgg+foldJobs
    # /PlanGantt.jsx) ถ้าแปะเลขซ้ำลงชีทแผนจะอ่านได้ 4+4 = 8 → Gantt โชว์เครื่องเป็นเท่าตัว
    # และเครื่องว่างของพูลถูกหักเกิน (เคสจริง 23 ก.ย. 2569: 5 ก้อน หักเกินรวม 18 ตัว)
    # → เขียนแบบเดียวกับที่ Planning เขียนแถวที่ใช้เครื่องร่วมกัน: ใบที่ยอดมากสุดถือเครื่อง
    #   ของทั้งก้อน ใบที่เหลือ ACTUAL_MC/NEW_MC/CARRYOVER/MC_BOOKING/SETUP_DAYS = 0
    #   แล้วประกาศว่าเกาะเครื่องชุดนั้นด้วย MC_SHARED = เครื่องของก้อน (สูตรบวกจึงไม่ซ้ำ)
    _mcg = _str_col(d, "MC_GROUP")
    _gauge = _s(d, "GUAGE").map(_norm_gauge)
    _blk_key = (week.astype(str) + "|" + year.astype(str) + "|" + _item_c.str.upper()
                + "|" + _mcg.str.upper() + "|" + _gauge.astype(str))
    _owner = pd.Series(False, index=d.index)
    _owner.loc[_qty.groupby(_blk_key).idxmax().to_numpy()] = True
    _shared_rows = int((~_owner).sum())
    if _shared_rows:
        print(f"ℹ️ {source}: ใบสั่งที่ใช้เครื่องชุดเดียวกับใบอื่น {_shared_rows} แถว → "
              f"ให้ใบที่ยอดมากสุดถือเครื่องของก้อน ที่เหลือเป็น MC_SHARED (กันนับเครื่องซ้ำ)")
    ceil_own = ceil.where(_owner, 0)
    new_own = new_mc.where(_owner, 0)

    out = pd.DataFrame(index=range(len(d)))
    src = {
        "ITEM_CODE": _item_c,
        "SC_SO_NO": _so,
        "ORDERS_QTY": _order_qty,
        "PENDING_PLAN": _order_qty,     # booking = ของที่ยังไม่ได้ทอ → ค้างเท่ากับยอดใบ
        "ORDER_TYPE": _otype,
        "ORDER_DATE": _s(d, "PO_IN_DATE").map(_last_po_in_date),
        "FG_WEEK": _fg,
        "RDD_WEEK": pd.to_numeric(_s(d, "RDD_WEEK"), errors="coerce"),
        "CUSTOMER": _str_col(d, "KNIT_CUS_NAME"),
        "PO_NO": _str_col(d, "PO_NO"),
        "KP_NO": _str_col(d, "KP_NO"),      # เลขใบจองของงาน core ที่ยกจาก booking ลงแผนตรง ๆ
        "SC_LINE_ID": _s(d, "SC_LINE_ID").map(_line_text),
        "SUB_COLOR": _str_col(d, "COLOR"),
        "NAY_COLOR": _str_col(d, "NAY_COLOR"),
        "MC_GROUP": _mcg,
        "MC_GUAGE": _gauge,
        "FACTORY_TYPE": _str_col(d, "FACTORY"),
        "PLAN_WEEK": week,
        "PLAN_YEAR": year,
        "PRODUCE_QTY": _qty,
        # setup/เครื่อง = ของ "ทั้งก้อน" → ใบที่ถือก้อนเท่านั้นที่เขียน ใบร่วมเครื่องเขียน 0
        # (REQUIRED_MC คงค่าต้นทางไว้ทุกใบ — เป็นตัวเลข "ก้อนนี้ต้องใช้กี่เครื่อง" ที่การ์ดใช้แสดง
        #  ไม่มีใครเอาไปรวมเป็นเครื่องของพูล)
        "SETUP_DAYS": setup_days.where(_owner, 0),
        "REQUIRED_MC": ceil,
        "ACTUAL_MC": ceil_own,
        "NEW_MC": new_own,
        "CARRYOVER_MC": (ceil_own - new_own).clip(lower=0),
        "MC_BOOKING": ceil_own,             # = ACTUAL_MC → ไม่กินเครื่องเพิ่มจากพูล
        "MC_SHARED": ceil.where(~_owner, 0),
        "FACTORY_WORKING_DAYS": wd,
        "ACTUAL_WORKING_DAYS": wd,
        "AVAILABLE_DAYS": wd,
        "DAILY_CAPACITY": pd.to_numeric(_s(d, "CAP_KNIT"), errors="coerce"),
        "CAT": _str_col(d, "CAT"),
        "TEAM": _str_col(d, "TEAM_NAME"),
        "MATERIAL_CONTENT": _str_col(d, "MATERIAL_CONTENT"),
        "YARN_USED": _str_col(d, "YARN_USED"),
        "PO_IN_DATE": _s(d, "PO_IN_DATE"),
        "IS_CORE_ITEM": True if is_core else "",
        "PLAN_SOURCE": source,
        "REMARK": remark,
        # ธงก้อนสี — แถวพวกนี้ไม่ผ่านลูปวางของ Planning จึงไม่มีใครแจกธงให้ ถ้าปล่อยว่าง
        # หน้าเว็บจะระบายเป็น "ยังไม่ให้สี" ทั้งที่ใบนั้นเปิดสีมาแล้ว (เคสจริง 48 แถว KP
        # 29 ก.ย. 2569) ขณะที่ฝั่ง Planning มองว่าเป็นงานสีอยู่แล้วผ่าน _main_row_is_color
        # → เขียนธงด้วยเกณฑ์เดียวกันเป๊ะ: ใบนั้นมียอดเปิดสีไหม (ผ้า core ไม่เข้ากติกาสี)
        "COLOR_PART": pd.Series("", index=d.index) if is_core else _color_part_of_orders(_so, _item_c),
        "DUE_BASIS": pd.Series("", index=d.index).where(~_fg_by_rdd, "RDD"),
    }
    for col in (plan_cols if plan_cols is not None else list(src)):
        v = src.get(col)
        out[col] = (v.to_numpy() if isinstance(v, pd.Series) else v) if v is not None else pd.NA
    return out


# =========================
# LOCK_PLAN ฝั่งแผน: ตรึงงานที่ user กดล็อกไว้ให้อยู่สัปดาห์เดิม
# =========================
# ล็อก SOURCE=PLAN ไม่ได้ชี้ไปที่ booking แต่ชี้ไปที่ "แถวในไฟล์แผน" → ต้องทำ 3 อย่าง
#   1. ย้อนไปอ่านแถวจริงจากไฟล์แผน (ไฟล์ไหนดูจากคอลัมน์ SHEET ของล็อก)
#   2. หักยอดของแถวนั้นออกจากดีมานด์ SC ไม่ให้ Planning วางซ้ำอีกก้อน
#   3. กันเครื่องของสัปดาห์นั้นไว้ แล้วแปะแถวเดิมกลับเข้าผลลัพธ์
PLAN_SHEET = "PLAN"


def _latest_plan_file() -> Path | None:
    """ไฟล์ production_plan_*.xlsx ที่ใหม่สุด (ไฟล์ที่หน้าเว็บเปิดอยู่ตอนกดล็อก)"""
    cands = sorted(DATA_PLAN_DIR.glob("production_plan_*.xlsx"),
                   key=lambda p: p.stat().st_mtime, reverse=True)
    return cands[0] if cands else None


def _plan_source_file(sheet: str) -> Path | None:
    """ชีทที่กดล็อก → ไฟล์ที่ต้องย้อนไปอ่านแถวจริง

    ชีท PLAN กับ REPLAN บนหน้าเว็บอ่านคนละไฟล์ (ดู plan_view.py) — ถ้าอ่านผิดใบจะหา
    แถวไม่เจอทั้งที่ล็อกถูกต้อง ; REPLAN = ผลรอบก่อน (ไฟล์นี้จะถูกเขียนทับตอนจบรอบนี้)
    """
    if sheet == "REPLAN":
        return RESULT_FILE if RESULT_FILE.exists() else None
    return _latest_plan_file()


def load_pinned_plan_rows(locks: list) -> pd.DataFrame:
    """ล็อกฝั่งแผน → แถวจริงจากไฟล์แผน (คอลัมน์ครบตามชีท PLAN)

    จับคู่ด้วย (PLAN_WEEK, ITEM_CODE, MC_GROUP, MC_GUAGE) + SC_SO_NO ถ้าล็อกระบุใบไว้
    ล็อกที่หาแถวไม่เจอ (แผน gen ใหม่แล้วงานย้ายสัปดาห์) → เตือนแล้วข้าม ไม่ตรึงมั่ว
    """
    if not locks:
        return pd.DataFrame()

    by_sheet: dict = {}
    for lk in locks:
        by_sheet.setdefault(lk["sheet"], []).append(lk)

    frames, missed = [], []
    for sheet, group in by_sheet.items():
        src = _plan_source_file(sheet)
        if src is None:
            print(f"⚠️ ล็อกแผน {len(group)} รายการอ้างชีท {sheet} แต่ไม่พบไฟล์แผน → ข้าม")
            missed += group
            continue
        try:
            plan = pd.read_excel(src, sheet_name=PLAN_SHEET)
        except Exception as e:
            print(f"⚠️ อ่าน {src.name} ชีท {PLAN_SHEET} ไม่ได้ ({e}) → ข้ามล็อก {len(group)} รายการ")
            missed += group
            continue
        plan.columns = plan.columns.str.strip().str.upper()
        # สัปดาห์ของแถวแผน → YW (ไฟล์แผนแยก PLAN_WEEK/PLAN_YEAR) ให้เทียบกับคีย์ล็อกได้ตรงปี
        _to_yw(plan, "PLAN_WEEK", "PLAN_YEAR")
        _to_yw(plan, "TARGET_KNIT", "TARGET_YEAR")

        wk = pd.to_numeric(_s(plan, "PLAN_WEEK"), errors="coerce")
        item = _str_col(plan, "ITEM_CODE").str.upper()
        mcg = _str_col(plan, "MC_GROUP").str.upper()
        gauge = _s(plan, "MC_GUAGE").map(_norm_gauge)
        so_raw = _str_col(plan, "SC_SO_NO").str.upper()

        for lk in group:
            w, it, g_mc, g_g = lk["key"]
            m = (wk == w) & (item == it) & (mcg == g_mc) & (gauge == g_g)
            if lk["so"]:
                m &= so_raw.map(lambda s, want=lk["so"]: bool(want & _so_tokens(s))).astype(bool)
            if not m.any():
                missed.append(lk)
                continue
            frames.append(plan[m].copy())
        print(f"📄 ล็อกแผนชีท {sheet}: อ่านแถวจาก {src.name}")

    if missed:
        print(f"⚠️ ล็อกแผนที่หาแถวในไฟล์แผนไม่เจอ {len(missed)} รายการ "
              f"(แผนถูกวางใหม่แล้วงานย้ายสัปดาห์/เครื่อง?) → ไม่ถูกตรึง: "
              + ", ".join(f"{yw_label(k['key'][0])} {k['key'][1]} {k['key'][2]} G{k['key'][3]}"
                          for k in missed[:10])
              + (" ..." if len(missed) > 10 else ""))
    if not frames:
        return pd.DataFrame()

    pinned = pd.concat(frames, ignore_index=True)
    # ล็อกหลายรายการชี้แถวเดียวกันได้ (เช่นล็อกทั้งช่องแล้วล็อกเจาะ SO ซ้ำ) — ตัดซ้ำทิ้ง
    # ไม่งั้นยอดถูกหักออกจากดีมานด์สองเท่า และเครื่องถูกกันเกินจริง
    pinned = pinned.drop_duplicates()
    print(f"📌 ตรึงงานในแผนไว้ {len(pinned)} แถว "
          f"รวม {pd.to_numeric(_s(pinned, 'PRODUCE_QTY'), errors='coerce').fillna(0).sum():,.0f} กก.")
    return pinned


def _mc_cat_pool_maps(booking_src: Path) -> tuple:
    """(MC_GROUP, GUAGE) → MC_CAT และ → MC_POOL

    MC_CAT มาจากชีทแรกของ MasterMC (คอลัมน์ MC / Guage / MC_CAT) — เป็นแหล่งเดียวกับที่
    Planning ใช้ ; MC_POOL มาจากชีท MC_POOL_MAP ของไฟล์ booking (มีเฉพาะเครื่องที่อยู่ในพูล)
    ต้องได้ทั้งคู่เพราะ SUMMARY_MC_REMAIN แยกเป็นแถวพูล (จับด้วย MC_POOL) กับแถวปกติ
    (จับด้วย MC_CAT + GUAGE) — จับผิดฝั่ง เครื่องที่กันไว้จะไปหักผิดแถว
    """
    cat_map: dict = {}
    p = _master_mc_path()
    if p is not None and p.exists():
        try:
            mm = pd.read_excel(p)
            mm.columns = mm.columns.str.strip()
            _cat_col = "MC_CAT" if "MC_CAT" in mm.columns else "Type_1"
            if {"MC", "Guage"} <= set(mm.columns) and _cat_col in mm.columns:
                for _, r in mm.iterrows():
                    mc = str(r.get("MC") or "").strip().upper()
                    if mc:
                        cat_map[(mc, _norm_gauge(r.get("Guage")))] = str(r.get(_cat_col) or "").strip()
        except Exception as e:
            print(f"⚠️ อ่าน MasterMC เพื่อหา MC_CAT ไม่ได้ ({e}) → เครื่องของงานที่ตรึงอาจกันไม่ครบ")

    pool_map: dict = {}
    try:
        pm = pd.read_excel(booking_src, sheet_name="MC_POOL_MAP")
        pm.columns = pm.columns.str.strip().str.upper()
        for _, r in pm.iterrows():
            mc = str(r.get("MC_GROUP") or "").strip().upper()
            if mc:
                pool_map[(mc, _norm_gauge(r.get("GUAGE")))] = str(r.get("MC_POOL") or "").strip()
    except Exception:
        pass                                   # ไฟล์เก่าไม่มีชีทนี้ = ไม่มีพูล ใช้ MC_CAT ล้วน
    return cat_map, pool_map


def pinned_mc_use(pinned: pd.DataFrame, booking_src: Path):
    """เครื่องที่งาน "ตรึง" กินไว้ → (ยอดต่อพูล, ยอดต่อ non-pool) รูปเดียวกับ _locked_mc_use

    ต้องกันทุกสัปดาห์ ไม่ใช่เฉพาะสัปดาห์ที่ปลดล็อก — งานในแผนไม่เคยถูกหักออกจาก
    booking capacity มาก่อน (Planning นับเครื่องของแผนเองในหน่วยความจำ) ถ้าไม่กันไว้
    Planning รอบ replan จะแจกเครื่องพวกนี้ให้งานอื่นทั้งที่งานเดิมยังใช้อยู่

    ⚠️ 1 บล็อกงาน = (สัปดาห์, item, กลุ่มเครื่อง, เกจ) อาจมีหลายแถว → ยุบ 2 ชั้น
      • แถวซ้ำของ "ใบสั่งเดียวกัน" เขียนเครื่องชุดเดียวซ้ำทุกแถว → max
        (F6DRYTECH201/05A0 IBLTA G22 W37 สองแถวเขียน ACTUAL_MC=8 ทั้งคู่ = carry 6 + ใหม่ 2)
      • "คนละใบสั่ง" = คนละชุดเครื่อง → บวก ; ใบที่ใช้เครื่องร่วมกับใบอื่นเขียนเฉพาะส่วนที่เพิ่ม
        (ที่เหลืออยู่ใน MC_SHARED) จึงบวกได้โดยไม่ซ้ำ
    = สูตรเดียวกับ _pool_blocks_from_plan ใน Planning.py และ planMcAgg ใน PlanGantt.jsx
    (เดิมใช้ max ทั้งบล็อก → แถว 2+1+0 ของ 3 ใบอ่านได้ 2 ทั้งที่บล็อกใช้ 3 = กันเครื่องขาด)
    """
    empty = ({}, {})
    if pinned.empty:
        return empty
    cat_map, pool_map = _mc_cat_pool_maps(booking_src)

    d = pinned.copy()
    d["_WK"] = pd.to_numeric(_s(d, "PLAN_WEEK"), errors="coerce")
    d["_MC"] = pd.to_numeric(_s(d, "ACTUAL_MC"), errors="coerce").fillna(0)
    d = d[d["_WK"].notna() & (d["_MC"] > 0)]
    if d.empty:
        return empty
    d["_G"] = _s(d, "MC_GUAGE").map(_norm_gauge)
    d["_MCG"] = _str_col(d, "MC_GROUP").str.upper()
    d["_ITEM"] = _str_col(d, "ITEM_CODE").str.upper()

    d["_SC"] = _str_col(d, "SC_SO_NO")

    # ยุบหลายแถวของบล็อกเดียวกันให้เหลือ "จำนวนเครื่องของบล็อก" ก่อน แล้วค่อยบวกข้ามบล็อก
    # (max ภายในใบสั่ง → บวกข้ามใบสั่ง)
    blocks = (
        d.groupby(["_WK", "_ITEM", "_MCG", "_G", "_SC"], as_index=False)["_MC"].max()
        .groupby(["_WK", "_ITEM", "_MCG", "_G"], as_index=False)["_MC"].sum()
    )
    _dup = len(d) - len(blocks)
    if _dup:
        print(f"ℹ️ งานที่ตรึงไว้: ยุบแถวซ้ำของบล็อกเดียวกัน {_dup} แถว "
              f"(เครื่องนับตามบล็อก {int(d['_MC'].sum())} → {int(blocks['_MC'].sum())} ตัว)")

    pool: dict = {}
    nonpool: dict = {}
    for _, r in blocks.iterrows():
        k = (r["_MCG"], r["_G"])
        w = int(r["_WK"])
        pk = pool_map.get(k, "")
        if pk:
            pool[(pk, w)] = pool.get((pk, w), 0.0) + float(r["_MC"])
        else:
            cat = cat_map.get(k, "")
            if not cat:
                print(f"⚠️ ไม่รู้จัก MC_CAT ของ {r['_MCG']} เกจ {r['_G']} → เครื่องของงานที่ตรึง "
                      f"{yw_label(w)} ไม่ถูกกันไว้")
                continue
            nonpool[(cat, r["_G"], w)] = nonpool.get((cat, r["_G"], w), 0.0) + float(r["_MC"])
    return pool, nonpool


# กลุ่มที่ไม่ใช่เครื่องทอ — ไม่กินโควตา job (ชุดเดียวกับ _JOB_EXCLUDE_MC ใน Planning.py)
# อ่านค่าเดียวกันจากหน้า "ตั้งค่าระบบ" → แก้ที่เดียวแล้วตรงกันทั้งสองฝั่งเสมอ
JOB_EXCLUDE_MC = set(plan_config.get_list("KP_JOB_EXCLUDE_MC", upper=True))


def pinned_setup_jobs(pinned: pd.DataFrame) -> dict:
    """งานที่ตรึงไว้กินโควตา job setup ของสัปดาห์นั้นไปเท่าไร → {(WEEK, MC_GROUP): jobs}

    เครื่องของงานที่ตรึงถูกกันไว้ผ่าน SUMMARY_MC_REMAIN แล้ว (ดู pinned_mc_use) แต่โควตา
    job ต่อสัปดาห์ (PHET_DOUBLE 33 / PHET_SINGLE 44 / OM 13) เป็นคนละบัญชี ซึ่ง Planning
    pre-populate จากชีท DETAIL เท่านั้น — งานที่ตรึงไม่ได้อยู่ใน DETAIL (มันมาจากไฟล์แผน)
    ถ้าไม่ส่งยอดนี้ไปด้วย รอบ replan จะเห็น job slot ว่างเกินจริง แล้วแจก setup เกิน cap
    เท่ากับจำนวน setup ของงานที่ตรึงในสัปดาห์นั้น

    นับเฉพาะ NEW_MC (เครื่องที่ต้องตั้งใหม่จริง — carryover ไม่กินโควตา) และยุบด้วย max
    ต่อ (สัปดาห์, item, กลุ่มเครื่อง, เกจ) เพราะ Planning เขียน NEW_MC ซ้ำทุกแถวของเครื่อง
    ชุดเดียวกัน บวกตรง ๆ จะเกินจริง — สูตรเดียวกับ _preempt_job_slots ใน Planning.py
    """
    jobs: dict = {}
    for (_w, _it, _mcg, _g), v in pinned_setup_jobs_by_item(pinned).items():
        k = (_w, _mcg)
        jobs[k] = jobs.get(k, 0) + v
    return jobs


def pinned_setup_jobs_by_item(pinned: pd.DataFrame) -> dict:
    """เหมือน pinned_setup_jobs แต่ไม่ยุบ → {(WEEK, ITEM, MC_GROUP, GAUGE): เครื่องที่ตั้งใหม่}

    Planning ใช้หักโควตาตั้งเครื่องรายผ้า (2 ตัว/ผ้า/สัปดาห์) — งานที่ตรึงไว้ตั้งผ้า A ไปแล้ว
    ใบใหม่ของผ้า A สัปดาห์เดียวกันต้องเห็นว่าโควตาถูกใช้ไปแล้ว (user 30 ก.ย. 2569)
    ส่งผ่านชีท PINNED_JOBS_ITEM แยกจาก PINNED_JOBS เพื่อไม่เปลี่ยนรูปชีทเดิม
    """
    if pinned.empty:
        return {}
    d = pinned.copy()
    d["_WK"] = pd.to_numeric(_s(d, "PLAN_WEEK"), errors="coerce")
    d["_NEW"] = pd.to_numeric(_s(d, "NEW_MC"), errors="coerce").fillna(0)
    d["_MCG"] = _str_col(d, "MC_GROUP").str.upper()
    # งานจ้างทอที่ถูกตรึงไว้ไม่ได้ตั้งเครื่องในบ้าน → ไม่กินโควตา job
    # (MC_GROUP ของแถวจ้างทอเป็น "ชื่อผู้รับจ้าง" ซึ่งไม่มีใน MasterMC → ฝั่ง Planning
    #  _get_type_key_for_mc คืน default PHET_DOUBLE แล้วไปกินโควตาของโรงเพชรแทน)
    d["_OUT"] = (
        _str_col(d, "S9_ROUTING").str.lower().isin(["1", "1.0", "true", "yes"])
        | (_str_col(d, "S9_SUPP") != "")
        | (_str_col(d, "FACTORY_TYPE").str.upper() == "OUTSOURCE")
        | (d["_MCG"] == "COMKN")
        # ไฟล์แผนรุ่นเก่าอาจไม่มีคอลัมน์ S9_*/FACTORY_TYPE เลย → ดูจากชื่อกลุ่มเครื่องด้วย
        # (ผู้รับจ้างเป็นชื่อไทย ส่วนเครื่องในบ้านเป็นรหัส ASCII ล้วน)
        | d["_MCG"].str.contains(r"[^\x00-\x7F]", regex=True, na=False)
    )
    d = d[
        d["_WK"].notna() & (d["_NEW"] > 0) & (d["_MCG"] != "")
        & ~d["_MCG"].isin(JOB_EXCLUDE_MC) & ~d["_OUT"]
    ]
    if d.empty:
        return {}
    d["_ITEM"] = _str_col(d, "ITEM_CODE").str.upper()
    d["_G"] = _s(d, "MC_GUAGE").map(_norm_gauge)

    jobs: dict = {}
    per_key = d.groupby(["_WK", "_ITEM", "_MCG", "_G"])["_NEW"].max()
    for (_w, _it, _mcg, _g), v in per_key.items():
        if int(v) > 0:
            jobs[(int(_w), _it, _mcg, _g)] = int(v)
    return jobs


def _demand_keysets(demand: pd.DataFrame) -> list:
    """แต่ละแถวดีมานด์ → เซ็ตคีย์ '<เลข SC>|<item>' ที่แถวนั้นครอบ

    ต้องรวมทั้ง SC_NO และทุกใบใน SO_NO — ดีมานด์ฝั่ง booking ยุบหลายใบไว้แถวเดียวได้
    ('S721574,721436') ถ้าเทียบแค่ใบแรก งานที่ตรึงของใบที่สองจะหักยอดไม่ติด
    """
    item = demand["FABRIC_ITEM"].astype(str).str.strip().str.upper()
    return [
        {f"{n}|{it}" for n in ({_sc_norm(sc)} | {_sc_norm(t) for t in _so_tokens(so)}) if n}
        for sc, so, it in zip(_s(demand, "SC_NO"), _str_col(demand, "SO_NO"), item)
    ]


BOOKING_DEMAND_SOURCE = "booking (replan)"   # SOURCE_FILE ของแถวดีมานด์ที่ยกมาจาก booking


def deduct_pinned_from_demand(demand: pd.DataFrame, pinned: pd.DataFrame,
                              label: str = "งานที่ตรึงไว้",
                              sc_side_only: bool = False) -> pd.DataFrame:
    """หักยอดที่ตรึงไว้ออกจากดีมานด์รวม (แถวที่เหลือ 0 = ตัดทิ้ง)

    ใช้กับ "ทุกแถวที่ถูกแปะกลับเข้าแผนโดยไม่ผ่านลูปวางแผน" — งานที่ user ล็อก, core ที่เบรก
    และใบที่ออกเลข KP ; label มีไว้บอกใน log ว่ากำลังหักของชนิดไหน

    ⚠️ ด่านกันงาน (KP/core/ล็อก) ทำงานระดับ "แถว booking" แต่ดีมานด์คิดระดับ "ใบ" —
    กันแถวออกไม่ได้ลดยอดของใบเลย ถ้าไม่หักตรงนี้ Planning จะวางเต็มยอดใบแล้วแถวที่แปะกลับ
    ไปบวกทับอีกก้อน (เคสจริง 23 ก.ย. 2569: ใบ 717921 ยอด 1,800 กก. ถูกวางรวม 2,904 กก.)

    ต้องหักจาก "ดีมานด์รวม" หลัง dedup ไม่ใช่จาก order_ready อย่างเดียว — งานเดียวกัน
    มาได้ 2 ทาง คือใบ SC (order_ready) กับยอด booking ที่ยกมาวางใหม่ ; เวลามีทั้งคู่
    dedup จะตัดฝั่ง SC ทิ้งแล้วเหลือแถว booking ถ้าไปหักที่ orders อย่างเดียว ยอดของงาน
    ที่ตรึงไว้จะรอดมาถูกวางซ้ำอีกก้อนทับงานเดิม (เจอจากการทดสอบจริง)

    เลือกแถวที่จะหักโดยดูสัปดาห์ก่อน — แถวที่ครอบสัปดาห์ของงานที่ตรึง (BOOKING_WEEK_FROM..TO)
    ถูกหักก่อน เพื่อให้ยอดที่หายไปตรงกับก้อนงานจริง ไม่ไปลดก้อนของสัปดาห์อื่น

    sc_side_only=True — ใช้กับงานที่ถูกกันออกตั้งแต่ระดับแถว booking แล้ว (ใบที่ออกเลข KP /
    core ที่เบรกไว้): ยอดของแถวพวกนี้ไม่เคยเข้าดีมานด์ฝั่ง booking อยู่แล้ว (build_booking_demand
    รับแต่แถวที่ปลดล็อก) → หักได้เฉพาะแถวฝั่ง SC ซึ่งเป็นยอดค้างทั้งใบที่รวมส่วนนั้นไว้
    ถ้าไปหักแถว booking ด้วย = หักซ้ำ ยอดหายจากแผน (เคสจริง 29 ก.ย. 2569: 31 ใบ 46,348 กก.
    เช่น 720274 FD3BASMI64A0 booking 12,530 + KP 850 ถูกวางแค่ 11,680 + 850)
    """
    if demand.empty or pinned.empty:
        return demand

    pins = []
    qty = pd.to_numeric(_s(pinned, "PRODUCE_QTY"), errors="coerce").fillna(0)
    wks = pd.to_numeric(_s(pinned, "PLAN_WEEK"), errors="coerce")
    items = _str_col(pinned, "ITEM_CODE").str.upper()
    sos = _str_col(pinned, "SC_SO_NO")
    for so, it, q, w in zip(sos, items, qty, wks):
        if q <= 0:
            continue
        keys = {f"{n}|{it}" for n in {_sc_norm(t) for t in _so_tokens(so)} if n}
        if not keys:
            print(f"⚠️ แถวแผน{label} ({it}) ไม่มีเลข SC → หักยอดออกจากดีมานด์ไม่ได้")
            continue
        pins.append({"keys": keys, "qty": float(q), "week": (int(w) if pd.notna(w) else None),
                     "item": it})
    if not pins:
        return demand

    out = demand.copy().reset_index(drop=True)
    keysets = _demand_keysets(out)
    pend = pd.to_numeric(out["PENDING_PLAN"], errors="coerce").fillna(0)
    w_from = pd.to_numeric(_s(out, "BOOKING_WEEK_FROM"), errors="coerce")
    w_to = pd.to_numeric(_s(out, "BOOKING_WEEK_TO"), errors="coerce")
    from_booking = (_str_col(out, "SOURCE_FILE") == BOOKING_DEMAND_SOURCE) if sc_side_only \
        else pd.Series(False, index=out.index)

    cut, short, touched, n_booking_only = 0.0, [], set(), 0
    for pin in sorted(pins, key=lambda p: (p["week"] is None, p["week"] or 0)):
        cands = [i for i in out.index if keysets[i] & pin["keys"] and pend.at[i] > 0]
        if sc_side_only:
            _sc = [i for i in cands if not from_booking.at[i]]
            if cands and not _sc:
                n_booking_only += 1   # ใบนี้เหลือแต่ดีมานด์ booking = ไม่มียอดของงานนี้อยู่แล้ว
                continue
            cands = _sc
        if not cands:
            short.append(pin)
            continue

        def _same_week(i):
            """0 = แถวนี้ครอบสัปดาห์ของงานที่ตรึง (หักก่อน), 1 = แถวอื่น"""
            if pin["week"] is None or pd.isna(w_from.at[i]) or pd.isna(w_to.at[i]):
                return 1
            return 0 if w_from.at[i] <= pin["week"] <= w_to.at[i] else 1

        left = pin["qty"]
        for i in sorted(cands, key=_same_week):
            if left <= 0:
                break
            take = min(left, float(pend.at[i]))
            pend.at[i] -= take
            left -= take
            cut += take
            touched.add(i)
        if left > 0.001:
            short.append({**pin, "qty": left})

    out["PENDING_PLAN"] = pend
    # ตัดเฉพาะแถวที่ "เราหักจนหมด" เท่านั้น — ดีมานด์ปกติมีแถวยอด 0/ติดลบอยู่แล้วบางส่วน
    # (ใบที่วางครบแล้ว) ซึ่ง Replan รอบปกติปล่อยผ่าน ถ้าเหมาตัดด้วยจะเปลี่ยนผลนอกเรื่องล็อก
    drop = [i for i in touched if pend.at[i] <= 0.001]
    if drop:
        out = out.drop(index=drop)
    print(f"➖ หักยอด{label}ออกจากดีมานด์ {cut:,.0f} กก. "
          f"(ตัดแถวที่เหลือ 0 ออก {len(drop)} แถว → เหลือ {len(out)} แถว)")
    if n_booking_only:
        print(f"   ↳ {label} {n_booking_only} ก้อนไม่หัก: ใบนั้นเหลือแต่ดีมานด์ booking "
              f"ซึ่งไม่รวมยอดของงานนี้อยู่แล้ว (กันหักซ้ำ)")
    if short:
        print(f"⚠️ หักยอดไม่ครบ (ดีมานด์รอบนี้ไม่มีก้อนของ{label}เหลือให้หัก): "
              + ", ".join(f"{p['item']} {yw_label(p['week'])} {p['qty']:,.0f} กก." for p in short[:10])
              + (" ..." if len(short) > 10 else ""))
    return out.reset_index(drop=True)


def _keep_rows_yw(df: pd.DataFrame) -> pd.DataFrame:
    """สำเนาของแถว "คงแผนเดิม" ที่ PLAN_WEEK เป็น YW — ใช้ส่งเข้า deduct_pinned_from_demand
    (core_plan_rows คืนรูปไฟล์แผน: สัปดาห์เปล่า + PLAN_YEAR ; ดีมานด์คิดเป็น YW)"""
    if df is None or df.empty:
        return df if df is not None else pd.DataFrame()
    return _to_yw(df.copy(), "PLAN_WEEK", "PLAN_YEAR")


def _plan_row_keys(df: pd.DataFrame) -> list:
    """แถวแผน → ลิสต์ของเซ็ตคีย์ (เลขใบ, รหัสผ้า, สัปดาห์ YW) ทีละแถว
    ช่องใบที่ยุบหลายเลขไว้ด้วยกัน (' | ' / ',') แตกเป็นหลายคีย์ — ตรงกันสักเลขก็ถือว่าแถวเดียวกัน"""
    out = []
    wk = pd.to_numeric(_s(df, "PLAN_WEEK"), errors="coerce")
    items = _str_col(df, "ITEM_CODE").str.upper()
    sos = _s(df, "SC_SO_NO")
    for w, it, so in zip(wk, items, sos):
        ks = set()
        if pd.notna(w) and it:
            for t in _so_tokens(str(so or "").replace("|", ",")):
                n = _sc_norm(t)
                if n:
                    ks.add((n, it, int(w)))
        out.append(ks)
    return out


def drop_pinned_rebuilt(pinned: pd.DataFrame, rebuilt: list) -> pd.DataFrame:
    """ทิ้งแถวล็อกที่ "รอบนี้ถูกสร้างใหม่จาก booking อยู่แล้ว" (ใบที่ออกเลข KP / core ที่เบรก)

    ถ้าไม่ทิ้ง งานชุดเดียวเข้าแผน 2 ทางพร้อมกัน = ยอดเบิ้ลเป๊ะ ๆ สองเท่า
    (ทดสอบจริง 23 ก.ย. 2569: ล็อกแถว KP 3 ใบ → วาง 10,990/7,480/6,145 จากยอดใบ
     5,495/3,740/3,072 ; ด่านหักดีมานด์ช่วยไม่ได้เพราะใบถูกกันออกจากดีมานด์ไปแล้ว)

    ใช้ของที่สร้างจาก booking เป็นตัวจริง — สดกว่าแถวในไฟล์แผนรอบก่อนเสมอ
    งานไม่หลุดจากการล็อก เพราะงาน KP/core ที่เบรก ถูกตรึงด้วยตัวมันเองอยู่แล้ว
    ส่วนแถวล็อกที่ "ไม่ถูกสร้างใหม่" (เช่น user เลิกเบรก core แล้ว) ยังทำงานตามเดิม
    """
    if pinned is None or pinned.empty:
        return pinned
    built = set()
    for df in rebuilt:
        df = _keep_rows_yw(df)
        if df is None or df.empty:
            continue
        for ks in _plan_row_keys(df):
            built |= ks
    if not built:
        return pinned

    keys = _plan_row_keys(pinned)
    keep = pd.Series([not (ks & built) for ks in keys], index=pinned.index)
    n_drop = int((~keep).sum())
    if n_drop:
        _d = pinned[~keep]
        print(f"♻️ ทิ้งแถวล็อก {n_drop} แถว (รวม "
              f"{pd.to_numeric(_s(_d, 'PRODUCE_QTY'), errors='coerce').fillna(0).sum():,.0f} กก.) "
              f"— งานชุดนี้ถูกสร้างใหม่จาก booking รอบนี้แล้ว (ใบที่ออกเลข KP / core ที่เบรก) "
              f"→ กันยอดเบิ้ล")
    return pinned[keep]


PINNED_REMARK = "🔐 ล็อกไว้ (Lock_Plan) — คงแผนเดิม ไม่ถูกจัดใหม่"


def pinned_result_rows(pinned: pd.DataFrame) -> pd.DataFrame:
    """แถวงานที่ตรึงไว้ → แถวที่พร้อมแปะกลับเข้าแผน (Planning รอบ replan เป็นคนแปะ)

    ไม่แปะกลับ = งานหายไปจากหน้าแผนทั้งที่ยังต้องผลิต (ถูกกันออกจากดีมานด์ไปแล้ว)
    """
    if pinned.empty:
        return pd.DataFrame()
    add = pinned.copy()
    # 🔐 = เครื่องหมายเดียวกับที่ Gantt ใช้กับงานที่ user ล็อก (🔒 สงวนไว้ให้สัปดาห์ freeze)
    # แถวที่ตรึงข้ามหลายรอบมีป้ายนี้ติดมาจากไฟล์แผนรอบก่อนแล้ว → ไม่ต่อซ้ำให้ยาวสะสม
    _rm = _str_col(add, "REMARK")
    _has = _rm.str.startswith(PINNED_REMARK)
    add["REMARK"] = _rm.where(_has, (PINNED_REMARK + " | " + _rm).str.rstrip(" |"))
    # เครื่องของงานที่ตรึงถูกกันไว้ใน SUMMARY_MC_REMAIN ไปแล้ว (ดู pinned_mc_use) → แถวที่แปะกลับ
    # ต้องประกาศตัวเองว่า "นั่งบนเครื่องที่กันไว้" ด้วย MC_BOOKING = จำนวนเครื่องของบล็อก ไม่งั้น
    # หน้าเว็บคิดเครื่องที่แผนกินเพิ่ม = ACTUAL_MC − MC_BOOKING (PlanGantt.planMcAgg / mcRecalc)
    # แล้วหักเครื่องชุดเดิมซ้ำอีกรอบ — เคสจริง 9 ก.ย. 2569: SBP เกจ 26 (SINGLE-36) W41 ว่าง 12 ตัว
    # แผนใช้จริง 12 แต่โชว์ "ไม่ว่าง (เกิน 3)" เพราะงานล็อก FD3BASKD21/03A0 3 ตัวถูกนับสองรอบ
    # เครื่องของบล็อก (สัปดาห์, item, กลุ่มเครื่อง, เกจ) = max ภายในใบสั่ง แล้วบวกข้ามใบสั่ง
    # = สูตรเดียวกับที่ pinned_mc_use กันไว้ (ใช้คนละสูตร หน้าเว็บจะเห็นงานล็อกกินเครื่องเพิ่ม/ขาด)
    # กติกาเดียวกับแถว core ที่ core_plan_rows ตั้ง MC_BOOKING = ACTUAL_MC ไว้แล้ว
    _k4 = ["_WK", "_ITEM", "_MCG", "_G"]
    _t = pd.DataFrame({
        "_WK": pd.to_numeric(_s(add, "PLAN_WEEK"), errors="coerce"),
        "_ITEM": _str_col(add, "ITEM_CODE").str.upper(),
        "_MCG": _str_col(add, "MC_GROUP").str.upper(),
        "_G": _s(add, "MC_GUAGE").map(_norm_gauge),
        "_SC": _str_col(add, "SC_SO_NO"),
        "_MC": pd.to_numeric(_s(add, "ACTUAL_MC"), errors="coerce").fillna(0),
    }, index=add.index)
    _blk_sum = (
        _t.groupby(_k4 + ["_SC"], dropna=False)["_MC"].max()
        .groupby(level=_k4, dropna=False).sum().rename("_BLK").reset_index()
    )
    _blk = _t[_k4].merge(_blk_sum, on=_k4, how="left")["_BLK"].fillna(0)
    _blk.index = add.index
    _bk = pd.to_numeric(_s(add, "MC_BOOKING"), errors="coerce").fillna(0)
    add["MC_BOOKING"] = _bk.where(_bk >= _blk, _blk)
    _fix = int((add["MC_BOOKING"] > _bk).sum())
    if _fix:
        print(f"🔧 งานที่ตรึงไว้: ตั้ง MC_BOOKING = เครื่องของบล็อก {_fix} แถว "
              f"(เครื่องถูกกันไว้ใน capacity แล้ว — กันหน้าเว็บหักซ้ำ)")
    return add


def write_result_append(pinned_rows: pd.DataFrame, core_rows: pd.DataFrame,
                        kp_rows: pd.DataFrame = None) -> None:
    """เขียนแถวที่ต้องแปะเข้าแผนตรง ๆ ลง replan_append.xlsx ให้ Planning รอบ replan อ่านไปต่อท้าย

    ทำไมไม่แปะหลัง Planning จบเหมือนเดิม: แถว core ยังขาด TARGET_KNIT / TEAM_PROGRAM ซึ่งต้องใช้
    ปฏิทิน + offset ตาม FOB_TYPE + ชีท Program ที่อยู่ใน Planning — เดิม Planning รอบหลักเติมให้
    ตอนแปะชีท REPLAN แต่รอบหลักเลิกรันแล้ว จึงต้องให้รอบ replan แปะ + เติมเองก่อนเขียนไฟล์
    เขียนทุกรอบแม้ว่าง — กันไฟล์ของรอบก่อนค้างแล้วถูกแปะซ้ำ
    """
    # งานที่ตรึงมี PLAN_WEEK/TARGET_KNIT เป็น YW (แปลงตอนอ่านไฟล์แผน) → แยกกลับเป็นรูปไฟล์แผน
    pinned_rows = _split_yw(pinned_rows, [("PLAN_WEEK", "PLAN_YEAR"), ("TARGET_KNIT", "TARGET_YEAR")])
    kp_rows = kp_rows if kp_rows is not None else pd.DataFrame()
    with pd.ExcelWriter(APPEND_FILE, engine="openpyxl") as w:
        (pinned_rows if not pinned_rows.empty else pd.DataFrame(columns=["ITEM_CODE"])).to_excel(
            w, sheet_name=APPEND_PINNED_SHEET, index=False)
        (core_rows if not core_rows.empty else pd.DataFrame(columns=["ITEM_CODE"])).to_excel(
            w, sheet_name=APPEND_CORE_SHEET, index=False)
        (kp_rows if not kp_rows.empty else pd.DataFrame(columns=["ITEM_CODE"])).to_excel(
            w, sheet_name=APPEND_KP_SHEET, index=False)
    print(f"💾 แถวที่แปะเข้าแผนโดยไม่จัดใหม่: ตรึงไว้ {len(pinned_rows)} แถว + core ที่เบรก {len(core_rows)} แถว "
          f"+ ใบที่ออกเลข KP {len(kp_rows)} แถว → {APPEND_FILE.name}")


# คอลัมน์ท้ายชีท PLAN บอกว่าแถวมาจากไหน — "SC ใหม่" = ใบจาก View_SC (order_ready) ที่ไม่ได้
# มาจาก booking ; "Booking" = งานเดิมใน booking ที่ถูกยกมาวางใหม่ (รวม core ที่คงแผนเดิม)
SC_STATUS_COL = "SC_STATUS"
SC_STATUS_NEW = "SC ใหม่"
SC_STATUS_BOOKING = "Booking"
# หมายเหตุ: แถว "กันเครื่อง" (HOLD_COLOR / HOLD_TARGET) เคยถูกยกเว้นไม่คิด SC_STATUS เพราะ
# ไม่มีข้อมูลใบเป็นของตัวเอง — ตั้งแต่ 21/9/2569 แถวพวกนี้ยกใบสั่งที่ทำให้ต้องกันเครื่องมาครบแล้ว
# (Planning._hold_order_fields) จึงคิดสถานะได้เหมือนแถวปกติ


def _sc_item_keys(so_values, items) -> set:
    """(SO_NO, ITEM) หลายแถว → {'720188|ITEM', ...}
    แตกใบที่ยุบไว้ช่องเดียว (คั่นคอมมา หรือ ' | ') ก่อน แล้ว normalize เลข SC ให้เทียบกันได้"""
    out = set()
    for so, it in zip(so_values, items):
        it = str(it or "").strip().upper()
        if not it:
            continue
        for t in _so_tokens(str(so or "").replace("|", ",")):
            n = _sc_norm(t)
            if n:
                out.add(f"{n}|{it}")
    return out


def _sc_status(so, item, source, remark, demand_keys: set, booking_keys: set) -> str:
    """ค่า SC_STATUS ของแถวแผนหนึ่งแถว

    แถวยุบหลายใบ มีใบใดมาจาก booking = Booking ทั้งแถว (กติกาเดียวกับหน้าเว็บ KnitPlan.jsx)
    แถวที่ตรึง 🔐 ถูกหักออกจากดีมานด์ไปแล้ว → เทียบกับ booking ทั้งไฟล์แทน
    """
    src = str(source or "").strip().upper()
    if src in (CORE_PLAN_SOURCE, KP_PLAN_SOURCE):   # ยกมาจาก booking ตรง ๆ ทั้งคู่
        return SC_STATUS_BOOKING
    keys = _sc_item_keys([so], [item])
    ref = booking_keys if str(remark or "").startswith(PINNED_REMARK) else demand_keys
    return SC_STATUS_BOOKING if keys & ref else SC_STATUS_NEW


def _add_sc_status(ws, demand_keys: set, booking_keys: set) -> int:
    """เติมคอลัมน์ SC_STATUS ต่อท้ายชีท PLAN (openpyxl — สีแถวเดิมอยู่ครบ) → คืนจำนวน SC ใหม่"""
    from copy import copy
    header = [c.value for c in ws[1]]
    filled = [i for i, v in enumerate(header, 1) if v not in (None, "")]
    if not filled:
        return 0
    pos = {str(v).strip().upper(): i for i, v in enumerate(header, 1) if v not in (None, "")}
    col = pos.get(SC_STATUS_COL, max(filled) + 1)

    def _v(r, name):
        i = pos.get(name)
        return ws.cell(r, i).value if i else None

    def _put(r, value):
        c = ws.cell(r, col, value)
        if col > 1:                       # ยืมสไตล์ช่องข้าง ๆ (หัวตัวหนา / สีแถว) ให้กลืนกับตาราง
            c._style = copy(ws.cell(r, col - 1)._style)

    _put(1, SC_STATUS_COL)
    n_new = 0
    for r in range(2, ws.max_row + 1):
        if _v(r, "ITEM_CODE") in (None, ""):
            continue
        st = _sc_status(_v(r, "SC_SO_NO"), _v(r, "ITEM_CODE"), _v(r, "PLAN_SOURCE"),
                        _v(r, "REMARK"), demand_keys, booking_keys)
        _put(r, st)
        n_new += st == SC_STATUS_NEW
    return n_new


def _move_col_last(ws, name: str) -> None:
    """ย้ายคอลัมน์ชื่อ name ไปท้ายสุดของชีท (openpyxl — ค่า + สไตล์ย้ายตามไปด้วย)

    KP_NO ต้องอยู่หลังสุด (user 22 ก.ย. 2569): ผู้ใช้ภายนอกอ่านชีท PLAN ตามตำแหน่งคอลัมน์
    ถ้า SC_STATUS ต่อท้ายหลัง KP_NO ก็ยังถือว่าไม่ใช่คอลัมน์สุดท้าย → ย้ายหลังเติม SC_STATUS
    """
    from copy import copy
    header = [c.value for c in ws[1]]
    filled = [i for i, v in enumerate(header, 1) if v not in (None, "")]
    src = next((i for i, v in enumerate(header, 1) if str(v or "").strip().upper() == name), None)
    if src is None or src == max(filled):
        return
    dst = max(filled) + 1
    for r in range(1, ws.max_row + 1):
        s = ws.cell(r, src)
        d = ws.cell(r, dst, s.value)
        d._style = copy(s._style)
    ws.delete_cols(src)


def publish_plan_file(demand_keys: set | None = None,
                      booking_keys: set | None = None) -> Path | None:
    """replan_result.xlsx → production_plan_<วัน-เดือน-ปีพ.ศ._เวลา>.xlsx (ไฟล์ที่หน้าเว็บเปิด)

    ชื่อไฟล์รูปแบบเดียวกับที่ Planning รอบหลักเคยเขียน — หน้าเว็บ/scheduler/ตัวลบไฟล์เก่า glob ชื่อนี้
    คัดลอกผ่าน openpyxl (ไม่ใช่ pandas) เพื่อเก็บสีแถวที่ Planning ใส่ไว้ในชีท PLAN
    demand_keys / booking_keys = คีย์ 'SC|ITEM' ของยอด booking ที่ยกมาวาง / ของ booking ทั้งไฟล์
    → ใช้เติมคอลัมน์ SC_STATUS ท้ายชีท PLAN
    """
    if not RESULT_FILE.exists():
        print(f"❌ ไม่พบ {RESULT_FILE.name} → ไม่มีไฟล์แผนผลิตรอบนี้")
        return None
    from openpyxl import load_workbook
    now = datetime.now()
    PLAN_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = PLAN_OUTPUT_DIR / (f"production_plan_{now.day}-{now.month}-{now.year + 543}"
                             f"_{now.hour:02d}{now.minute:02d}.xlsx")
    try:
        wb = load_workbook(RESULT_FILE)
        if COMPARE_SHEET in wb.sheetnames:
            wb[COMPARE_SHEET].title = PUBLISHED_COMPARE_SHEET
        if "PLAN" in wb.sheetnames:
            n_new = _add_sc_status(wb["PLAN"], demand_keys or set(), booking_keys or set())
            print(f"🆕 คอลัมน์ {SC_STATUS_COL}: {SC_STATUS_NEW} {n_new} แถว (มาจาก View_SC ไม่ใช่ booking)")
            _move_col_last(wb["PLAN"], "KP_NO")
        # เขียนลงไฟล์ชั่วคราวก่อนแล้วค่อยสลับชื่อ — ห้ามให้ไฟล์ชื่อ production_plan_*.xlsx
        # โผล่ขึ้นมาตั้งแต่ยังเขียนไม่เสร็จ เพราะหน้าเว็บถือว่า "ไฟล์แผนล่าสุด" คือไฟล์นี้ทันที
        # ผู้ใช้ที่กำลังลาก Gantt อยู่จะเซฟทับไฟล์ที่ยังเขียนไม่จบ = แผนเสียทั้งไฟล์
        with file_guard.atomic_save(out) as tmp:
            wb.save(tmp)
        wb.close()
    except Exception as e:
        print(f"❌ เขียนไฟล์แผน {out.name} ไม่สำเร็จ: {e}")
        return None
    print(f"📄 ไฟล์แผนผลิต → {out}")
    return out


# (TIMELINE ล็อก vs ปลดล็อก ใช้ _timeline() ด้านบน — ปฏิทินเดียวกับ Planning, คีย์เป็น YW)

# =========================
# แปลง booking → ดีมานด์
# =========================
def _s(df: pd.DataFrame, col: str) -> pd.Series:
    """คอลัมน์ที่อาจไม่มีในไฟล์ (booking เก่าไม่มี FR_CODE/FG_WEEK) → คืน series ว่างแทน KeyError"""
    if col in df.columns:
        return df[col]
    return pd.Series([None] * len(df), index=df.index, dtype=object)


def _str_col(df: pd.DataFrame, col: str) -> pd.Series:
    """คอลัมน์ข้อความที่ทน NaN — ต้อง fillna ก่อน astype(str)
    ไม่งั้น pandas ใหม่คืน dtype 'str' ที่ยังมี NaN ค้าง แล้ว .str accessor พัง"""
    s = _s(df, col).fillna("").astype(str).str.strip()
    return s.replace({"nan": "", "None": "", "NaT": "", "<NA>": ""}).fillna("")


def _line_tokens(v) -> list:
    """ค่าเลขบรรทัด (เลขเดี่ยว หรือยุบหลายบรรทัดด้วยคอมมาจาก AVA_MC._uniq_lines) → list เลขที่ใช้ได้จริง

    view ส่ง 0 มาแทน null สำหรับงาน CL-ORDERS/YD-ORDERS (วัดจริง 21/9/2569: CL 1,316/1,316 แถว
    และ YD 50/50 แถวของ view_booking เป็น 0 ทั้งหมด) — user 21/9/2569: "ถ้าข้อมูลมาแบบนั้น
    ก็ใส่ข้อมูลไปแบบนั้น" → ไม่ทิ้ง 0 อีกต่อไป ใบที่ต้นทางมีแค่ 0 ก็ลง 0 ในชีทแผนตามจริง

    ยกเว้นแถวที่ยุบหลายใบแล้วมีทั้งเลขจริงและ 0 ("774607,0") → เก็บเฉพาะเลขจริง
    เพราะ 0 ของอีกใบไม่ได้เพิ่มข้อมูล มีแต่ทำให้อ่านเหมือนมีบรรทัดเลข 0 อยู่จริง
    """
    real, zeros = [], []
    for t in ColorOpen.so_text(v).split(","):
        t = t.strip()
        if not t:
            continue
        bucket = real if t.strip("0") else zeros
        if t not in bucket:
            bucket.append(t)
    return real or zeros[:1]


def _first_text(x) -> str:
    """ค่าแรกที่ "ไม่ว่าง" ของก้อน — ใช้แทน pandas "first" กับคอลัมน์ข้อความทุกตัว

    ⚠ "first" ของ pandas ข้ามแค่ NaN แต่คอลัมน์พวกนี้ผ่าน _str_col มาแล้วจึงเป็นสตริงว่าง ""
    ไม่ใช่ NaN → แถวแรกของก้อนที่บังเอิญว่างจะบังค่าจริงของแถวอื่นทั้งก้อน
    (วัดจริง 21/9/2569: YARN_USED ว่างในแผน 242 แถว ทั้งที่ต้นทางมีค่าให้ 200 แถว)
    """
    for v in x:
        s = "" if v is None else str(v).strip()
        if s and s.lower() not in ("nan", "none", "nat", "<na>"):
            return s
    return ""


def _uniq_kp_text(x) -> str:
    """เลขใบจอง (KP_NO) ของก้อน → รวมไม่ซ้ำคั่น ' | ' (ตัวคั่นเดียวกับ AVA_MC._uniq_values)"""
    out = []
    for v in x:
        for t in str(v if v is not None else "").split("|"):
            t = t.strip()
            if t and t.lower() not in ("nan", "none") and t not in out:
                out.append(t)
    return " | ".join(out)


def _uniq_line_text(x) -> str:
    """หลายค่าเลขบรรทัด → รวมไม่ซ้ำคั่นคอมมา (ตัวรวมตอน groupby ของ build_booking_demand)

    กติกา 0 เหมือน _line_tokens แต่ต้องตัดสินที่ระดับ "ทั้งก้อน": หลายใบมารวมเป็นแถวเดียว
    ถ้ามีใบที่มีเลขจริงปนอยู่ 0 ของใบอื่นต้องไม่ติดมาด้วย
    """
    _all = list(dict.fromkeys(t for v in x for t in _line_tokens(v)))
    _real = [t for t in _all if t.strip("0")]
    return ",".join(_real or _all[:1])


def _line_text(v) -> str:
    """เลขบรรทัดใบจาก booking → ข้อความตามค่าต้นทาง (รวม '0' ของงาน CL/YD — ดู _line_tokens)"""
    return ",".join(_line_tokens(v))


def _last_po_in_date(v) -> object:
    """PO_IN_DATE เป็นวันที่หลายล็อตคั่นคอมมา → ใช้วันหลังสุด = ด้ายเข้าครบ"""
    s = str(v or "").strip()
    if not s or s.lower() in ("nan", "nat", "none"):
        return pd.NaT
    parts = [p.strip() for p in s.split(",") if p.strip()]
    dates = [d for d in (pd.to_datetime(p, dayfirst=True, errors="coerce") for p in parts) if pd.notna(d)]
    return max(dates) if dates else pd.NaT


def _fr_parse(v) -> tuple:
    """FR_CODE → (SC ใบแรก, ORDER_TYPE)

    FR_CODE 1 ใบ = 'S{SC_NO}-{ORDER_TYPE}' แต่ AVA_MC ยุบหลาย SO ที่ลงเครื่องเดียวกันมาไว้
    ช่องเดียว คั่นด้วย ' | ' (ห้ามใช้ comma — ชื่อลูกค้ามีจุลภาคในตัว ดู AVA_MC._uniq_values)
        'S721282-SC-ORDERS | S721257-SC-ORDERS'
    ⚠️ ห้าม split('-') ตรง ๆ กับทั้งช่อง: จะได้ ORDER_TYPE = 'SC-ORDERS | S721257-SC-ORDERS'
    ซึ่ง Planning ไม่รู้จัก แล้วโยนงานทั้งก้อนเข้า UNPLANNED เงียบ ๆ
    (รอบ 17/8/2026 บนเซิร์ฟเวอร์: หายไป 100 แถว ≈ 330 ตัน ด้วยเหตุนี้)
    → ต้องแตก '|' ก่อน แล้วค่อยแยก '-' ทีละใบ (กติกาเดียวกับ plan_view._order_type)

    ชนิดปนกันได้ (เช่น 'S721247-ผ้าเติม | S721142-SC-ORDERS') → เลือกตัวแรกที่ Planning
    รองรับ ; ไม่มีสักตัว → คืนตัวแรกที่อ่านได้ ให้ไปโดนเตือน/UNPLANNED ตามจริง
    """
    scs, types = [], []
    for part in str(v or "").split("|"):
        p = part.strip()
        if not p or p.lower() in ("nan", "none"):
            continue
        # ใบที่ไม่มีเลข SC นำหน้า ('SC-ORDERS', '-SC-ORDERS' — เจอจริงในไฟล์ booking เมื่อ
        # view ส่ง SC_NO ว่างมา) ต้องอ่านทั้งใบเป็น ORDER_TYPE ; ถ้า partition('-') ตรง ๆ
        # จะได้ชนิด 'ORDERS' ซึ่ง Planning ไม่รู้จักเหมือนเดิม
        m = re.match(r"^(S?\d+)\s*-\s*(.*)$", p)
        if m:
            sc, t = m.group(1).strip(), m.group(2).strip()
        else:
            sc, t = "", p.lstrip("-").strip()
        if sc and sc not in scs:
            scs.append(sc)
        if t and t not in types:
            types.append(t)
    otype = next((t for t in types if t in PLANNABLE_ORDER_TYPES), types[0] if types else "")
    return (scs[0] if scs else ""), otype


def _sc_norm(v) -> str:
    """เลข SC ให้อยู่รูปเดียวกันก่อนเทียบ: 'S720140' / '720140.0' / ' 720140 ' → '720140'
    (order_ready เก็บ 'S720140' ส่วน SO_NO ในบางแถวถูกอ่านเป็น float — ดู _so_tokens)"""
    t = str(v or "").strip().upper()
    if t.startswith("S"):
        t = t[1:]
    if re.fullmatch(r"\d+\.0+", t):
        t = t.split(".")[0]
    return t


# คอลัมน์ที่เป็นของ "ทั้งบล็อก" ไม่ใช่ของใบ (เครื่อง/setup/วันทำงาน) — DETAIL_SC ไม่มีเพราะ
# AVA_MC คำนวณหลังยุบใบแล้ว → ยกมาจาก DETAIL ด้วยคีย์บล็อก ค่าเดียวซ้ำทุกใบในบล็อกนั้น
# (กติกาเดียวกับที่ Planning เขียนแถวแผน: หลาย SC แบ่งเวลาบนเครื่องชุดเดิม → นับแบบ max ไม่บวก)
_BLOCK_COLS = [
    "MC_USE", "MC_USE_CEIL", "_SETUP_DAYS", "_PREV_WEEK", "_WEEK_GAP", "_IS_NEW_SETUP",
    "_PREV_MC_USE_CEIL", "_MC_INCREASE", "_HAS_MC_INCREASE", "_EFFECTIVE_WORKING_DAYS",
    "WORKING_DAY", "FACTORY", "POOL_TYPE", "TOTAL_MC", "MC_USE_CUM", "TOTAL_MC_REMAIN",
    "_CARRY_KEY", "_LOCK_WARM_MC", "_LOCK_SAME_WEEK_MC",
]
_BLOCK_KEY = ["MC_GROUP", "GUAGE", "ITEM_CODE", "WEEK"]


def load_detail_sc(booking_file: Path, detail: pd.DataFrame):
    """ชีท DETAIL_SC (1 แถว = 1 SC) พร้อมคอลัมน์ระดับบล็อกจาก DETAIL

    ทำไมต้องมี: DETAIL ยุบทุก SO ของ (item, สัปดาห์, เครื่อง, เกจ) เป็นแถวเดียว — จำเป็นต่อ
    การนับเครื่อง แต่ FG_WEEK/RDD_WEEK เหลือได้ค่าเดียวคือใบที่เร็วสุด ถ้าวางแผนจากตารางนี้
    ยอดทั้งก้อนจะถูกวางตาม deadline ของใบอื่น (ดู DETAIL_SC ใน AVA_MC.py)

    คืน None ถ้าไฟล์ booking สร้างก่อนมีชีทนี้ → ผู้เรียกกลับไปใช้ DETAIL เหมือนเดิม
    """
    try:
        d = pd.read_excel(booking_file, sheet_name="DETAIL_SC")
    except (ValueError, KeyError):
        return None
    d.columns = d.columns.str.strip().str.upper()
    if d.empty or "SO_NO" not in d.columns:
        return None
    _to_yw(d)  # WEEK → YW ให้ตรงกับ detail ที่ main แปลงไว้แล้ว (merge ด้วยคีย์บล็อกที่มีสัปดาห์)
    add = [c for c in _BLOCK_COLS if c in detail.columns and c not in d.columns]
    if add and all(c in detail.columns for c in _BLOCK_KEY) and all(c in d.columns for c in _BLOCK_KEY):
        blk = detail[_BLOCK_KEY + add].copy()
        for _df in (blk, d):
            _df["WEEK"] = pd.to_numeric(_df["WEEK"], errors="coerce")
            _df["GUAGE"] = _df["GUAGE"].map(_norm_gauge)
            for _c in ("MC_GROUP", "ITEM_CODE"):
                _df[_c] = _str_col(_df, _c).str.upper()
        d = d.merge(blk.drop_duplicates(_BLOCK_KEY), on=_BLOCK_KEY, how="left")
    return d


def _block_keys(df: pd.DataFrame) -> list:
    """คีย์บล็อก (เครื่อง, เกจ, item, สัปดาห์) แบบ normalize แล้ว — ใช้เทียบข้ามชีท"""
    wk = pd.to_numeric(_s(df, "WEEK"), errors="coerce").map(
        lambda v: int(v) if pd.notna(v) else None)
    return list(zip(
        _str_col(df, "MC_GROUP").str.upper(),
        _s(df, "GUAGE").map(_norm_gauge),
        _str_col(df, "ITEM_CODE").str.upper(),
        wk,
    ))


def locked_detail_sc(detail_sc, locked_detail: pd.DataFrame):
    """DETAIL_SC เฉพาะใบของบล็อกที่ยังล็อกอยู่ — ชีทคู่ของ DETAIL ในไฟล์ capacity รอบ replan

    ทำไมต้องมี: Planning รอบ replan อ่าน DETAIL_SC จาก booking_final_replan.xlsx
    (Planning.py:373) เพื่อแตกแถว OLD ทีละใบสั่ง ถ้าไม่เขียนชีทนี้ไปด้วย แถว OLD จะยุบใบ
    → SC_SO_NO กลายเป็น '708936,719941' กำหนดส่ง/ลูกค้า/PO ของแต่ละใบหาย และหน้า
    CoreStock กลับไปนับต่อบล็อก

    คัดด้วย "คีย์บล็อก" ไม่ใช่รันมาสก์ล็อกซ้ำ — สองชีทในไฟล์เดียวกันต้องเห็นงานชุดเดียวกันเป๊ะ
    (Planning merge ด้วยคีย์นี้ บล็อกที่ไม่มีใบจะหายทั้งก้อน ส่วนใบที่ไม่มีบล็อกจะได้คอลัมน์
    เครื่อง/setup เป็นค่าว่าง)

    คืนค่าคีย์บล็อกตามรูปที่อยู่ใน DETAIL เสมอ เพราะ load_detail_sc normalize GUAGE ไปแล้ว
    ('22' ขณะที่ DETAIL อาจเป็น '22.0') ถ้าปล่อยไว้ merge ฝั่ง Planning จะไม่ติด
    """
    if detail_sc is None or detail_sc.empty or locked_detail.empty:
        return None
    if not all(c in detail_sc.columns for c in _BLOCK_KEY) or \
       not all(c in locked_detail.columns for c in _BLOCK_KEY):
        return None

    canon: dict = {}
    for k, row in zip(_block_keys(locked_detail),
                      locked_detail[_BLOCK_KEY].to_dict("records")):
        canon.setdefault(k, row)

    sc_keys = _block_keys(detail_sc)
    hit = [k in canon for k in sc_keys]
    out = detail_sc[pd.Series(hit, index=detail_sc.index)].copy()
    for col in _BLOCK_KEY:
        out[col] = [canon[k][col] for k, ok in zip(sc_keys, hit) if ok]

    miss = set(canon) - {k for k, ok in zip(sc_keys, hit) if ok}
    if miss:
        # บล็อกที่ล็อกไว้แต่ไม่มีใบใน DETAIL_SC — ใส่แถวระดับบล็อกแทน ไม่งั้นแถว OLD หายทั้งก้อน
        fb = locked_detail[pd.Series([k in miss for k in _block_keys(locked_detail)],
                                     index=locked_detail.index)]
        out = pd.concat([out, fb], ignore_index=True)
        print(f"⚠️ {len(fb)} บล็อกที่ล็อกไว้ไม่มีใบใน DETAIL_SC → ใส่แถวระดับบล็อกแทน "
              f"(แถว OLD ของบล็อกนั้นจะยุบใบเหมือนเดิม)")
    return out


def build_booking_demand(unlocked: pd.DataFrame) -> pd.DataFrame:
    """แถว DETAIL ของสัปดาห์ที่ปลดล็อก → ดีมานด์ schema เดียวกับ order_ready

    1 แถว = (SC, item, กำหนดส่ง) ; ยอด = ผลรวม KP_WEIGHT ของสัปดาห์ที่ปลดล็อกเท่านั้น
    (ไม่ใช่ยอดทั้งใบ — ส่วนที่ล็อก/ทอไปแล้วต้องไม่ถูกวางซ้ำ)
    """
    if unlocked.empty:
        return pd.DataFrame()

    d = unlocked.copy()
    # ใบที่ปิดแล้ว (Closed) ไม่ต้องทำ → ไม่ยกมาเป็นดีมานด์ (เครื่องของสัปดาห์นั้นถูกคืนไปแล้ว)
    # booking เก่าที่ยังไม่มีคอลัมน์ STATUS = ถือว่า Pending ทั้งหมด (พฤติกรรมเดิม)
    if "STATUS" in d.columns:
        _st = d["STATUS"].astype(str).str.strip().str.upper()
        _drop = int((~_st.isin(["PENDING", "", "NAN", "NONE"])).sum())
        if _drop:
            print(f"🧹 ตัดงานที่ไม่ใช่ Pending ออกจากดีมานด์: {_drop} แถว "
                  f"({', '.join(sorted(set(_st[~_st.isin(['PENDING', '', 'NAN', 'NONE'])])))})")
        d = d[_st.isin(["PENDING", "", "NAN", "NONE"])]
        if d.empty:
            return pd.DataFrame()
    fr = _str_col(d, "FR_CODE")
    so = _str_col(d, "SO_NO")

    # FR_CODE = 'S{SC_NO}-{ORDER_TYPE}' (ยุบหลายใบด้วย ' | ' ได้) — แหล่งเดียวที่มี ORDER_TYPE
    # ไม่มีคอลัมน์ FR_CODE (booking ไฟล์เก่า) = เดาจาก SO_NO และถือเป็น SC-ORDERS ตามเดิม
    if "FR_CODE" not in d.columns:
        print("❌ ไฟล์ booking ไม่มีคอลัมน์ FR_CODE (สร้างก่อนอัปเดต AVA_MC) → ORDER_TYPE ทุกแถว "
              "ถูกถือเป็น SC-ORDERS ; งาน YD-ORDERS/ผ้าเติม/LAB-DIP จะถูกวางผิดกติกาของตัวเอง "
              "(YD จะไม่ถูกกันไม่ให้ทอก่อนย้อมเสร็จ) — ให้รัน AVA_MC ใหม่ก่อน replan")
    parsed = fr.map(_fr_parse)
    sc_no = parsed.map(lambda t: t[0])
    sc_no = sc_no.where(sc_no != "", "S" + so)
    order_type = parsed.map(lambda t: t[1])
    order_type = order_type.where(order_type != "", "SC-ORDERS")

    work = pd.DataFrame({
        "SC_NO": sc_no,
        "SO_NO": so,
        "ORDER_TYPE": order_type,
        "FABRIC_ITEM": _str_col(d, "ITEM_CODE"),
        "MC_CAT": _str_col(d, "CAT").where(_str_col(d, "CAT") != "", _str_col(d, "MC_CAT")),
        "MC_GROUP": _str_col(d, "MC_GROUP"),
        "MC_GUAGE": _str_col(d, "GUAGE"),
        "FG_YEAR": pd.to_numeric(_s(d, "FG_YEAR"), errors="coerce"),
        "FG_WEEK": pd.to_numeric(_s(d, "FG_WEEK"), errors="coerce"),
        "PERIOD_WEEK": _str_col(d, "RDD_WEEK"),
        "FOB_TYPE": _str_col(d, "FOB_TYPE"),
        "PO_NO": _str_col(d, "PO_NO"),
        # เลขบรรทัดใบจากต้นทาง (view มีคอลัมน์นี้แล้ว) — แถวไหนว่าง (รวมค่า 0 ของงาน CL/YD)
        # ค่อยให้ attach_sc_line_id เดาย้อนจากใบ SC เหมือนเดิม
        "SC_LINE_ID": _s(d, "SC_LINE_ID").map(_line_text),
        # เลขใบจอง — มีเฉพาะฝั่ง booking (แถวที่มาจากใบ SC จะว่าง) ; ต้องไหลถึงชีท PLAN
        "KP_NO": _str_col(d, "KP_NO"),
        "CUSTOMER_NAME": _str_col(d, "KNIT_CUS_NAME").map(_pick_core_customer),
        "SUB_COLOR": _str_col(d, "COLOR"),
        "NAY_COLOR": _str_col(d, "NAY_COLOR"),
        "YARN_ITEM": _str_col(d, "YARN_USED"),
        "MATERIAL_CONTENT": _str_col(d, "MATERIAL_CONTENT"),
        "CAP_PER_DAY": pd.to_numeric(_s(d, "CAP_KNIT"), errors="coerce"),
        "REVOLUTION_WEIGHT": pd.to_numeric(_s(d, "REVOLUTION_WEIGHT"), errors="coerce"),
        "DATE_IN": _s(d, "PO_IN_DATE").map(_last_po_in_date),
        "KP_WEIGHT": pd.to_numeric(_s(d, "KP_WEIGHT"), errors="coerce").fillna(0),
        "WEEK": pd.to_numeric(_s(d, "WEEK"), errors="coerce"),
    })

    # ── เติมลูกค้าให้งาน core ที่ booking ไม่ได้ส่งชื่อลูกค้ามา ─────────────────
    # ⚠ ขาดตาข่ายนี้ไม่ได้: Planning ตัดสิน CORE ITEM จากคู่ (ผ้า, ลูกค้า) ในทะเบียน
    # (Planning.itemcore_pairs) — ลูกค้าว่าง = งาน core กลายเป็นงานสั่งปกติที่วางตาม RDD
    # แบบเงียบ ๆ รอบผลิตสต๊อกพังทั้งชุด ; เจอจริงตอนทดสอบ: ไฟล์ booking ที่สร้างก่อน AVA_MC
    # เพิ่มคอลัมน์ KNIT_CUS_NAME ทำให้ดีมานด์ core 92 แถวไม่มีลูกค้าเลยสักแถว
    # ทะเบียนเก็บทีมคู่กับลูกค้าไว้อยู่แล้ว (RTS→ที คัลเจอร์, NYK1-1→CENTER DOMESTIC) และ
    # booking มี TEAM_NAME ทุกแถว → ย้อนจากทีมได้ ; ทีมที่ยุบหลายค่า ("RTS,NYK1-1") หรือทีมนอก
    # ทะเบียนไม่เดา — ปล่อยว่างไว้เหมือนเดิม (ทีมอื่นสั่งผ้ารหัสเดียวกัน = งานสั่งปกติจริง ๆ)
    _core_team_cus = load_core_item_teams()
    _core_item_cus = load_core_item_customers()
    if _core_team_cus or _core_item_cus:
        _cus_blank = work["CUSTOMER_NAME"].astype(str).str.strip().str.upper().isin(["", "NAN", "NONE"])
        if _cus_blank.any():
            _team = _str_col(d, "TEAM_NAME").str.upper()
            _filled, _filled_solo, _left = 0, 0, {}
            for _i in work.index[_cus_blank]:
                _it = str(work.at[_i, "FABRIC_ITEM"]).strip().upper()
                if _it not in _core_item_cus:
                    continue
                _tm = str(_team.get(_i, "")).strip()
                _cu = _core_team_cus.get((_it, _tm))
                if _cu:
                    work.at[_i, "CUSTOMER_NAME"] = _cu
                    _filled += 1
                    continue
                # ทีมยุบหลายค่า/ไม่ตรงทะเบียน — ผ้าที่มีลูกค้า core คนเดียวยังเติมได้แน่นอน
                _solo = _core_item_cus.get(_it, set())
                if len(_solo) == 1:
                    work.at[_i, "CUSTOMER_NAME"] = next(iter(_solo))
                    _filled_solo += 1
                else:
                    _left[_it] = _tm
            if _filled or _filled_solo:
                print(f"👥 เติมลูกค้าให้งาน core {_filled + _filled_solo} แถว "
                      f"(จากทีม {_filled} / จากผ้าที่มีลูกค้า core คนเดียว {_filled_solo}) "
                      f"— booking ไม่ได้ส่ง KNIT_CUS_NAME มา")
            if _left:
                _sample = ", ".join(f"{k}[{v or 'ไม่มีทีม'}]" for k, v in list(_left.items())[:8])
                print(f"⚠️ งาน core {len(_left)} รหัสยังไม่มีลูกค้า (ทีมไม่ตรงทะเบียน/ยุบหลายทีม) "
                      f"→ จะถูกวางเป็นงานสั่งปกติตาม RDD ไม่ใช่รอบผลิตสต๊อก: {_sample}"
                      + (" ..." if len(_left) > 8 else ""))

    # YD-ORDERS (ย้อมด้าย): คนละสี = คนละด้าย ทอต่อกันรวดเดียวไม่ได้ ต้องเปลี่ยนด้าย (เสีย 1 วัน)
    # → ห้ามยุบหลายสีเป็นดีมานด์แถวเดียว ต้องแยกสีเหมือนที่ Planning แยกไว้ใน _MERGE_COLOR_KEY
    # order type อื่น = ผ้าดิบย้อมทีหลัง สีไม่มีผลตอนทอ → คีย์ว่าง = รวมได้เหมือนเดิม
    work["_COLOR_KEY"] = work["SUB_COLOR"].where(
        work["ORDER_TYPE"].astype(str).str.strip().str.upper() == "YD-ORDERS", ""
    )
    keys = ["SC_NO", "SO_NO", "ORDER_TYPE", "FABRIC_ITEM", "FG_YEAR", "FG_WEEK", "PERIOD_WEEK", "_COLOR_KEY"]
    agg = (
        work.groupby(keys, dropna=False)
        .agg(
            PENDING_PLAN=("KP_WEIGHT", "sum"),
            BOOKING_WEEK_FROM=("WEEK", "min"),
            BOOKING_WEEK_TO=("WEEK", "max"),
            # ทุกตัวใช้ _first_text ไม่ใช่ "first" — ดูเหตุผลที่ _first_text (ช่องว่างบังค่าจริง)
            MC_CAT=("MC_CAT", _first_text),
            MC_GROUP=("MC_GROUP", _first_text),
            MC_GUAGE=("MC_GUAGE", _first_text),
            FOB_TYPE=("FOB_TYPE", _first_text),
            PO_NO=("PO_NO", _first_text),
            # เลขใบจอง — หายทั้งคอลัมน์มาตลอดเพราะไม่ได้อยู่ในรายการนี้ (บั๊กเดียวกับ SC_LINE_ID
            # วัดจริง 21/9/2569: ว่าง 986/986 แถวในชีทแผน ทั้งที่ booking มีให้ 331 แถว)
            KP_NO=("KP_NO", _uniq_kp_text),
            # ⚠ ต้องอยู่ในรายการนี้ ไม่งั้นคอลัมน์หายทั้งคอลัมน์ตอนยุบแถว แล้ว
            # _align_to_order_schema จะเติมกลับมาเป็นค่าว่าง → เลขบรรทัดที่ booking ส่งมาครบ
            # ถูกทิ้งทั้งหมด เหลือแต่ที่ attach_sc_line_id เดาย้อนได้ (วัดจริง 21/9/2569:
            # "จาก booking view 0 | เดาย้อนจากใบ SC 182" ทั้งที่ DETAIL มีเลขครบ 1,106/1,126 แถว)
            # 1 แถวดีมานด์ยุบได้หลายบรรทัด → รวมไม่ซ้ำคั่นคอมมา (กติกาเดียวกับ AVA_MC._uniq_lines)
            SC_LINE_ID=("SC_LINE_ID", _uniq_line_text),
            CUSTOMER_NAME=("CUSTOMER_NAME", _first_text),
            SUB_COLOR=("SUB_COLOR", _first_text),
            NAY_COLOR=("NAY_COLOR", _first_text),
            YARN_ITEM=("YARN_ITEM", _first_text),
            MATERIAL_CONTENT=("MATERIAL_CONTENT", _first_text),
            CAP_PER_DAY=("CAP_PER_DAY", "max"),
            REVOLUTION_WEIGHT=("REVOLUTION_WEIGHT", "max"),
            DATE_IN=("DATE_IN", "max"),
            # POINT DATE = วันด้ายเข้าครบของก้อนนี้ — ทอก่อนวันนี้ไม่ได้ (ไม่มีด้าย)
            # แยกคอลัมน์จาก DATE_IN เพราะ DATE_IN มีทั้งฝั่ง SC ด้วย (Planning ใช้เรียงลำดับ order)
            # ส่วน POINT_DATE มีเฉพาะแถวที่ยกมาจาก booking → Planning บังคับเฉพาะแถวเหล่านี้
            POINT_DATE=("DATE_IN", "max"),
        )
        .reset_index()
    )
    agg = agg[agg["PENDING_PLAN"] > 0].copy()
    _yd_rows = int((agg["_COLOR_KEY"].astype(str) != "").sum())
    if _yd_rows:
        print(f"🎨 YD-ORDERS แยกตามสี: {_yd_rows} แถว (สีต่างกัน = ดีมานด์คนละแถว)")
    agg = agg.drop(columns=["_COLOR_KEY"])

    # สรุปงานที่ด้ายยังไม่เข้า — Planning รอบ replan จะเลื่อนไปเริ่มหลังวันด้ายเข้า (ดู get_point_date_start_idx)
    _pt = pd.to_datetime(agg["POINT_DATE"], errors="coerce")
    # กติกาตั้งที่หน้า Monitor Yarn (KP_YARN_POIN_*) — Planning อ่านค่าชุดเดียวกัน
    _pt_late = int((_pt > pd.Timestamp.today().normalize()).sum())
    if _pt_late:
        if plan_config.get("KP_YARN_POIN_ON"):
            _lag = int(plan_config.get("KP_YARN_POIN_LAG_WEEKS"))
            print(f"🧵 ด้ายยังไม่เข้า {_pt_late} แถว (POINT DATE ถึง {_pt.max():%d/%m/%Y}) "
                  f"— เริ่มทอได้หลังสัปดาห์ที่ด้ายเข้า {_lag} สัปดาห์")
        else:
            print(f"🧵 ด้ายยังไม่เข้า {_pt_late} แถว (POINT DATE ถึง {_pt.max():%d/%m/%Y}) "
                  f"— ปิดกติกาวันด้ายเข้าไว้ (หน้า Monitor Yarn) วางตามกำลังเครื่องได้เลย")

    # FGWEEK = composite YYYYWW เหมือน order_ready (Planning ใช้ตัวนี้เป็นกำหนดส่ง/คิด TARGET_KNIT)
    # กำหนดส่ง: FG ของต้นทางก่อน ; ไม่มี FG → ใช้ RDD (PERIOD_WEEK) คิด TARGET_KNIT แทน
    # (user 24/9/2569 — เดิม 22/9 ไม่มี FG = ไม่วาง ทำให้ booking 323 แถว ≈ 1,157 ตัน หลุดแผนทั้งก้อน)
    #   ช่อง FG_WEEK ของแผนจึงโชว์ค่า RDD (user 24/9/2569: "เอาค่า RDD ไปใส่ใน FG" — ไม่ให้ว่าง)
    #   ธง DUE_FROM_RDD → Planning ติดป้าย DUE_BASIS='RDD' ให้แยกออกจาก FG ที่มาจากต้นทางจริง
    #   ช่อง RDD_WEEK ของแผนมาจาก PERIOD_WEEK ตรง ๆ ไม่ถูกแตะ
    # ไม่มีทั้ง FG และ RDD = ไม่วาง → ชีท NO_FG ของ booking_demand.xlsx ให้ส่งต้นทางแก้
    def _fgweek(r):
        y, w = r["FG_YEAR"], r["FG_WEEK"]
        if pd.notna(y) and pd.notna(w):
            return int(y) * 100 + int(w)
        return pd.NA

    def _rddweek(v):
        """PERIOD_WEEK ('2026/16' หรือ 202616) → YYYYWW ; อ่านไม่ได้ = NA (กติกาเดียวกับ
        Planning._rdd_week_of)"""
        s = str(v if v is not None else "").strip()
        m = (re.match(r"^(\d{4})\D+(\d{1,2})$", s)
             or re.match(r"^(\d{4})(\d{2})(?:\.0+)?$", s))
        if m and 1 <= int(m.group(2)) <= 53:
            return int(m.group(1)) * 100 + int(m.group(2))
        return pd.NA

    if agg.empty:
        return agg
    agg["FGWEEK"] = agg.apply(_fgweek, axis=1)
    _use_rdd = agg["FGWEEK"].isna()
    agg["DUE_FROM_RDD"] = 0
    if _use_rdd.any():
        agg.loc[_use_rdd, "FGWEEK"] = agg.loc[_use_rdd, "PERIOD_WEEK"].map(_rddweek)
        _by_rdd = _use_rdd & agg["FGWEEK"].notna()
        agg.loc[_by_rdd, "DUE_FROM_RDD"] = 1
        if _by_rdd.any():
            print(f"📅 booking ไม่มี FG week {int(_by_rdd.sum())} แถว "
                  f"({float(agg.loc[_by_rdd, 'PENDING_PLAN'].sum()):,.0f} กก.) → วางแผนตาม RDD แทน "
                  f"(ช่อง FG ในแผน = RDD, ติดป้าย DUE_BASIS=RDD)")
    no_fg = agg[agg["FGWEEK"].isna()].drop(columns=["FGWEEK", "DUE_FROM_RDD"])
    if not no_fg.empty:
        print(f"⚠️ booking ไม่มีทั้ง FG และ RDD จากต้นทาง {len(no_fg)} แถว "
              f"({float(no_fg['PENDING_PLAN'].sum()):,.0f} กก.) → ไม่วาง ; "
              f"รายชื่ออยู่ชีท NO_FG ของ {DEMAND_FILE.name} (ต้องแก้ที่ต้นทาง)")
        agg = agg[agg["FGWEEK"].notna()].copy()
    agg["FGWEEK"] = agg["FGWEEK"].astype("int64")
    # เตือนตั้งแต่ตอนสร้างดีมานด์ ถ้าชนิดออร์เดอร์ไม่ใช่ตัวที่ Planning มี routing ให้ —
    # ปล่อยผ่านไปจะเงียบจนกลายเป็นแถวใน UNPLANNED ("routing ไม่ครอบคลุม") ซึ่งไล่ย้อนยาก
    _ot = agg["ORDER_TYPE"].astype(str).str.strip()
    _bad = agg[~_ot.isin(PLANNABLE_ORDER_TYPES)]
    if not _bad.empty:
        _kinds = sorted(set(_bad["ORDER_TYPE"].astype(str)))
        print(f"⚠️ ORDER_TYPE ที่ Planning ไม่รองรับ {len(_bad)} แถว "
              f"({float(_bad['PENDING_PLAN'].sum()):,.0f} กก.) — จะตกเป็น UNPLANNED: "
              f"{', '.join(_kinds[:5])}" + (" ..." if len(_kinds) > 5 else ""))

    agg["ORDERS_QTY"] = agg["PENDING_PLAN"]
    agg["PLAN_QTY"] = 0
    agg["UOM"] = "KG"
    agg["SOURCE_FILE"] = BOOKING_DEMAND_SOURCE
    agg = agg.drop(columns=["FG_YEAR", "FG_WEEK"])
    # แถวที่ไม่ได้วางเพราะไม่มี FG — ผู้เรียกเขียนลงชีท NO_FG (อ่านทันทีหลังเรียก ก่อนแปลงตารางต่อ)
    agg.attrs["no_fg"] = no_fg.drop(columns=["FG_YEAR", "FG_WEEK"])
    return agg


def _align_to_order_schema(demand: pd.DataFrame, orders: pd.DataFrame) -> pd.DataFrame:
    """เติมคอลัมน์ที่ order_ready มีแต่ booking ไม่มี ให้ว่าง แล้วเรียงคอลัมน์ให้ตรงกัน
    (คอลัมน์เสริมของ replan เช่น BOOKING_WEEK_FROM/TO ต่อท้ายไว้ ไม่รบกวน Planning)"""
    if demand.empty:
        return demand
    for c in orders.columns:
        if c not in demand.columns:
            demand[c] = pd.NA
    extra = [c for c in demand.columns if c not in orders.columns]
    return demand[list(orders.columns) + extra]


def _dye_map(df: pd.DataFrame) -> dict:
    """ตาราง SC → {(SC, item): วันย้อมด้ายเสร็จช้าสุด} ; แถวที่ไม่มีวันย้อมข้ามไป"""
    out: dict = {}
    if df is None or df.empty or "DYE_END_DATE" not in df.columns:
        return out
    _item_col = next((c for c in ("FABRIC_ITEM", "ITEM_CODE") if c in df.columns), None)
    if _item_col is None or "SC_NO" not in df.columns:
        return out
    dye = pd.to_datetime(df["DYE_END_DATE"], errors="coerce")
    key = df["SC_NO"].map(_sc_norm) + "|" + df[_item_col].astype(str).str.strip().str.upper()
    for k, v in zip(key, dye):
        if pd.isna(v):
            continue
        out[k] = max(out[k], v) if k in out else v
    return out


def _dye_map_from_sc_files() -> dict:
    """วันย้อมด้ายจากไฟล์ SC ดิบในโฟลเดอร์ Order (ก่อน Order.py กรอง)

    ใบ YD ที่ถูกกรองออกจาก order_ready (CL/FQC/blocklist) ยังถูก booking ยกมาวางใหม่ได้
    → ถ้าไม่อ่านไฟล์ดิบด้วย ใบพวกนี้จะกลายเป็น YD ที่ไม่รู้วันย้อมทั้งที่ต้นทางมีข้อมูลอยู่
    """
    out: dict = {}
    if not ORDER_SRC_DIR.exists():
        return out
    for f in sorted(ORDER_SRC_DIR.iterdir()):
        if f.suffix.lower() not in (".xlsx", ".xls", ".csv"):
            continue
        try:
            df = pd.read_csv(f) if f.suffix.lower() == ".csv" else pd.read_excel(f)
        except Exception as e:
            print(f"⚠️ อ่านไฟล์ SC ดิบไม่ได้ (ข้ามไป): {f.name} — {e}")
            continue
        df.columns = df.columns.str.strip().str.upper()
        for k, v in _dye_map(df).items():
            out[k] = max(out[k], v) if k in out else v
    return out


def attach_dye_end_date(demand: pd.DataFrame, orders: pd.DataFrame) -> pd.DataFrame:
    """เติม DYE_END_DATE (วันย้อมด้ายเสร็จ) ให้ดีมานด์ที่ยกมาจาก booking

    booking view ไม่มีฟิลด์วันย้อมด้ายเลย (ดู View_Booking.SELECT_COLUMNS) → งาน YD-ORDERS
    ที่ยกมาวางใหม่จะไม่มีค่านี้สักแถว ทั้งที่มันคือกฎหลักของ YD ("ห้ามเริ่มทอก่อนย้อมเสร็จ
    +1 สัปดาห์" — Planning) พอไม่มี Planning จะถอยไปใช้ FG week − 3 แทน ซึ่งเร็วกว่าความจริงได้
    วันย้อมมีอยู่ฝั่งใบ SC → ยกมาให้ด้วยคีย์ (SC, item) จาก 2 แหล่งตามลำดับ:
        1. order_ready (ใบที่เข้าแผนปกติ)  2. ไฟล์ SC ดิบในโฟลเดอร์ Order (ใบที่ถูกกรองออก)

    ทุกใบ YD ในระบบ SC มีวันย้อมครบ (ตรวจ view_sc: YD 100% มีค่า ส่วนใบที่ว่างคือชนิดอื่น)
    → YD ที่ยังหาไม่เจอหลังจากนี้ = ผิดปกติ ไม่ใช่เคสปกติ ต้องเตือนให้ไปตรวจต้นทาง

    1 แถว booking ครอบได้หลายใบ (AVA_MC ยุบ SO ที่ลงเครื่องเดียวกันมาช่องเดียว) → ต้องเทียบ
    ทุกใบในช่อง SO_NO เหมือน dedup_against_sc ; จับคู่ได้หลายใบ → เอา "วันช้าสุด" เพราะต้อง
    ย้อมเสร็จครบทุกใบก่อนถึงทอก้อนนี้ได้ทั้งก้อน

    ⚠️ ต้องเรียกก่อน dedup_against_sc: ถ้า booking ชนะ ใบ SC ฝั่งนั้นถูกตัดทิ้ง วันย้อมจะหายไปด้วย
    """
    if demand.empty:
        return demand

    dye_by_key = _dye_map(orders)
    _n_ready = len(dye_by_key)
    for k, v in _dye_map_from_sc_files().items():
        if k not in dye_by_key:
            dye_by_key[k] = v

    if "DYE_END_DATE" not in demand.columns:
        demand["DYE_END_DATE"] = pd.NaT
    cur = pd.to_datetime(demand["DYE_END_DATE"], errors="coerce")
    d_item = demand["FABRIC_ITEM"].astype(str).str.strip().str.upper()
    # เติมเฉพาะ YD-ORDERS — order type อื่นไม่ได้ใช้วันย้อมตัดสินอะไร (ผ้าดิบย้อมทีหลัง)
    # เติมไปก็ไม่มีใครอ่าน แต่ทำให้ค่าที่โผล่ในไฟล์ตีความยากว่ามาจากใบไหน
    _yd = demand["ORDER_TYPE"].astype(str).str.strip().str.upper() == "YD-ORDERS"

    filled = []
    for sc, so, it, have, is_yd in zip(
        demand["SC_NO"], _str_col(demand, "SO_NO"), d_item, cur, _yd
    ):
        if pd.notna(have) or not is_yd:
            filled.append(have)
            continue
        cands = [
            dye_by_key[f"{n}|{it}"]
            for n in ({_sc_norm(sc)} | {_sc_norm(t) for t in _so_tokens(so)})
            if n and f"{n}|{it}" in dye_by_key
        ]
        filled.append(max(cands) if cands else pd.NaT)
    demand["DYE_END_DATE"] = pd.Series(filled, index=demand.index)

    # รายงานเฉพาะ YD — order type อื่นไม่ได้ใช้วันย้อมตัดสินสัปดาห์เริ่มทอ
    _is_yd = demand["ORDER_TYPE"].astype(str).str.strip().str.upper() == "YD-ORDERS"
    n_yd = int(_is_yd.sum())
    if n_yd:
        n_known = int((_is_yd & demand["DYE_END_DATE"].notna()).sum())
        print(f"🧶 YD-ORDERS จาก booking {n_yd} แถว → รู้วันย้อมด้ายเสร็จ {n_known} แถว "
              f"(คีย์ SC+item จาก order_ready {_n_ready} ใบ + ไฟล์ SC ดิบ {len(dye_by_key) - _n_ready} ใบ)")
        if n_known < n_yd:
            _miss = demand[_is_yd & demand["DYE_END_DATE"].isna()]
            _miss_scs = sorted(set(_miss["SC_NO"].astype(str)))
            # ปกติ: view_sc เก็บเฉพาะใบที่ยังค้างวางแผน ใบที่ confirm แล้วหลุดออกไป แต่ booking
            # ยังมีงานค้างอยู่ → หาวันย้อมไม่เจอเป็นเรื่องธรรมดา ไม่ใช่ข้อมูลพัง
            # งานพวกนี้ยังมีพื้นกันอยู่: วันด้ายเข้า (POINT_DATE = PO_IN_DATE ของ booking)
            # ซึ่งมีครบทุกแถวและช้ากว่าวันย้อมเสร็จอยู่แล้วโดยธรรมชาติ (ย้อมเสร็จ → ส่งเข้าโรงทอ)
            print(f"ℹ️ YD ที่ไม่มีวันย้อมในไฟล์ SC {len(_miss)} แถว "
                  f"({float(pd.to_numeric(_miss['PENDING_PLAN'], errors='coerce').fillna(0).sum()):,.0f} กก.) "
                  f"— ใบที่ confirm แล้วหลุดจาก view_sc ; ใช้วันด้ายเข้า (PO_IN_DATE) เป็นพื้นแทน: "
                  f"{', '.join(_miss_scs[:8])}" + (" ..." if len(_miss_scs) > 8 else ""))
    return demand


SC_REGISTRY_FILE = DATA_PLAN_DIR / "sc_registry.xlsx"   # ทะเบียนสะสมของ View_SC (ไม่ถูก purge)


def _sc_line_maps(df: pd.DataFrame, item_col: str) -> dict:
    """ตาราง SC → {'sc|item|po': [line,...]} ; แถวที่ไม่มีเลขบรรทัดข้ามไป

    ⚠ ต้องผ่าน _line_tokens (ไม่ใช่ so_text ดิบ): ทำความสะอาดรูปแบบ (".0" ของ pandas) และ
    แตกค่าที่ยุบคอมมามาแล้ว ; เลข 0 ของใบ CL/YD ในทะเบียน sc_registry (วัดจริง 626/1,353 แถว)
    ถูกส่งต่อไปตามต้นทางตามที่ user สั่ง ไม่ทิ้ง

    ⚠ คีย์ต้องมี PO เสมอ — ไม่มีตารางสำรองแบบ (SC, item) แล้ว (user สั่ง 21/9/2569)
    ใบเดียวแตกหลายบรรทัดตาม PO ได้ การ match แบบไม่ดู PO จะยัดเลขบรรทัดของ PO อื่นเข้ามาด้วย
    = เดาข้อมูลขึ้นเอง ; ไม่ตรง PO ให้ปล่อยว่างตามจริง
    """
    by_po: dict = {}
    if df is None or df.empty or item_col not in df.columns or "SC_LINE_ID" not in df.columns:
        return by_po
    for sc, it, po, lns in zip(
        _s(df, "SC_NO").map(_sc_norm),
        df[item_col].astype(str).str.strip().str.upper(),
        _str_col(df, "PO_NO").str.upper(),
        df["SC_LINE_ID"].map(_line_tokens),
    ):
        if not sc or not lns:
            continue
        k = f"{sc}|{it}|{po}"
        for ln in lns:
            if ln not in by_po.setdefault(k, []):
                by_po[k].append(ln)
    return by_po


def attach_sc_line_id(demand: pd.DataFrame, orders: pd.DataFrame) -> pd.DataFrame:
    """เติม SC_LINE_ID (เลขบรรทัดใบ) ให้ดีมานด์ booking ที่ยังไม่มีเลขบรรทัด

    booking view มีคอลัมน์ SC_LINE_ID แล้ว (เพิ่มทีหลัง) และ build_booking_demand ยกมาให้ตรง ๆ
    → ปกติแถวส่วนใหญ่มีค่ามาแต่ต้น ตัวนี้เหลือหน้าที่เป็น "ตาข่ายรอง" 2 กรณี: ไฟล์ booking รุ่นเก่า
    ที่สร้างก่อนเพิ่มคอลัมน์ (เดิมว่าง 777/800 แถว) และแถวที่ต้นทางไม่ได้กรอกเลขบรรทัดมา
    แถวที่มีค่าอยู่แล้วจะถูกข้าม (`if have`) — ค่าจากต้นทางชนะเสมอ
    → ยกมาจากฝั่งใบ SC ตามลำดับแหล่ง: 1. order_ready  2. ทะเบียน sc_registry (ใบที่หลุด view ไปแล้ว)
    คีย์เดียวคือ (SC, item, PO) ตรงเป๊ะ — ไม่ตรงก็ปล่อยว่าง ห้ามถอยไปเดาด้วย (SC, item)
    เพราะใบเดียวแตกหลายบรรทัดตาม PO ได้ การ match หลวม ๆ จะได้เลขบรรทัดของ PO อื่นติดมา
    1 แถว booking ครอบหลายใบได้ → รวมเลขบรรทัดที่ match ตรงคีย์คั่นคอมมา

    ⚠️ ต้องเรียกก่อน dedup_against_sc เหมือน attach_dye_end_date
    """
    if demand.empty:
        return demand

    sources = [_sc_line_maps(orders, "FABRIC_ITEM")]
    if SC_REGISTRY_FILE.exists():
        try:
            sources.append(_sc_line_maps(pd.read_excel(SC_REGISTRY_FILE, sheet_name="SC_REGISTRY"), "ITEM_CODE"))
        except Exception as e:
            print(f"⚠️ อ่าน {SC_REGISTRY_FILE.name} ไม่ได้ ({e}) — หาเลขบรรทัดจาก order_ready อย่างเดียว")

    # ค่าที่ติดมาจากต้นทางต้องผ่าน _line_text เหมือนทุกจุด — เลข 0 ของงาน CL/YD ไม่ใช่ "มีค่าแล้ว"
    cur = _str_col(demand, "SC_LINE_ID").map(_line_text)
    d_item = demand["FABRIC_ITEM"].astype(str).str.strip().str.upper()
    filled = []
    for sc, so, it, po, have in zip(demand["SC_NO"], _str_col(demand, "SO_NO"), d_item,
                                    _str_col(demand, "PO_NO").str.upper(), cur):
        if have:
            filled.append(have)
            continue
        scs = [n for n in dict.fromkeys([_sc_norm(sc)] + [_sc_norm(t) for t in sorted(_so_tokens(so))]) if n]
        pos = [p.strip() for p in po.split(",") if p.strip()]
        lines: list = []
        for by_po in sources:
            for n in scs:
                hit = [ln for p in pos for ln in by_po.get(f"{n}|{it}|{p}", [])]
                lines += [ln for ln in hit if ln not in lines]
            if lines:
                break
        filled.append(",".join(lines))
    demand["SC_LINE_ID"] = pd.Series(filled, index=demand.index, dtype=object)

    n_miss = int((demand["SC_LINE_ID"] == "").sum())
    n_src = int((cur != "").sum())         # มาจาก booking view ตรง ๆ
    n_join = len(demand) - n_miss - n_src  # ยกจากใบ SC ด้วยคีย์ SC+item+PO ที่ตรงเป๊ะ
    print(f"🔢 SC_LINE_ID ของดีมานด์ booking: มีเลขบรรทัด {len(demand) - n_miss}/{len(demand)} แถว "
          f"(จาก booking view {n_src} | จากใบ SC คีย์ SC+item+PO {n_join})"
          + (f" — ไม่มีเลขบรรทัดที่ต้นทาง {n_miss} แถว (ปล่อยว่างตามจริง ไม่เดาให้)" if n_miss else ""))
    return demand


def dedup_against_sc(demand: pd.DataFrame, orders: pd.DataFrame):
    """ตัดงานซ้ำระหว่าง booking กับ SC ที่ระดับ (SC, FABRIC_ITEM)

    1 แถว booking ครอบได้หลายใบ (AVA_MC ยุบ SO ที่ลงเครื่องเดียวกันมาไว้ช่องเดียว
    'S721282-SC-ORDERS | S721257-SC-ORDERS' → SO_NO = '721282,721257')
    → ต้องเทียบทุกใบในช่อง SO_NO ไม่ใช่แค่ SC_NO ใบแรก ไม่งั้นใบที่ 2 ที่มีใน view_sc ด้วย
    จะรอด dedup มาถูกวางซ้ำอีกก้อน

    คืน (orders ที่ตัดแล้ว, demand ที่ตัดแล้ว, ตาราง audit)
    """
    if demand.empty or orders.empty:
        return orders, demand, pd.DataFrame()

    o_item = orders["FABRIC_ITEM"].astype(str).str.strip().str.upper()
    o_key = orders["SC_NO"].map(_sc_norm) + "|" + o_item

    d_item = demand["FABRIC_ITEM"].astype(str).str.strip().str.upper()
    d_keysets = [
        {f"{n}|{it}" for n in ({_sc_norm(sc)} | {_sc_norm(t) for t in _so_tokens(so)}) if n}
        for sc, so, it in zip(demand["SC_NO"], _str_col(demand, "SO_NO"), d_item)
    ]
    dup = set(o_key) & set().union(*d_keysets) if d_keysets else set()
    if not dup:
        return orders, demand, pd.DataFrame()

    o_qty = orders.assign(_k=o_key).groupby("_k")["PENDING_PLAN"].sum()
    # ยอด booking ต่อคีย์ — แถวที่ครอบหลายใบถูกนับให้ทุกใบของแถวนั้น (ตัวเลขไว้ไล่ตรวจ
    # ไม่ได้เอาไปคำนวณต่อ) ; การตัดจริงทำที่ฝั่ง orders ด้วย dup เท่านั้น
    d_qty: dict = {}
    _dp = pd.to_numeric(demand["PENDING_PLAN"], errors="coerce").fillna(0)
    for ks, q in zip(d_keysets, _dp):
        for k in ks:
            d_qty[k] = d_qty.get(k, 0.0) + float(q)
    audit = pd.DataFrame({
        "KEY": sorted(dup),
        "SC_NO": [k.split("|")[0] for k in sorted(dup)],
        "ITEM_CODE": [k.split("|")[1] for k in sorted(dup)],
        "QTY_SC": [float(o_qty.get(k, 0)) for k in sorted(dup)],
        "QTY_BOOKING_UNLOCKED": [float(d_qty.get(k, 0)) for k in sorted(dup)],
    })
    audit["QTY_DIFF"] = audit["QTY_BOOKING_UNLOCKED"] - audit["QTY_SC"]
    audit["USED"] = "BOOKING" if BOOKING_WINS_ON_DUPLICATE else "SC"

    if BOOKING_WINS_ON_DUPLICATE:
        orders = orders[~o_key.isin(dup)].copy()
    else:
        _hit = pd.Series([bool(ks & dup) for ks in d_keysets], index=demand.index)
        demand = demand[~_hit].copy()
    return orders, demand, audit


# =========================
# capacity คู่ขนาน
# =========================
def _locked_mc_use(locked_detail: pd.DataFrame, unlocked_weeks: set):
    """เครื่องที่ยัง "ถูกจองอยู่" ในสัปดาห์ที่ปลดล็อก = ของงานที่ผู้ใช้ล็อกไว้

    คืน (ยอดต่อพูล, ยอดต่อ non-pool) ด้วยสูตรเดียวกับที่ AVA_MC สร้าง SUMMARY_MC_REMAIN:
        pool     → group by POOL_TYPE + WEEK
        non-pool → group by MC_CAT + GUAGE + WEEK
    ถ้าใช้สูตรอื่น ตัวเลขเครื่องจะไม่ตรงกับไฟล์ต้นทางแล้วแผนจะจองเครื่องเกิน/ขาด
    """
    empty = ({}, {})
    if locked_detail.empty or not unlocked_weeks:
        return empty
    d = locked_detail.copy()
    d["_WK"] = pd.to_numeric(_s(d, "WEEK"), errors="coerce")
    d = d[d["_WK"].isin(unlocked_weeks)]
    if d.empty:
        return empty
    d["_MC"] = pd.to_numeric(_s(d, "MC_USE_CEIL"), errors="coerce").fillna(0)
    d["_POOL"] = _str_col(d, "POOL_TYPE")
    d["_CAT"] = _str_col(d, "MC_CAT")
    d["_G"] = _s(d, "GUAGE").map(_norm_gauge)

    pool = d[d["_POOL"] != ""].groupby(["_POOL", "_WK"])["_MC"].sum()
    nonpool = d[d["_POOL"] == ""].groupby(["_CAT", "_G", "_WK"])["_MC"].sum()
    return (
        {(k[0], int(k[1])): float(v) for k, v in pool.items()},
        {(k[0], k[1], int(k[2])): float(v) for k, v in nonpool.items()},
    )


def _reserved_weekly_from_detail(det: pd.DataFrame) -> pd.DataFrame:
    """MC_RESERVED_WEEKLY (เครื่อง POLY/COTTON ที่ booking กินไปแล้วต่อสัปดาห์) จาก detail
    ที่คืนเครื่องแล้ว — นับเฉพาะงานที่ยัง "ตรึง" อยู่ ตรรกะเดียวกับ AVA_MC._reserved_weekly_usage

    ⚠️ ห้ามก็อปชีทเดิมจาก booking_final มาตรง ๆ : งานที่ถูกยกออกมาวางใหม่จะยังอยู่ในชีทนั้น
    แล้วโผล่เป็นแถวแผนอีกรอบ → หน้าเว็บ (plan_view._reserved_used_by_cat_week) นับซ้ำ
    ชิป 🔒 จึงโชว์ "เกิน" ทั้งที่แผนไม่ได้วางเกินโควตา
    """
    cols = ["MC_CAT", "GUAGE", "MC_POOL", "WEEK", "POLY_USED", "COTTON_USED"]
    if det is None or det.empty or "POOL_TYPE" not in det.columns:
        return pd.DataFrame(columns=cols)

    def _usage(sfx: str, name: str) -> pd.DataFrame:
        sub = det[det["POOL_TYPE"].astype(str).str.endswith(sfx)].copy()
        if sub.empty:
            return pd.DataFrame(columns=["MC_CAT", "GUAGE", "MC_POOL", "WEEK", name])
        # MC_POOL = POOL_TYPE ตัด suffix :POLY/:COTTON ออก → ป้ายพูลเดียวกับ SUMMARY
        sub["MC_POOL"] = sub["POOL_TYPE"].astype(str).str[:-len(sfx)]
        return (
            sub.groupby(["MC_CAT", "GUAGE", "MC_POOL", "WEEK"], as_index=False)["MC_USE_CEIL"]
            .sum().rename(columns={"MC_USE_CEIL": name})
        )

    out = _usage(":POLY", "POLY_USED").merge(
        _usage(":COTTON", "COTTON_USED"),
        on=["MC_CAT", "GUAGE", "MC_POOL", "WEEK"], how="outer",
    )
    if out.empty:
        return pd.DataFrame(columns=cols)
    for _c in ("MC_CAT", "GUAGE", "MC_POOL"):
        out[_c] = out[_c].astype(str).str.strip()
    out["WEEK"] = pd.to_numeric(out["WEEK"], errors="coerce").fillna(0).astype(int)
    for _c in ("POLY_USED", "COTTON_USED"):
        out[_c] = pd.to_numeric(out.get(_c), errors="coerce").fillna(0).astype(int)
    return out[cols].sort_values(["MC_CAT", "GUAGE", "MC_POOL", "WEEK"])


def write_replan_capacity(src: Path, locked_detail: pd.DataFrame, unlocked_weeks: set,
                          unlocked_yw: set | None = None, pinned_mc=None,
                          pinned_jobs: dict | None = None,
                          locked_sc: pd.DataFrame | None = None,
                          pinned_item_jobs: dict | None = None) -> None:
    """สร้างไฟล์ booking_final_replan.xlsx — เครื่องของสัปดาห์ที่ปลดล็อกถูกคืนเป็นว่าง

    ⚠️ คืนเฉพาะเครื่องของงานที่ "ถูกยกออกมาวางใหม่" เท่านั้น — งานที่ผู้ใช้ล็อกไว้
    (ชีท Lock_Plan) ยังอยู่ใน locked_detail และยังกินเครื่องของสัปดาห์นั้นอยู่
    ถ้าคืนเหมาทั้งสัปดาห์เหมือนเดิม เครื่องของงานที่ล็อกจะถูกงานอื่นแย่งไปใช้ = ชนกันจริง

    Planning อ่านจากไฟล์นี้ 6 ชีท: DETAIL / DETAIL_SC / SUMMARY_MC_REMAIN / MC_POOL_MAP / PINNED_JOBS /
    PINNED_JOBS_ITEM (= PINNED_JOBS แยกรายผ้า สำหรับโควตาตั้งเครื่อง 2 ตัว/ผ้า/สัปดาห์)
    (MC_RESERVED ก็อปมาเฉย ๆ เพราะเป็นโควตาคงที่ ; MC_RESERVED_WEEKLY ต้องคำนวณใหม่จาก
    locked_detail — หน้าเว็บใช้ชีทนี้เทียบโควตา 🔒 ถ้าก็อปของเดิมจะนับงานที่ยกมาวางใหม่ซ้ำ)

    ⚠️ ห้ามยัดงานที่ตรึงไว้เข้า DETAIL เพื่อให้มันกินโควตา job — Planning ใช้ DETAIL อีกหลายที่
    (seed carryover/last_production, สร้างแถว OLD, นับเครื่องที่ booking ใช้อยู่) แถวเทียมจะ
    ไปโผล่ทุกจุดนั้นและเครื่องจะถูกกันสองเท่ากับ SUMMARY → จึงส่งผ่านชีท PINNED_JOBS แยก
    """
    xls = pd.ExcelFile(src)
    summary = xls.parse("SUMMARY_MC_REMAIN")
    summary.columns = summary.columns.str.strip().str.upper()
    # คีย์สัปดาห์ทั้งหมดในฟังก์ชันนี้เป็น YW (locked_detail / unlocked_weeks / งานที่ตรึง แปลงมาแล้ว)
    _to_yw(summary)

    pool_keep, np_keep = _locked_mc_use(locked_detail, unlocked_weeks)

    wk = pd.to_numeric(summary["WEEK"], errors="coerce")
    # WEEK เป็น YW → สัปดาห์เลขเดียวกันคนละปีแยกกันเองอยู่แล้ว (unlocked_yw เก็บไว้เพื่อความเข้ากันได้)
    free = wk.isin(unlocked_weeks)
    n_before = int(pd.to_numeric(summary.loc[free, "MC_USE_CEIL"], errors="coerce").fillna(0).sum())

    s_pool = summary["MC_POOL"].fillna("").astype(str).str.strip() if "MC_POOL" in summary.columns \
        else pd.Series("", index=summary.index)
    s_cat = summary["MC_CAT"].fillna("").astype(str).str.strip()
    s_g = summary["GUAGE"].map(_norm_gauge)

    if pool_keep or np_keep:
        keep = pd.Series(0.0, index=summary.index)
        for i in summary.index[free]:
            w = int(wk.at[i])
            keep.at[i] = (pool_keep.get((s_pool.at[i], w), 0.0) if s_pool.at[i]
                          else np_keep.get((s_cat.at[i], s_g.at[i], w), 0.0))
        summary.loc[free, "MC_USE_CEIL"] = keep[free]
    else:
        summary.loc[free, "MC_USE_CEIL"] = 0

    summary.loc[free, "TOTAL_MC_REMAIN"] = (
        pd.to_numeric(summary.loc[free, "TOTAL_MC"], errors="coerce")
        - pd.to_numeric(summary.loc[free, "MC_USE_CEIL"], errors="coerce").fillna(0)
    )
    n_keep = int(pd.to_numeric(summary.loc[free, "MC_USE_CEIL"], errors="coerce").fillna(0).sum())
    print(f"🔓 คืนเครื่องให้สัปดาห์ที่ปลดล็อก: {int(free.sum())} แถวใน SUMMARY_MC_REMAIN "
          f"(เดิมถูกจอง {n_before} เครื่อง → เหลือ {n_keep} เครื่องของงานที่ล็อกไว้)")

    # เครื่องของ "งานในแผนที่ถูกตรึง" — บวกเพิ่มทุกสัปดาห์ ไม่ใช่เฉพาะสัปดาห์ที่ปลดล็อก
    # เพราะงานพวกนี้ไม่เคยถูกหักออกจาก booking capacity มาก่อน (ดู pinned_mc_use)
    if pinned_mc and (pinned_mc[0] or pinned_mc[1]):
        p_pool, p_np = pinned_mc
        add = pd.Series(0.0, index=summary.index)
        for i in summary.index:
            w = wk.at[i]
            if pd.isna(w):
                continue
            w = int(w)
            add.at[i] = (p_pool.get((s_pool.at[i], w), 0.0) if s_pool.at[i]
                         else p_np.get((s_cat.at[i], s_g.at[i], w), 0.0))
        hit = add > 0
        if hit.any():
            summary.loc[hit, "MC_USE_CEIL"] = (
                pd.to_numeric(summary.loc[hit, "MC_USE_CEIL"], errors="coerce").fillna(0) + add[hit]
            )
            summary.loc[hit, "TOTAL_MC_REMAIN"] = (
                pd.to_numeric(summary.loc[hit, "TOTAL_MC"], errors="coerce")
                - pd.to_numeric(summary.loc[hit, "MC_USE_CEIL"], errors="coerce").fillna(0)
            )
            # นับจากยอดที่กันไว้จริง ไม่ใช่ผลรวมข้ามแถว — พูลหนึ่งมีได้หลายแถวสมาชิก
            # (เช่น 'PHET|IIP 20G , RL 18G' มีทั้งแถว IIP และแถว RL) และทุกแถวถือยอดรวม
            # ของพูลเท่ากันตามที่ AVA_MC สร้างไว้ บวกข้ามแถวจะได้ตัวเลขเกินจริง
            print(f"📌 กันเครื่องให้งานที่ตรึงไว้: "
                  f"{int(sum(p_pool.values()) + sum(p_np.values()))} เครื่อง "
                  f"(แตะ {int(hit.sum())} แถวใน SUMMARY_MC_REMAIN)")

    # โควตา job ที่งานตรึงกินไปแล้ว — เขียนเสมอ (แม้ว่าง) ให้ Planning เห็นชีทเป็นรูปเดียวกันทุกครั้ง
    _pj_df = pd.DataFrame(
        [{"WEEK": _w, "MC_GROUP": _m, "JOBS": _j}
         for (_w, _m), _j in sorted((pinned_jobs or {}).items())],
        columns=["WEEK", "MC_GROUP", "JOBS"],
    )
    if not _pj_df.empty:
        print(f"📌 โควตา job ของงานที่ตรึงไว้: {int(_pj_df['JOBS'].sum())} jobs "
              f"({len(_pj_df)} คู่ สัปดาห์×กลุ่มเครื่อง) → ชีท PINNED_JOBS")
    # ชุดเดียวกันแต่แยกรายผ้า — Planning หักโควตาตั้งเครื่องรายผ้า (เขียนเสมอแม้ว่าง เหมือน PINNED_JOBS)
    _pji_df = pd.DataFrame(
        [{"WEEK": _w, "ITEM_CODE": _it, "MC_GROUP": _m, "GUAGE": _g, "JOBS": _j}
         for (_w, _it, _m, _g), _j in sorted((pinned_item_jobs or {}).items())],
        columns=["WEEK", "ITEM_CODE", "MC_GROUP", "GUAGE", "JOBS"],
    )

    # เขียนออกเป็นรูปไฟล์เดิม: WEEK (เลขสัปดาห์) + YEAR — Planning / หน้าเว็บอ่านรูปนี้
    _wy = [("WEEK", "YEAR")]
    with pd.ExcelWriter(BOOKING_REPLAN_FILE, engine="openpyxl") as w:
        _split_yw(locked_detail, _wy).to_excel(w, sheet_name="DETAIL", index=False)
        # DETAIL_SC = ใบสั่งของบล็อกใน DETAIL — Planning ใช้แตกแถว OLD ทีละใบ ต้องเขียนคู่กันเสมอ
        # ไม่มี = แถว OLD ยุบใบ (กำหนดส่ง/ลูกค้า/PO ของแต่ละใบหาย) ; ไม่กระทบการนับเครื่อง
        if locked_sc is not None and not locked_sc.empty:
            _split_yw(locked_sc, _wy).to_excel(w, sheet_name="DETAIL_SC", index=False)
            print(f"🧾 DETAIL_SC ของงานที่ล็อกไว้: {len(locked_sc)} ใบ "
                  f"(จาก {len(locked_detail)} บล็อก) → {BOOKING_REPLAN_FILE.name}")
        _split_yw(summary, _wy).to_excel(w, sheet_name="SUMMARY_MC_REMAIN", index=False)
        _split_yw(_pj_df, _wy).to_excel(w, sheet_name="PINNED_JOBS", index=False)
        _split_yw(_pji_df, _wy).to_excel(w, sheet_name="PINNED_JOBS_ITEM", index=False)
        # MC_RESERVED / MC_POOL_MAP = ค่าคงที่ (โควตาที่กันไว้ + แผนที่พูล) ก็อปมาได้ตรง ๆ
        for sh in ("MC_RESERVED", "MC_POOL_MAP"):
            if sh in xls.sheet_names:
                xls.parse(sh).to_excel(w, sheet_name=sh, index=False)
        # MC_RESERVED_WEEKLY ต้องคำนวณใหม่จาก locked_detail — ห้ามก็อปของเดิม
        _split_yw(_reserved_weekly_from_detail(locked_detail), _wy).to_excel(
            w, sheet_name="MC_RESERVED_WEEKLY", index=False)
    print(f"💾 capacity คู่ขนาน → {BOOKING_REPLAN_FILE.name}")


# =========================
# ด่านก่อนรัน: สัปดาห์ที่ยุบรวมแล้วต้องระบุวัน + ชั่วโมงทำงาน
# =========================
def merge_gaps() -> list:
    """สัปดาห์ที่ยุบรวมกันแล้วยังกรอกไม่ครบ/เกินเพดาน → [(สัปดาห์, เพดาน, {กอง: [ที่ขาด]})]

    สัปดาห์ที่รวมกันแล้วมักเดินไม่เท่าสัปดาห์ปกติ (เช่น ยุบ W31 เข้า W32 แล้วเร่งเป็น 24 ชม.)
    ปล่อยว่างไว้ระบบจะรวมวันให้เองและใช้ชั่วโมงมาตรฐานของกองเงียบ ๆ ซึ่งอาจไม่ใช่สิ่งที่
    ผู้วางแผนตั้งใจ → หยุดไว้ก่อน (ข้ามได้ด้วย KP_ALLOW_MERGE_HOUR_GAP=1)

    เพดานวัน = จำนวนวันที่ปฏิทินเปิดรวมทั้งก้อน (ยุบ 2 สัปดาห์ที่เปิด 6+6 → ไม่เกิน 12)

    อ่านไฟล์ตรง ๆ เพราะ Planning รันเป็น subprocess แยก (WorkDay ของโปรเซสนี้ยังไม่ init)
    กติกาเดียวกับ WorkDay.merge_gaps() / mc_hub._merge_gaps()
    คีย์สัปดาห์ = YW ทั้งหมด (การยุบเป็นของปีนั้น) ; ตรวจเฉพาะสัปดาห์ยุบตั้งแต่สัปดาห์ปัจจุบัน"""
    from Calendar import load_week_merge, load_workday, pin_merge_years
    tl = _timeline()
    wd = load_workday()
    merges = pin_merge_years(load_week_merge(), tl)
    targets = sorted({d for d in merges.values() if tl.pos.get(d, len(tl.keys)) >= tl.today_idx})
    if not targets:
        return []

    mc_path = _master_mc_path()
    if mc_path is None or not Path(mc_path).exists():
        return []                      # ไม่มี MasterMC ให้เทียบ → ไม่ด่วนสรุปว่าขาด

    def _k(v) -> str:
        s = str(v if v is not None else "").strip().upper()
        return s.split(".", 1)[0] if re.fullmatch(r"\d+\.0+", s) else s

    mm = pd.read_excel(mc_path, sheet_name="Master MC")
    mm.columns = [str(c).strip() for c in mm.columns]
    col = {str(c).strip().upper(): c for c in mm.columns}
    need = ("FACTORY", "MC_CAT", "GUAGE", "TOTAL MC")
    if any(c not in col for c in need):
        return []

    totals: dict = {}
    for r in mm.to_dict("records"):
        key = (_k(r.get(col["FACTORY"])), _k(r.get(col["MC_CAT"])), _k(r.get(col["GUAGE"])))
        if not any(key):
            continue
        n = pd.to_numeric(r.get(col["TOTAL MC"]), errors="coerce")
        totals[key] = totals.get(key, 0) + (0 if pd.isna(n) else float(n))

    # วันที่ปฏิทินเปิดของแต่ละสัปดาห์ — คีย์ YW เพราะปฏิทินมีหลายปี (รวมตามเลขเปล่า W41 = 6×5 ปี = 30)
    weekly = calendar_week_map(load_calendar(_CALENDAR_LOCAL_PATH, sheet_name="Sheet1"))
    cal_open: dict = {}
    for r in weekly.itertuples():
        yw = int(r.year) * 100 + int(r.week)
        cal_open[yw] = cal_open.get(yw, 0) + int(r.working_days)
    src_of: dict = {}
    for s, d in merges.items():
        src_of.setdefault(d, []).append(s)

    # แถวรายสัปดาห์ที่ไม่มีปี (แก้ไฟล์มือ) → ครั้งที่ใกล้วันนี้สุด เหมือน WorkDay._wk
    def _yw_keys(d: dict) -> dict:
        return {k[:3] + (int(k[3]) if int(k[3]) >= 100000 else tl.resolve_bare(int(k[3])),): v
                for k, v in d.items()}

    wk, wh = _yw_keys(wd.get("week", {})), _yw_keys(wd.get("week_hour", {}))
    out = []
    for w in targets:
        bucket = [w] + src_of.get(w, [])
        cap = sum(cal_open.get(b, 0) for b in bucket) or None
        miss: dict = {}
        for key, n in totals.items():
            if n <= 0:
                continue
            days = wk.get(key + (w,))
            bad = []
            if days is None:
                bad.append("วัน")
            elif cap is not None and float(days) > cap + 1e-9:
                bad.append(f"วัน {days:g} เกินเพดาน {cap:g}")
            if key + (w,) not in wh:
                bad.append("ชั่วโมง")
            if bad:
                miss["·".join(key)] = bad
        if miss:
            out.append((w, cap, miss))
    return out


# =========================
# รัน Planning รอบ replan
# =========================
def run_planning_replan() -> bool:
    env_extra = {
        "KP_REPLAN": "1",
        "KP_REPLAN_ORDER_FILE": str(ORDER_REPLAN_FILE),
        "KP_REPLAN_MC_FILE": str(BOOKING_REPLAN_FILE),
        "KP_REPLAN_OUTPUT": str(RESULT_FILE),
        "KP_REPLAN_APPEND_FILE": str(APPEND_FILE),
        # ชีท NO_FG = งาน booking ไม่มี FG week → Planning ลงชีท UNPLANNED พร้อมเหตุผล
        "KP_REPLAN_NO_FG_FILE": str(DEMAND_FILE),
    }
    print("=" * 70)
    print("▶ รัน Planning รอบ replan (ดีมานด์รวม booking + SC บน capacity ที่คืนเครื่องแล้ว)")
    print("=" * 70)
    env = os.environ.copy()
    env.update(env_extra)
    # log ของ Replan ค้างใน buffer → พิมพ์ออกก่อน ไม่งั้นไปแทรกกลาง log ของ Planning (เห็นชัดตอนวาง 2 รอบ)
    sys.stdout.flush()
    rc = subprocess.run([sys.executable, str(BASE_DIR / "Planning.py")], cwd=str(BASE_DIR), env=env).returncode
    return rc == 0


def yarn_supply_cap(combined: pd.DataFrame) -> pd.DataFrame:
    """แผนวางได้ไม่เกินด้ายที่มี = สต๊อก + PO Pending (สวิตช์ KP_YARN_SUPPLY_ON หน้า Monitor Yarn) — ดู YarnStockCover.supply_cap

    user 15 ต.ค. 2569: "มันต้องไม่มีติดลบ เราต้องวางแผนแค่ที่มีใน stock แล้วที่เหลือก็ไปตาม POIN" + "ต้องดู PO Pending ด้วย"
    ทำก่อนเขียน order_replan: ใบที่ด้ายมาไม่ครบถูกแตก — ก้อนที่รอ PO ได้ YARN_READY (Planning ห้ามวางก่อน) ;
    ก้อนที่ไม่มีด้ายเลย → ชีท YARN_SHORT ของ booking_demand.xlsx → Planning ลงชีท UNPLANNED พร้อมเหตุผล
    คิดไม่ได้ (ไม่มีไฟล์ / error) = ดีมานด์เดิม แผนยังออกได้ แค่ไม่คุมด้าย"""
    if not plan_config.get("KP_YARN_SUPPLY_ON"):
        print("🧶 วางแผนได้ไม่เกินด้ายที่มี: ปิดอยู่ (หน้า Monitor Yarn)")
        return combined
    import YarnStockCover as ysc
    lag = int(plan_config.get("KP_YARN_POIN_LAG_WEEKS")) if plan_config.get("KP_YARN_POIN_ON") else 0
    try:
        res = ysc.supply_cap(combined, _timeline(), FREEZE_WEEKS_AHEAD, lag,
                             exempt_kw=plan_config.get("KP_YARN_SUPPLY_EXEMPT") or [],
                             local=bool(plan_config.get("KP_YARN_SUPPLY_EXEMPT_LOCAL")))
    except Exception as e:
        print(f"⚠️ [YARN SUPPLY] คิดด้ายที่มีไม่ได้ ({e}) → วางตามดีมานด์เดิม ไม่คุมด้าย")
        return combined
    if "stats" not in res:
        print(f"⚠️ [YARN SUPPLY] {res.get('note')} → วางตามดีมานด์เดิม ไม่คุมด้าย")
        return combined
    if res.get("note"):
        print(f"⚠️ [YARN SUPPLY] {res['note']}")
    st = res["stats"]
    try:
        with pd.ExcelWriter(DEMAND_FILE, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
            (res["audit"] if not res["audit"].empty else pd.DataFrame(columns=["SC"])).to_excel(
                w, sheet_name="YARN_SUPPLY", index=False)
            (res["short"] if not res["short"].empty else pd.DataFrame(columns=["SC_NO"])).to_excel(
                w, sheet_name="YARN_SHORT", index=False)
    except Exception as e:
        # ไม่มีชีท YARN_SHORT = Planning ไม่รู้ว่ามีงานที่ตัดออก → ยอดหายเงียบ ; ยอมวางตามดีมานด์เดิมดีกว่า
        print(f"⚠️ [YARN SUPPLY] เขียนชีท YARN_SUPPLY/YARN_SHORT ไม่ได้ ({e}) → วางตามดีมานด์เดิม ไม่คุมด้าย")
        return combined
    print(f"🧶 วางแผนได้ไม่เกินด้ายที่มี (สต๊อก + PO Pending ; ไม่คุมด้ายในประเทศ {st['exempt_yarns']} ตัว "
          f"= Internal {st['exempt_yarns'] - st['exempt_local']} · Local {st['exempt_local']}): "
          f"ใบที่ต้องใช้ด้ายจากสต๊อก/PO {st['orders']} ใบ — ด้ายพอ {st['full']} · "
          f"รอ PO {st['wait']} ใบ {st['wait_kg']:,.0f} kg · แตกใบ {st['split']} · "
          f"ไม่มีด้าย {st['short_orders']} ใบ {st['short_kg']:,.0f} kg (→ UNPLANNED) "
          f"; รายละเอียดชีท YARN_SUPPLY ใน {DEMAND_FILE.name}")
    return res["demand"]


def yarn_stock_second_pass(combined: pd.DataFrame) -> None:
    """ใช้ด้ายในสต๊อกแทนการรอ PO IN (สวิตช์ KP_YARN_STOCK_ON หน้า Monitor Yarn) — ดู YarnStockCover.py

    สต๊อกลดลงตามที่แผนเอาด้ายไปใช้แต่ละสัปดาห์ (user: "ต้องไม่แย่งกัน เส้นด้ายจะเปลี่ยนไปตามแผน")
    → ต้องมีแผนรอบแรกก่อน (ใบรอ PO IN ตามเดิม) จึงรู้ว่าเหลือด้ายให้ยืมเท่าไหร่ แล้ววางรอบสอง
    รอบสองล้ม = คืนแผนรอบแรก (แผนยังออกได้ แค่ไม่มีใบไหนได้ยืมสต๊อก)"""
    if not plan_config.get("KP_YARN_STOCK_ON"):
        print("🧶 ใช้ด้ายในสต๊อกแทนการรอ PO IN: ปิดอยู่ (หน้า Monitor Yarn)")
        return
    if not plan_config.get("KP_YARN_POIN_ON"):
        print("🧶 ใช้ด้ายในสต๊อก: ปิดกติกาวันด้ายเข้าอยู่ ไม่มีใบไหนต้องรอ PO IN — ไม่ต้องวางรอบสอง")
        return
    import YarnStockCover as ysc
    lag = int(plan_config.get("KP_YARN_POIN_LAG_WEEKS"))
    tl = _timeline()
    try:
        cover = ysc.plan_cover(RESULT_FILE, combined, tl, FREEZE_WEEKS_AHEAD, lag)
    except Exception as e:
        print(f"⚠️ [YARN STOCK] คิดสต๊อกด้ายไม่ได้ ({e}) → ใบรอ PO IN ตามเดิม")
        return
    if cover.get("note"):
        print(f"⚠️ [YARN STOCK] {cover['note']} → ใบรอ PO IN ตามเดิม")
        return
    if cover.get("warn"):
        print(f"⚠️ [YARN STOCK] {cover['warn']}")
    audit = cover["audit"]
    if not audit.empty:
        try:
            with pd.ExcelWriter(DEMAND_FILE, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
                audit.to_excel(w, sheet_name="YARN_STOCK", index=False)
        except Exception as e:
            print(f"⚠️ [YARN STOCK] เขียนชีท YARN_STOCK ไม่ได้ ({e})")
    picks = cover["picks"]
    print(f"🧶 ใบที่รอ PO IN {len(audit)} ใบ → ยืมด้ายในสต๊อกได้ {len(picks)} ใบ "
          f"(สต๊อก ณ {cover.get('stock_at') or '-'} ; รายละเอียดชีท YARN_STOCK ใน {DEMAND_FILE.name})")
    for (_so_no, _item), _note in list(picks.items())[:40]:
        print(f"   🧶 SC {_so_no} {_item}: {_note}")
    if not picks:
        return

    # รอบสอง: ใบที่ยืมได้ไม่ต้องรอ PO IN — สำรองผลรอบแรกไว้คืนถ้ารอบสองล้ม
    backup = RESULT_FILE.with_name(RESULT_FILE.stem + "_pass1" + RESULT_FILE.suffix)
    shutil.copy2(RESULT_FILE, backup)
    ysc.mark(combined, picks).to_excel(ORDER_REPLAN_FILE, index=False)
    print("=" * 70)
    print(f"▶ วางแผนรอบสอง — {len(picks)} ใบใช้ด้ายในสต๊อก ไม่ต้องรอ PO IN")
    if not run_planning_replan():
        shutil.copy2(backup, RESULT_FILE)
        combined.to_excel(ORDER_REPLAN_FILE, index=False)
        print("⚠️ [YARN STOCK] Planning รอบสองล้มเหลว → ใช้แผนรอบแรก (ใบรอ PO IN ตามเดิม)")
        return
    try:
        warn = ysc.verify(RESULT_FILE, cover, tl, lag)
    except Exception as e:
        print(f"⚠️ [YARN STOCK] ตรวจบัญชีด้ายหลังวางรอบสองไม่ได้ ({e})")
        return
    if warn:
        print(f"⚠️ [YARN STOCK] หลังวางรอบสอง ด้ายที่ยืมทำให้คงเหลือติดลบ {len(warn)} ตัว "
              f"(แผนรอบสองขยับงานอื่นมาใช้ด้ายเร็วขึ้น): " + " · ".join(warn[:10]))
    else:
        print("✅ [YARN STOCK] ตรวจบัญชีด้ายหลังวางรอบสอง: ไม่มีด้ายตัวไหนติดลบเพราะการยืม")


# =========================
# ตารางเทียบ booking เดิม vs แผนใหม่
# =========================
def build_compare(unlocked: pd.DataFrame, demand_keys: set) -> pd.DataFrame:
    """long format 1 แถว = งานหนึ่งก้อนในสัปดาห์หนึ่ง — ROW_KIND บอกว่าเป็นของเดิมหรือของใหม่"""
    old = pd.DataFrame()
    if not unlocked.empty:
        old = pd.DataFrame({
            "ROW_KIND": "OLD",
            "ORIGIN": "BOOKING",
            "ITEM_CODE": _str_col(unlocked, "ITEM_CODE"),
            "SC_SO_NO": _str_col(unlocked, "SO_NO"),
            "MC_GROUP": _str_col(unlocked, "MC_GROUP"),
            "MC_GUAGE": _str_col(unlocked, "GUAGE"),
            "WEEK": pd.to_numeric(_s(unlocked, "WEEK"), errors="coerce"),
            "QTY": pd.to_numeric(_s(unlocked, "KP_WEIGHT"), errors="coerce"),
            "MC": pd.to_numeric(_s(unlocked, "MC_USE_CEIL"), errors="coerce"),
        })

    new = pd.DataFrame()
    if RESULT_FILE.exists():
        try:
            plan = pd.read_excel(RESULT_FILE, sheet_name="PLAN")
        except Exception as e:
            print(f"⚠️ อ่านชีท PLAN ของ replan ไม่ได้: {e}")
            plan = pd.DataFrame()
        if not plan.empty:
            _to_yw(plan, "PLAN_WEEK", "PLAN_YEAR")  # ให้เทียบ/เรียงกับฝั่ง OLD (YW) บนแกนเดียวกัน
            item = _str_col(plan, "ITEM_CODE")
            sc = _str_col(plan, "SC_SO_NO")
            origin = [
                "BOOKING" if f"{s}|{i}".upper() in demand_keys else "SC"
                for s, i in zip(sc, item)
            ]
            new = pd.DataFrame({
                "ROW_KIND": "NEW",
                "ORIGIN": origin,
                "ITEM_CODE": item,
                "SC_SO_NO": sc,
                "MC_GROUP": _str_col(plan, "MC_GROUP"),
                "MC_GUAGE": _str_col(plan, "MC_GUAGE"),
                "WEEK": pd.to_numeric(_s(plan, "PLAN_WEEK"), errors="coerce"),
                "QTY": pd.to_numeric(_s(plan, "PRODUCE_QTY"), errors="coerce"),
                "MC": pd.to_numeric(_s(plan, "ACTUAL_MC"), errors="coerce"),
            })

    out = pd.concat([old, new], ignore_index=True) if (not old.empty or not new.empty) else pd.DataFrame()
    if not out.empty:
        out = out.sort_values(["ITEM_CODE", "SC_SO_NO", "ROW_KIND", "WEEK"]).reset_index(drop=True)
        # ชีทนี้ถูกอ่านต่อ (replan_notify / หน้าเว็บ) → เขียนเป็นเลขสัปดาห์ + คอลัมน์ YEAR
        out = _split_yw(out, [("WEEK", "YEAR")])
    return out


def main() -> int:
    start = datetime.now()
    if not ORDER_FILE.exists():
        print(f"❌ ไม่พบ {ORDER_FILE} — ต้องรัน Order.py ก่อน")
        return 1
    booking_file = _latest_booking_final()
    if not booking_file.exists():
        print(f"❌ ไม่พบไฟล์ booking_final — ต้องรัน AVA_MC.py ก่อน")
        return 1

    # ยุบสัปดาห์แล้วต้องระบุ "วัน + ชั่วโมง" ของสัปดาห์นั้น ไม่งั้นกำลังผลิตของก้อนที่รวมกัน
    # จะมาจากการรวมอัตโนมัติ/ชั่วโมงมาตรฐาน ซึ่งแยกไม่ออกจากค่าที่ผู้วางแผนตั้งใจ
    _gaps = merge_gaps()
    if _gaps and os.environ.get("KP_ALLOW_MERGE_HOUR_GAP", "") != "1":
        print("❌ มีสัปดาห์ที่ยุบรวมกันแล้วยังกรอกวัน/ชั่วโมงไม่ครบ — หยุดก่อนวางแผน")
        for _w, _cap, _miss in _gaps:
            _captxt = f" (เพดาน {_cap:g} วัน)" if _cap else ""
            print(f"   {yw_label(_w)}{_captxt}: ขาด {len(_miss)} Cat")
            for _k, _bad in list(_miss.items())[:8]:
                print(f"      · {_k} → ขาด{', '.join(_bad)}")
            if len(_miss) > 8:
                print(f"      · ... และอีก {len(_miss) - 8} Cat")
        print("   วัน: ตั้งค่าระบบ › วันทำงาน (ตารางรายสัปดาห์) · "
              "ชั่วโมง: ข้อมูลหลัก › Master MC → กดชื่อ Cat")
        print("   (ถ้าตั้งใจจะรันทั้งที่ยังไม่ครบ ให้ตั้ง KP_ALLOW_MERGE_HOUR_GAP=1)")
        return 1
    if _gaps:
        print(f"⚠️ ยุบสัปดาห์แล้วยังกรอกไม่ครบ {sum(len(m) for _, _, m in _gaps)} Cat "
              "— ข้ามด่านตาม KP_ALLOW_MERGE_HOUR_GAP=1 (Cat เหล่านี้ใช้ค่ารวมอัตโนมัติ)")

    tl = _timeline()
    today_idx = tl.today_idx
    unlock_idx = today_idx + FREEZE_WEEKS_AHEAD
    print(f"📅 สัปดาห์ปัจจุบัน {yw_label(tl.at(today_idx))} (idx {today_idx}) → "
          f"ปลดล็อกตั้งแต่ {yw_label(tl.at(unlock_idx))} (idx {unlock_idx})")

    detail = pd.read_excel(booking_file, sheet_name="DETAIL")
    detail.columns = detail.columns.str.strip().str.upper()
    _to_yw(detail)  # WEEK + YEAR → YW (คีย์สัปดาห์ของทั้งไฟล์นี้)
    print(f"📄 booking DETAIL: {len(detail)} แถว จาก {booking_file.name}")

    # ตำแหน่งบนปฏิทินของสัปดาห์ (YW ตรงตัว ; W99 ไม่อยู่บนปฏิทิน = None = ล็อก)
    w_idx = pd.to_numeric(detail["WEEK"], errors="coerce").map(lambda w: tl.idx(int(w)) if pd.notna(w) else None)
    is_unlocked = w_idx.map(lambda i: i is not None and i >= unlock_idx)

    # ล็อกที่ผู้ใช้กดเอง (ชีท Lock_Plan) ชนะเส้นเวลา — งานที่ล็อกอยู่ฝั่ง "ล็อก" เสมอ
    locks = load_plan_locks()
    user_locked = user_locked_mask(detail, locks)
    n_user_lock = int((is_unlocked & user_locked).sum())
    if locks:
        # คีย์ทั้งหมดที่มีจริงใน booking → ล็อกที่ไม่อยู่ในนี้คือจับคู่ไม่ได้ (ต้องเตือน user)
        _detail_keys = set(zip(
            pd.to_numeric(detail["WEEK"], errors="coerce").fillna(-1).astype(int),
            _str_col(detail, "ITEM_CODE").str.upper(),
            _str_col(detail, "MC_GROUP").str.upper(),
            _s(detail, "GUAGE").map(_norm_gauge),
        ))
        _miss = [k for k in locks if k not in _detail_keys]
        print(f"🔐 ล็อกจากชีท {LOCK_PLAN_SHEET}: {len(locks)} รายการ "
              f"→ กันงานไว้ {n_user_lock} แถว (ไม่ยกไปวางใหม่)")
        if _miss:
            print(f"⚠️ ล็อกที่จับคู่กับ booking ไม่ได้ {len(_miss)} รายการ (booking ย้ายสัปดาห์/เครื่องไปแล้ว?) "
                  f"→ {', '.join(f'{yw_label(k[0])} {k[1]} {k[2]} G{k[3]}' for k in _miss[:10])}"
                  + (" ..." if len(_miss) > 10 else ""))
    is_unlocked = is_unlocked & ~user_locked

    # ใบที่ออกเลข KP แล้ว = ยืนยันกับฝ่ายผลิตไปแล้ว → ตรึงสัปดาห์/เครื่องเดิม ไม่เข้ารอบวางใหม่
    # ต้องอยู่ก่อนด่าน core: core ที่มีเลขใบจองแล้วก็ห้ามขยับเหมือนกัน (ผู้ใช้สั่ง 21 ก.ย. 2569)
    kp_locked = kp_locked_mask(detail) & is_unlocked
    n_kp_lock = int(kp_locked.sum())
    if n_kp_lock:
        _kp_weeks = sorted(pd.to_numeric(detail.loc[kp_locked, "WEEK"], errors="coerce")
                           .dropna().astype(int).unique())
        print(f"🎫 มีเลขใบจอง (KP_NO) แล้ว {n_kp_lock} แถว "
              f"(สัปดาห์ {[yw_label(w) for w in _kp_weeks]}) → คงแผนเดิม ไม่ยกมาวางใหม่")
    elif "KP_NO" not in detail.columns:
        print("⚠️ ไฟล์ booking ไม่มีคอลัมน์ KP_NO (สร้างก่อนอัปเดต View_Booking/AVA_MC) "
              "→ ข้ามด่าน 'ใบที่ออกเลขแล้วไม่ replan'")
    is_unlocked = is_unlocked & ~kp_locked

    # core item เข้ารอบวางใหม่ตามปกติ — ยกเว้นตัวที่ user เบรกไว้ด้วย Core Status
    # (ตัวที่เบรก = คงแผนเดิม + ไม่คืนเครื่อง เหมือนงานที่ user กดล็อกเอง ; ดู section CORE ITEM)
    core_items = load_core_hold_items()
    is_core = core_item_mask(detail, core_items) & is_unlocked
    n_core = int(is_core.sum())
    if n_core:
        _core_weeks = sorted(pd.to_numeric(detail.loc[is_core, "WEEK"], errors="coerce")
                             .dropna().astype(int).unique())
        _core_items_hit = _str_col(detail[is_core], "ITEM_CODE").str.upper().nunique()
        print(f"📘 เบรก core item ไว้ {n_core} แถว ({_core_items_hit} item, "
              f"สัปดาห์ {[yw_label(w) for w in _core_weeks]}) → คงแผนเดิม ไม่ยกมาวางใหม่")
    is_unlocked = is_unlocked & ~is_core
    core_detail = detail[is_core].copy()
    kp_detail = detail[kp_locked].copy()    # ใช้เป็นตัวสำรองเมื่อไฟล์ booking ไม่มีชีท DETAIL_SC

    locked_detail = detail[~is_unlocked].copy()
    # ธงบอก plan_view ว่าแถวนี้ไปโผล่เป็น "แถวแผน" แล้ว → อย่า overlay เป็นบล็อกเก่าซ้อนอีกใบ
    # (ยังต้องอยู่ใน DETAIL ตามเดิม เพราะเป็นแหล่งของโควตา job / เครื่องที่ booking จองไว้)
    # ติดให้ทั้ง core ที่เบรก และใบที่ออกเลข KP — ทั้งคู่ถูกแปะกลับเป็นแถวแผน (write_result_append)
    locked_detail[CORE_PIN_FLAG] = (
        (is_core | kp_locked).reindex(locked_detail.index).fillna(False).astype(int)
    )
    unlocked = detail[is_unlocked].copy()
    unlocked_weeks = set(pd.to_numeric(unlocked["WEEK"], errors="coerce").dropna().astype(int))
    # คู่ (ปี, สัปดาห์) ของงานที่ปลดล็อก — ใช้คืนเครื่องให้ตรงปี (ดู write_replan_capacity)
    unlocked_yw = set()
    if "YEAR" in unlocked.columns:
        _uy = pd.to_numeric(unlocked["YEAR"], errors="coerce")
        _uw = pd.to_numeric(unlocked["WEEK"], errors="coerce")
        unlocked_yw = {(int(y), int(w)) for y, w in zip(_uy, _uw)
                       if pd.notna(y) and pd.notna(w)}
    print(f"🔒 ล็อก {len(locked_detail)} แถว (ผู้ใช้ล็อกเอง {n_user_lock}, มีเลขใบจอง {n_kp_lock}, "
          f"core ที่เบรกไว้ {n_core}) | "
          f"🔓 ปลดล็อกมาวางใหม่ {len(unlocked)} แถว (สัปดาห์ {[yw_label(w) for w in sorted(unlocked_weeks)]})")

    # ── ระดับ SC: ชุดงานเดียวกันแต่ไม่ยุบใบ ────────────────────────────────────
    # ดีมานด์และแถว core ต้องแยกต่อใบ ไม่งั้นทั้งก้อนถูกวางตาม deadline ของใบที่เร็วสุด
    # (booking มี 729 บล็อกที่ FG week ของใบข้างในต่างกันจริง)
    # เครื่อง/โควตา (locked_detail → write_replan_capacity) ยังคิดจาก DETAIL ระดับบล็อกเหมือนเดิม
    detail_sc = load_detail_sc(booking_file, detail)
    if detail_sc is None:
        print("⚠️ ไฟล์ booking ไม่มีชีท DETAIL_SC (สร้างก่อนอัปเดต AVA_MC) "
              "→ วางแผนจากยอดที่ยุบใบแล้วเหมือนเดิม")
        unlocked_sc, core_detail_sc, kp_detail_sc = unlocked, core_detail, kp_detail
    else:
        _w_sc = pd.to_numeric(detail_sc["WEEK"], errors="coerce").map(
            lambda w: tl.idx(int(w)) if pd.notna(w) else None)
        _unlocked_sc = _w_sc.map(lambda i: i is not None and i >= unlock_idx)
        _unlocked_sc = _unlocked_sc & ~user_locked_mask(detail_sc, locks)
        # ด่าน KP_NO ต้องมีที่นี่ด้วย ไม่งั้นใบที่ตรึงไว้ระดับบล็อก (DETAIL) จะยังถูกยกมา
        # เป็นดีมานด์จากระดับใบ (DETAIL_SC) = วางซ้ำทั้งที่เครื่องเดิมยังถูกจองอยู่
        _kp_sc = kp_locked_mask(detail_sc) & _unlocked_sc
        _unlocked_sc = _unlocked_sc & ~_kp_sc
        _core_sc = core_item_mask(detail_sc, core_items) & _unlocked_sc
        _unlocked_sc = _unlocked_sc & ~_core_sc
        unlocked_sc = detail_sc[_unlocked_sc].copy()
        core_detail_sc = detail_sc[_core_sc].copy()
        kp_detail_sc = detail_sc[_kp_sc].copy()
        print(f"🧾 DETAIL_SC: ปลดล็อก {len(unlocked_sc)} ใบ (จาก {len(unlocked)} บล็อก) | "
              f"มีเลขใบจอง {int(_kp_sc.sum())} ใบ | "
              f"core ที่เบรกไว้ {len(core_detail_sc)} ใบ (จาก {len(core_detail)} บล็อก)")

    def _yw_of_week(w: int):
        """สัปดาห์ → YW ตามปฏิทินจริง (หาปีของแถว core ที่ WEEK เป็นเลขเปล่า)
        WEEK ในไฟล์นี้เป็น YW อยู่แล้ว → คืนตัวเดิม ; เลขเปล่า = ครั้งที่ใกล้วันนี้สุด"""
        k = tl.key(int(w))
        return k if k is not None and k >= 100000 else None

    demand = build_booking_demand(unlocked_sc)
    demand_no_fg = demand.attrs.get("no_fg", pd.DataFrame())
    print(f"📦 ยอด booking ที่ยกมาเป็นดีมานด์: {len(demand)} แถว "
          f"รวม {float(demand['PENDING_PLAN'].sum()) if not demand.empty else 0:,.0f} กก.")

    # งานในแผนที่ user กดล็อกไว้ (SOURCE=PLAN) — ตรึงไว้ที่เดิม ไม่ยกเข้ารอบวางใหม่
    pinned = load_pinned_plan_rows(load_plan_row_locks())

    orders = pd.read_excel(ORDER_FILE)
    orders.columns = orders.columns.str.strip()
    # SO_NO ว่างในใบ SC ที่ยังไม่แปลงเป็น SO → pandas ทำทั้งคอลัมน์เป็น float64
    # แล้ว ".0"/"nan" ติดไปกับ order_replan.xlsx → เข้าชีทแผนต่อ (ดู ColorOpen.so_text)
    for _c_so in ("SO_NO", "SC_NO"):
        if _c_so in orders.columns:
            orders[_c_so] = orders[_c_so].map(ColorOpen.so_text)
    demand = _align_to_order_schema(demand, orders)
    # วันย้อมด้าย: ต้องเติมก่อน dedup — ถ้า booking ชนะ ใบ SC ที่ถือวันย้อมอยู่จะถูกตัดทิ้ง
    demand = attach_dye_end_date(demand, orders)
    # เลขบรรทัดใบ: booking ไม่มีฟิลด์นี้ → ยกจากฝั่ง SC ก่อน dedup ด้วยเหตุผลเดียวกับวันย้อม
    demand = attach_sc_line_id(demand, orders)
    orders_kept, demand_kept, audit = dedup_against_sc(demand, orders)
    if not audit.empty:
        print(f"🧹 งานซ้ำ SC↔booking {len(audit)} คู่ → ใช้ยอดจาก {'booking' if BOOKING_WINS_ON_DUPLICATE else 'SC'} "
              f"(ตัดอีกฝั่งออก) ; ผลต่างรวม {audit['QTY_DIFF'].sum():,.0f} กก.")

    with pd.ExcelWriter(DEMAND_FILE, engine="openpyxl") as w:
        (demand_kept if not demand_kept.empty else pd.DataFrame(columns=orders.columns)).to_excel(
            w, sheet_name="BOOKING_DEMAND", index=False)
        (audit if not audit.empty else pd.DataFrame(columns=["KEY"])).to_excel(
            w, sheet_name="DEDUP_AUDIT", index=False)
        # งาน booking ที่ต้นทางไม่ส่ง FG week มา → ไม่ได้วาง (ห้ามเติมกำหนดส่งเอง)
        (demand_no_fg if not demand_no_fg.empty else pd.DataFrame(columns=["SC_NO"])).to_excel(
            w, sheet_name="NO_FG", index=False)
    print(f"💾 ดีมานด์จาก booking → {DEMAND_FILE.name}")

    # งานที่ตรึงไว้ + core ที่ user เบรกไว้ + ใบที่ออกเลข KP ต้องกลับเข้าชีทแผน — ไม่ถูกจัดใหม่
    # Planning รอบ replan เป็นคนแปะ (อ่าน replan_append.xlsx) → แถวพวกนี้อยู่ในแผนก่อนทำตารางเทียบ
    # (core ที่ไม่ได้เบรก ไม่ผ่านทางนี้ — Planning วางให้เป็นแถวแผนปกติ)
    # core แตกเป็นแถวละใบ (ยอด/กำหนดส่งของใบนั้นจริง) — เครื่องยังเป็นของทั้งบล็อก
    # งาน KP เดิมไม่ถูกแปะกลับเลย → หายทั้งจาก Gantt และชีท PLAN ทั้งที่เครื่องยังถูกหักไว้
    # (แก้ 23 ก.ย. 2569) ; ต่างจาก core ตรงที่หน้าเว็บล็อกไม่ให้ลาก/แก้
    # ต้องสร้างก่อนหักดีมานด์ — ยอดของทุกแถวที่แปะกลับต้องถูกหักออก ไม่งั้นวางซ้ำ
    _yr_of = lambda w: (_yw_of_week(w) or 0) // 100 or None
    core_rows = core_plan_rows(core_detail_sc, None, _yr_of)
    kp_rows = core_plan_rows(kp_detail_sc, None, _yr_of,
                             source=KP_PLAN_SOURCE, remark=KP_REMARK, is_core=False)
    # งานที่ user กดล็อกไว้แต่รอบนี้ booking สร้างแถวให้อยู่แล้ว → ใช้ของ booking ทางเดียว
    # (ต้องทำก่อนทุกที่ที่ใช้ pinned: หักดีมานด์ / กันเครื่อง / โควตา job / แปะกลับ)
    pinned = drop_pinned_rebuilt(pinned, [kp_rows, core_rows])

    combined = pd.concat([orders_kept, demand_kept], ignore_index=True) if not demand_kept.empty else orders_kept
    # หักงานที่ตรึงไว้ตอนนี้ (หลัง dedup) — ต้องรู้ก่อนว่าฝั่งไหนรอดมาเป็นดีมานด์จริง
    # ทั้ง 3 ชนิดใช้ด่านเดียวกัน: แปะกลับเป็นแถวแผน = ต้องไม่เหลือยอดให้ Planning วางอีก
    # ⚠️ แถวจาก core_plan_rows เก็บ PLAN_WEEK เป็นเลขสัปดาห์เปล่า + PLAN_YEAR (รูปไฟล์แผน)
    # แต่ deduct เทียบสัปดาห์กับ BOOKING_WEEK_FROM/TO ของดีมานด์ซึ่งเป็น YW → ส่งสำเนาที่แปลงแล้ว
    # (ไม่งั้นเทียบไม่ติด แล้วไปหักยอดของก้อนสัปดาห์อื่นแทนก้อนจริง)
    combined = deduct_pinned_from_demand(combined, pinned)
    # KP/core ถูกกันออกตั้งแต่ระดับแถว booking (ไม่อยู่ใน unlocked_sc) → หักเฉพาะดีมานด์ฝั่ง SC
    # ส่วนงานที่ user ล็อกบนแผน (pinned) แถว booking ของสัปดาห์นั้นยังปลดล็อกอยู่ → หักได้ทั้งสองฝั่ง
    combined = deduct_pinned_from_demand(combined, _keep_rows_yw(kp_rows), label="ใบที่ออกเลข KP",
                                         sc_side_only=True)
    combined = deduct_pinned_from_demand(combined, _keep_rows_yw(core_rows), label="core ที่เบรกไว้",
                                         sc_side_only=True)
    # แผนวางได้ไม่เกินด้ายที่มี (สต๊อก + PO Pending) — หลังหักงานที่ตรึงแล้ว = ดีมานด์ที่ Planning จะวางจริง
    combined = yarn_supply_cap(combined)
    combined.to_excel(ORDER_REPLAN_FILE, index=False)
    print(f"💾 ดีมานด์รวม {len(combined)} แถว "
          f"(ก่อนหักงานที่ตรึง: SC {len(orders_kept)} + booking {len(demand_kept)}) "
          f"→ {ORDER_REPLAN_FILE.name}")

    write_replan_capacity(booking_file, locked_detail, unlocked_weeks, unlocked_yw,
                          pinned_mc=pinned_mc_use(pinned, booking_file),
                          pinned_jobs=pinned_setup_jobs(pinned),
                          locked_sc=locked_detail_sc(detail_sc, locked_detail),
                          pinned_item_jobs=pinned_setup_jobs_by_item(pinned))

    write_result_append(pinned_result_rows(pinned), core_rows, kp_rows)

    if not run_planning_replan():
        print("❌ Planning รอบ replan ล้มเหลว — ไม่มีแผนผลิตรอบนี้")
        return 1
    # ใบที่รอ PO IN แต่สต๊อกมีด้ายพอ → วางรอบสองโดยไม่ต้องรอ (หน้า Monitor Yarn)
    yarn_stock_second_pass(combined)

    demand_keys = set()
    if not demand_kept.empty:
        demand_keys = {
            f"{s}|{i}".upper()
            for s, i in zip(
                demand_kept["SO_NO"].astype(str).str.strip(),
                demand_kept["FABRIC_ITEM"].astype(str).str.strip(),
            )
        }
    # เทียบระดับใบ — แถว NEW มี SC_SO_NO เป็นใบเดี่ยว ถ้าฝั่ง OLD ยังยุบหลายใบไว้จะจับคู่ไม่ตรง
    compare = build_compare(unlocked_sc, demand_keys)
    if not compare.empty:
        with pd.ExcelWriter(RESULT_FILE, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
            compare.to_excel(w, sheet_name=COMPARE_SHEET, index=False)
        print(f"📊 ตารางเทียบ booking เดิม vs แผนใหม่: {len(compare)} แถว → {RESULT_FILE.name} ชีท {COMPARE_SHEET}")

    # SC_STATUS ท้ายชีท PLAN: คีย์ของยอด booking ที่ยกมาวางรอบนี้ + ของ booking ทั้งไฟล์ (ใช้กับแถวที่ตรึง)
    _bk = detail_sc if detail_sc is not None else detail
    if publish_plan_file(
        demand_keys=(_sc_item_keys(demand_kept["SO_NO"], demand_kept["FABRIC_ITEM"])
                     if not demand_kept.empty else set()),
        booking_keys=_sc_item_keys(_s(_bk, "SO_NO"), _s(_bk, "ITEM_CODE")),
    ) is None:
        return 1

    print(f"[DONE] Replan ใช้เวลา {datetime.now() - start}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
