"""
mc_hub.py — "เครื่องจักร" ที่เดียวจบ: จำนวนเครื่อง · วันทำงาน · ชั่วโมง · ค่าเฉพาะสัปดาห์ ·
             กันเครื่อง · เครื่องที่กันให้ POLY/COTTON

ทำไมต้องรวม: ข้อมูลของ "เครื่องหนึ่งกอง" เดิมกระจายอยู่ 5 ที่ คนละไฟล์ คนละคีย์
  MasterMC/Master MC   จำนวนเครื่อง            คีย์ Factory + MC_CAT + MC + Guage
  MasterMC/MC_Total    จำนวนเฉพาะสัปดาห์        คีย์ Week + Factory + MC_CAT + Guage (+MC)
  MasterMC/Lock_MC     กันเครื่อง               คีย์ Week + MC_CAT + Guage (+MC)
  MasterMC/MC Special  กันให้ POLY/COTTON       คีย์ Factory + MC_CAT + MC + Guage
  Calendar/Work Day    วันทำงาน + ชั่วโมง        คีย์ Factory + MC_CAT + Guage   ← ไม่มีชื่อเครื่อง
คีย์ไม่ตรงกันคือสาเหตุที่แก้ที่หนึ่งแล้วอีกที่หลุด (เคสจริง: เปลี่ยน SINGEL-16 → SINGLE-16
ในทะเบียน แต่แผงวันทำงานยังเป็นชื่อเดิม → ใช้ 6 วันแทน 5.5 โดยไม่มีใครรู้)

โครงที่ใช้: จัดเป็น "กอง" (Factory · MC_CAT · Guage) = หน่วยเดียวกับคีย์วันทำงาน
  • วันทำงาน / ชั่วโมง / กันเครื่อง / ค่าเฉพาะสัปดาห์ → ของทั้งกอง
  • จำนวนเครื่อง / POLY-COTTON                      → ของเครื่องแต่ละตัวในกอง
เขียนกลับลงไฟล์เดิมทั้งหมด (pipeline อ่านที่เดิม ไม่ต้องแก้อะไรฝั่ง engine)
"""
import livesync
import masters
import workday

MC_FILE = "MasterMC"
SH_MASTER = "Master MC"
SH_TOTAL = "MC_Total"
SH_LOCK = "Lock_MC"
SH_SPECIAL = "MC Special"


def _s(v) -> str:
    return str(v if v is not None else "").strip()


def _key(v) -> str:
    s = _s(v).upper().replace("\xa0", "")
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def _num(v):
    s = _s(v)
    if not s or s in ("-", "nan", "NaN"):
        return None
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return None


def _cols(columns, *names):
    """ชื่อคอลัมน์ → index (เทียบแบบไม่สนตัวพิมพ์/ช่องว่าง) ; ไม่เจอ = -1"""
    up = [_key(c) for c in columns]
    out = []
    for n in names:
        try:
            out.append(up.index(_key(n)))
        except ValueError:
            out.append(-1)
    return out


def _get(row, i):
    return row[i] if 0 <= i < len(row) else None


def pool_key(factory, mc_cat, guage) -> str:
    return f"{_key(factory)}|{_key(mc_cat)}|{_key(guage)}"


def _future_merge_targets(merges: dict, tl) -> list:
    """สัปดาห์ปลายทางของการยุบ (YW) ตั้งแต่สัปดาห์ปัจจุบัน — ของปีที่ผ่านไปแล้วไม่ต้องตรวจ
    (อยู่ในไฟล์ตลอดเพราะผูกปี ถ้าตรวจด้วย กองที่เพิ่มทีหลังจะทำให้บันทึกไม่ได้อีกเลย)"""
    targets = sorted({int(d) for d in (merges or {}).values()})
    if tl is None:
        return targets
    return [w for w in targets if tl.pos.get(w, len(tl.keys)) >= tl.today_idx]


def _day_cap(yw: int, merges: dict, cal_yw: dict):
    """เพดานวันของสัปดาห์ = วันที่ปฏิทินเปิดรวมสัปดาห์ที่ยุบเข้ามา (ของปีนั้น) ; ไม่มีข้อมูล = None"""
    bucket = [yw] + [int(s) for s, d in (merges or {}).items() if int(d) == yw]
    vals = [cal_yw.get(b) for b in bucket]
    if all(v is None for v in vals):
        return None
    return float(sum(v or 0 for v in vals))


def _pool_info() -> dict:
    """กองทั้งหมดจาก Master MC → {key: {key, factory, mc_cat, guage, total}} (total = ผลรวม Total MC)"""
    m = masters.read_sheet(MC_FILE, SH_MASTER)
    fi, ci, gi, toi = _cols(m["columns"], "Factory", "MC_CAT", "Guage", "Total MC")
    out: dict = {}
    for r in m["rows"]:
        k = pool_key(_get(r, fi), _get(r, ci), _get(r, gi))
        if k == "||":
            continue
        p = out.setdefault(k, {"key": k, "factory": _s(_get(r, fi)), "mc_cat": _s(_get(r, ci)),
                               "guage": _s(_get(r, gi)), "total": 0})
        p["total"] += _num(_get(r, toi)) or 0
    return out


def _keys_with_machines() -> list:
    """กองที่มีเครื่องอยู่จริง (Total MC รวมกันมากกว่า 0) — ใช้เป็นรายชื่อที่ต้องกรอกให้ครบ"""
    return sorted(k for k, p in _pool_info().items() if p["total"] > 0)


def _merge_gaps(week_hours: dict, weeks: dict, merges: dict, tl=None, cal_yw: dict = None) -> list:
    """สัปดาห์ยุบ (ตั้งแต่สัปดาห์นี้) ที่ยังกรอกไม่ครบ
    → [{"week": YW, "cap": วัน, "days": [key...], "hours": [key...], "over": [key...]}]

    ยุบสัปดาห์เข้าด้วยกันแล้ว "วันที่เดินได้จริง" กับ "ชั่วโมงต่อวัน" มักไม่เท่าสัปดาห์ปกติ ถ้าปล่อยว่าง
    ระบบจะรวมวันให้เอง/ใช้ชั่วโมงมาตรฐานของกองเงียบ ๆ (เคยทำให้ 12 กองคิดกำลังผลิตเกินจริง)
    → ทุกกองที่มีเครื่องต้องระบุทั้งวันและชั่วโมง ; วันต้องไม่เกินวันที่ปฏิทินเปิดรวมทั้งก้อน
    กองที่ไม่มีเครื่องเลยไม่ต้องกรอก เพราะวางแผนลงไม่ได้อยู่แล้ว
    ต้องตรงกับ workday.merge_day_gaps / WorkDay.merge_gaps / Replan.merge_gaps"""
    targets = _future_merge_targets(merges, tl)
    if not targets:
        return []
    if cal_yw is None:
        cal_yw = workday._cal_days_yw()
    keys = _keys_with_machines()
    week_hours, weeks = week_hours or {}, weeks or {}
    out = []
    for w in targets:
        cap = _day_cap(w, merges, cal_yw)
        no_day = [k for k in keys if weeks.get(f"{k}|{w}") is None]
        no_hour = [k for k in keys if week_hours.get(f"{k}|{w}") is None]
        over = [k for k in keys if cap is not None and weeks.get(f"{k}|{w}") is not None
                and float(weeks[f"{k}|{w}"]) > cap + 1e-9]
        if no_day or no_hour or over:
            out.append({"week": w, "label": workday._wlabel(w), "cap": cap,
                        "days": no_day, "hours": no_hour, "over": over})
    return out


def _gap_message(gaps: list) -> str:
    parts = []
    for g in gaps:
        bits = []
        if g["days"]:
            bits.append(f"ยังไม่ได้ระบุวัน {len(g['days'])} Cat")
        if g["hours"]:
            bits.append(f"ยังไม่ได้ระบุชั่วโมง {len(g['hours'])} Cat")
        if g["over"]:
            bits.append(f"วันเกินเพดาน {g['cap']:g} วัน {len(g['over'])} Cat")
        parts.append(f"{g['label']} (สัปดาห์ที่ยุบรวมมา) " + " · ".join(bits))
    return "บันทึกไม่ได้ — " + " · ".join(parts)


# =============================================================================
# จำนวนเครื่องรายสัปดาห์ (MC_Total)
# =============================================================================
# user สั่ง (25 ก.ย. 2569): จำนวนเครื่องมีผล "เฉพาะสัปดาห์ที่กรอก" — สัปดาห์ที่ไม่ได้กรอกใช้ Total MC
# ของ Master MC (เคยลองแบบยกค่าล่าสุดไปเรื่อย ๆ แล้ว user ไม่เอา: "ไม่ใช่เป็น 71 ยาว")
# = ความหมายเดียวกับที่ engine (AVA_MC / Planning) อ่านชีทนี้อยู่แล้ว — ไม่แตะ engine
# หน้านี้แปลงแถวเป็น {สัปดาห์: จำนวน} ให้ผู้ใช้แก้ทีละสัปดาห์/เป็นช่วง และตอนบันทึกรวมสัปดาห์ติดกันที่
# จำนวนเท่ากันเป็นแถวเดียว (Week..Week To + ปีเสมอ — เลขสัปดาห์เปล่าเปลี่ยนความหมายตามวันที่รัน)
# แถวที่ระบุคอลัมน์ MC (ยอดของพูลย่อย) ไม่ยุ่ง — แสดงให้เห็นอย่างเดียวเหมือนเดิม
def _timeline():
    """WeekTimeline ตัวเดียวกับที่ pipeline ใช้กางช่วง (Calendar.WeekTimeline) ; อ่านไม่ได้ = None"""
    try:
        import sys
        import config
        if str(config.REPO_DIR) not in sys.path:
            sys.path.insert(0, str(config.REPO_DIR))
        from Calendar import WeekTimeline, calendar_week_map, load_calendar
        weekly = calendar_week_map(load_calendar(workday._path("Calendar"), sheet_name="Sheet1"))
        tl = WeekTimeline(weekly)
        return tl if tl.keys else None
    except Exception:                                        # noqa: BLE001
        return None


def _total_cols(columns) -> dict:
    names = ("Week", "Factory", "MC_CAT", "Guage", "Total MC", "MC", "Remark",
             "Year", "Week To", "Year To")
    return dict(zip(names, _cols(columns, *names)))


def _is_step_row(r, c: dict) -> bool:
    """แถวที่หน้านี้ดูแล = ระบุ Factory และไม่ได้ระบุ MC (ยอดทั้งกอง)"""
    return bool(_s(_get(r, c["Factory"]))) and not _s(_get(r, c["MC"]))


def _row_weeks(r, c: dict, tl):
    """แถวเดียว → (list YW ที่แถวนี้ครอบ, ครอบคลุมครบไหม)
    ครบ = ทุกสัปดาห์ของแถวอยู่ในปฏิทิน → หน้านี้แทนแถวนี้ด้วยตารางรายสัปดาห์ได้
    ไม่ครบ = สัปดาห์อยู่นอกช่วงปฏิทิน (เช่นแถวของปีที่ผ่านไปแล้ว) → หน้านี้ไม่รู้จัก ห้ามแตะ"""
    weeks, _warn = tl.span(_get(r, c["Week"]), _get(r, c["Week To"]),
                           _get(r, c["Year"]), _get(r, c["Year To"]))
    known = [w for w in weeks if w in tl.pos]
    return known, bool(weeks) and len(known) == len(weeks)


def _week_values(rows: list, c: dict, pkey: str, tl) -> dict:
    """แถว MC_Total ของกองนี้ → {yw: (จำนวน, remark)} เฉพาะสัปดาห์ที่มีแถว
    (แถวหลังชนะแถวก่อน = กติกาเดียวกับ AVA_MC)"""
    out = {}
    for r in rows:
        if not _is_step_row(r, c):
            continue
        if pool_key(_get(r, c["Factory"]), _get(r, c["MC_CAT"]), _get(r, c["Guage"])) != pkey:
            continue
        tot = _num(_get(r, c["Total MC"]))
        if _num(_get(r, c["Week"])) is None or tot is None or tot < 0:
            continue
        for yw in _row_weeks(r, c, tl)[0]:
            out[yw] = (int(tot), _s(_get(r, c["Remark"])))
    return out


def _week_rows(vals: dict, c: dict, ncols: int, fac, cat, guage, tl) -> list:
    """{yw: (จำนวน, remark)} → แถว MC_Total — สัปดาห์ติดกัน (ตามปฏิทิน) ที่จำนวน "และหมายเหตุ"
    เหมือนกันรวมเป็นแถวเดียว

    ต้องเทียบหมายเหตุด้วย ไม่งั้นสัปดาห์ที่จำนวนบังเอิญเท่ากันแต่คนละเหตุผล (W41 "ถอดไปซ่อม"
    + W42 "รออะไหล่") ถูกยุบเป็นแถวเดียวแล้วเหลือหมายเหตุของสัปดาห์แรก = เหตุผลของสัปดาห์
    ที่เหลือหายไปเงียบ ๆ"""
    out, run = [], None                     # run = [ตำแหน่งแรก, ตำแหน่งท้าย, จำนวน, remark]
    def flush():
        if run is None:
            return
        a, b = tl.keys[run[0]], tl.keys[run[1]]
        row = [""] * ncols
        for name, v in (("Week", a % 100), ("Year", a // 100), ("Week To", b % 100),
                        ("Year To", b // 100), ("Factory", fac), ("MC_CAT", cat),
                        ("Guage", _num(guage) if _num(guage) is not None else guage),
                        ("Total MC", run[2]), ("Remark", run[3] or "")):
            if c.get(name, -1) >= 0:
                row[c[name]] = v
        out.append(row)
    for yw in sorted(vals, key=tl.pos.get):
        i, (v, rm) = tl.pos[yw], vals[yw]
        if run is not None and i == run[1] + 1 and v == run[2] and rm == run[3]:
            run[1] = i
            continue
        flush()
        run = [i, i, v, rm]
    flush()
    return out


def _save_weeks(edits: list, rev, user: str) -> int:
    """edits = [(pool, {yw: จำนวน | None})] (None = ไม่ตั้ง → ใช้ Total ของทะเบียน)
    → เขียนชีท MC_Total ครั้งเดียว ; คืนจำนวนกองที่แก้"""
    tl = _timeline()
    if tl is None:
        raise ValueError("อ่านปฏิทินไม่ได้ — บันทึกจำนวนเครื่องรายสัปดาห์ไม่ได้")
    t = masters.read_sheet(MC_FILE, SH_TOTAL)
    cols, rows = list(t["columns"]), [list(r) for r in t["rows"]]
    # ชีทเดิมอาจไม่มีคอลัมน์ช่วง/ปี → เติมท้าย (แถวเก่าเว้นว่าง = พฤติกรรมเดิม)
    for name in ("Year", "Week To", "Year To"):
        if _cols(cols, name)[0] < 0:
            cols.append(name)
            rows = [r + [""] for r in rows]
    c = _total_cols(cols)
    rows = [r + [""] * (len(cols) - len(r)) for r in rows]
    n = 0
    for p, ed in edits:
        vals = _week_values(rows, c, p["key"], tl)
        for yw, v in ed.items():
            if yw not in tl.pos:
                continue
            if v is None:
                vals.pop(yw, None)
            else:
                old = vals.get(yw)
                vals[yw] = (int(v), old[1] if old else f"แก้ที่หน้าเครื่องจักร ({user or '-'})")
        # แทนแถวเดิมของกองนี้ ณ ตำแหน่งแถวแรกที่เคยอยู่ (แถวกองอื่น/แถวระบุ MC ไม่แตะ)
        # ⚠️ ทิ้งได้เฉพาะแถวที่ตารางรายสัปดาห์ของหน้านี้ "อ่านเข้ามาแล้ว" เท่านั้น —
        # แถวที่สัปดาห์อยู่นอกช่วงปฏิทิน (เช่นของปีที่ผ่านไปแล้ว) ไม่เคยถูกอ่านเข้า vals
        # ถ้าทิ้งด้วยจะหายถาวรทันทีที่มีคนกดบันทึกกองนี้ โดยไม่มีใครรู้
        mine = [i for i, r in enumerate(rows)
                if _is_step_row(r, c)
                and pool_key(_get(r, c["Factory"]), _get(r, c["MC_CAT"]), _get(r, c["Guage"])) == p["key"]]
        at = mine[0] if mine else len(rows)
        drop = {i for i in mine if _row_weeks(rows[i], c, tl)[1]}
        keep = [r for i, r in enumerate(rows) if i not in drop]
        rows = keep[:at] + _week_rows(vals, c, len(cols), p["factory"], p["mc_cat"], p["guage"], tl) + keep[at:]
        n += 1
    if n:
        masters.write_sheet(MC_FILE, SH_TOTAL, cols, rows, user=user, base_rev=rev or t.get("rev"))
    return n


# =============================================================================
# อ่านรวม
# =============================================================================
def get() -> dict:
    """ทุกกองเครื่อง + ค่าที่เกี่ยวข้องทั้งหมด (พร้อมเลขรุ่นของแต่ละไฟล์ไว้กันเขียนทับ)"""
    m = masters.read_sheet(MC_FILE, SH_MASTER)
    fi, ti, ci, mi, gi, gri, sgi, toi = _cols(
        m["columns"], "Factory", "Type", "MC_CAT", "MC", "Guage", "Group", "SUBGroup", "Total MC")

    pools = {}
    for ri, r in enumerate(m["rows"]):
        cat, mc, g = _s(_get(r, ci)), _s(_get(r, mi)), _s(_get(r, gi))
        if not (cat or mc or g):
            continue
        fac = _s(_get(r, fi))
        k = pool_key(fac, cat, g)
        p = pools.setdefault(k, {
            "key": k, "factory": fac, "mc_cat": cat, "guage": g,
            "type": _s(_get(r, ti)), "machines": [],
            "work_days": None, "work_hours": None, "week_hours": {}, "week_days": {},
            "weeks": [], "locks": [], "reserved": [],
        })
        p["machines"].append({
            "row": ri, "mc": mc, "group": _s(_get(r, gri)),
            "total": _num(_get(r, toi)), "total_raw": _s(_get(r, toi)),
        })

    # ---- วันทำงาน + ชั่วโมง + ค่ารายสัปดาห์ (Calendar / Work Day)
    try:
        wd = workday.get_workday()
        wd_ok, wd_err = True, ""
    except Exception as e:                                   # noqa: BLE001
        wd = {"defaults": {}, "hours": {}, "weeks": {}, "week_hours": {}, "merges": {}}
        wd_ok, wd_err = False, str(e)
    for k, v in (wd.get("defaults") or {}).items():
        if k in pools:
            pools[k]["work_days"] = v
    for k, v in (wd.get("hours") or {}).items():
        if k in pools:
            pools[k]["work_hours"] = v
    # วัน/ชั่วโมง เฉพาะบางสัปดาห์ (คีย์ YW) — แก้ได้ที่หน้านี้ที่เดียว ; หน้าปฏิทินเหลือแค่การยุบสัปดาห์
    for k, v in (wd.get("weeks") or {}).items():
        base, _, w = k.rpartition("|")
        if base in pools:
            pools[base]["week_days"][str(_num(w) or w)] = v
    for k, v in (wd.get("week_hours") or {}).items():
        base, _, w = k.rpartition("|")
        if base in pools:
            pools[base]["week_hours"][str(_num(w) or w)] = v

    # ---- จำนวนเครื่องเฉพาะสัปดาห์ (MC_Total)
    t = masters.read_sheet(MC_FILE, SH_TOTAL)
    tw, tf, tc, tg, tt, tm, tr, ty, twt, tyt = _cols(
        t["columns"], "Week", "Factory", "MC_CAT", "Guage", "Total MC", "MC", "Remark",
        "Year", "Week To", "Year To")
    tl = _timeline()
    tcols = _total_cols(t["columns"])
    for ri, r in enumerate(t["rows"]):
        if _num(_get(r, tw)) is None:
            continue
        k = pool_key(_get(r, tf), _get(r, tc), _get(r, tg))
        if k not in pools:
            continue
        if tl is not None and _is_step_row(r, tcols):
            continue                        # แสดงเป็นตารางจุดเปลี่ยนแทน (คำนวณข้างล่าง)
        pools[k]["weeks"].append({
            "kind": "total", "row": ri, "week": _num(_get(r, tw)),
            "week_to": _num(_get(r, twt)), "year": _num(_get(r, ty)), "year_to": _num(_get(r, tyt)),
            "value": _num(_get(r, tt)), "mc": _s(_get(r, tm)), "remark": _s(_get(r, tr)),
        })

    # ---- กันเครื่อง (Lock_MC) — ชีทนี้ไม่มีคอลัมน์ Factory จึงจับเข้ากองด้วย CAT+เกจ
    lk = masters.read_sheet(MC_FILE, SH_LOCK)
    lw, lc, lm, lg, ll, li, lwt, ly, lyt = _cols(
        lk["columns"], "Week", "MC_CAT", "MC", "Guage", "MC Lock", "ITEM", "Week To", "Year", "Year To")
    for ri, r in enumerate(lk["rows"]):
        if _num(_get(r, lw)) is None:
            continue
        cat, g = _key(_get(r, lc)), _key(_get(r, lg))
        # ⚠️ ชีทนี้ไม่มีคอลัมน์ Factory และกติกาของ Planning คือ "ไม่ระบุ MC = หักทุกพูล"
        # (Planning.py:1307) → ถ้า CAT+เกจ เดียวกันมีหลายโรง รายการนี้กระทบทุกโรง
        # ต้องแสดงในทุกกองที่ตรง ไม่ใช่กองแรกกองเดียว ไม่งั้นบอกผู้ใช้ผิด
        hit = [p for p in pools.values()
               if _key(p["mc_cat"]) == cat and _key(p["guage"]) == g]
        mc_pick = _s(_get(r, lm))
        for p in hit:
            if mc_pick and not any(_key(x["mc"]) == _key(mc_pick) for x in p["machines"]):
                continue                        # ระบุเครื่องไว้ และเครื่องนั้นไม่ได้อยู่กองนี้
            p["locks"].append({
                "row": ri, "week": _num(_get(r, lw)), "week_to": _num(_get(r, lwt)),
                "year": _num(_get(r, ly)), "year_to": _num(_get(r, lyt)),
                "qty": _num(_get(r, ll)), "mc": mc_pick, "item": _s(_get(r, li)),
                # รายการเดียวกันโผล่หลายกอง = ของใช้ร่วม ต้องบอกให้รู้ว่าแก้แล้วกระทบที่อื่นด้วย
                "shared_with": [q["key"] for q in hit if q is not p] if not mc_pick else [],
            })

    # ---- เครื่องที่กันไว้ให้ POLY / COTTON (MC Special)
    sp = masters.read_sheet(MC_FILE, SH_SPECIAL)
    sf, sc, sm, sg, spo, sco, sd = _cols(
        sp["columns"], "Factory", "MC_CAT", "MC", "Guage", "POLY", "COTTON", "DESCRIPTION")
    for ri, r in enumerate(sp["rows"]):
        cat = _s(_get(r, sc))
        if not cat:
            continue
        k = pool_key(_get(r, sf), cat, _get(r, sg))
        if k not in pools:
            continue
        pools[k]["reserved"].append({
            "row": ri, "mc": _s(_get(r, sm)), "poly": _num(_get(r, spo)),
            "cotton": _num(_get(r, sco)), "note": _s(_get(r, sd)),
        })

    out = sorted(pools.values(), key=lambda p: (p["factory"].upper(), p["mc_cat"].upper(),
                                                str(p["guage"]).rjust(4)))
    for p in out:
        p["machines"].sort(key=lambda x: x["mc"].upper())
        p["total"] = sum(x["total"] or 0 for x in p["machines"])
        p["weeks"].sort(key=lambda x: (x["week"] or 0, x["kind"]))
        # จำนวนเครื่องเฉพาะสัปดาห์ที่กรอกไว้ [{yw, value, remark}] — สัปดาห์ที่ไม่มี = Total ของทะเบียน
        wv = _week_values(t["rows"], tcols, p["key"], tl) if tl is not None else {}
        p["mc_weeks"] = [{"yw": yw, "value": v, "remark": rm}
                         for yw, (v, rm) in sorted(wv.items(), key=lambda kv: tl.pos[kv[0]])]

    merges = {str(s): int(d) for s, d in (wd.get("merges") or {}).items()}
    cal_yw = {int(k): v for k, v in (wd.get("cal_days_yw") or {}).items()}
    return {
        "pools": out,
        "workday_ok": wd_ok, "workday_error": wd_err,
        # ยุบสัปดาห์ (ตั้งที่หน้าปฏิทิน) — หน้านี้ใช้บอกว่าสัปดาห์ไหนรับรวมมา
        # และบังคับให้ระบุวัน + ชั่วโมงของสัปดาห์นั้นก่อนบันทึก
        "merges": merges,
        "merge_targets": sorted({d for d in merges.values()}),
        "merge_gaps": _merge_gaps(wd.get("week_hours"), wd.get("weeks"), merges, tl, cal_yw),
        # วันที่ปฏิทินเปิดราย YW — เพดานวันของแต่ละสัปดาห์ (รวมสัปดาห์ที่ยุบเข้ามา)
        "cal_days_yw": {str(k): v for k, v in cal_yw.items()},
        # ไทม์ไลน์สัปดาห์ (YW = ปี*100+สัปดาห์) ของตารางจำนวนเครื่องรายสัปดาห์ ; ว่าง = อ่านปฏิทินไม่ได้
        "timeline": list(tl.keys) if tl is not None else [],
        "today_yw": tl.at(tl.today_idx) if tl is not None else None,
        "revs": {
            "master": m.get("rev"), "total": t.get("rev"),
            "lock": lk.get("rev"), "special": sp.get("rev"),
            "workday": livesync.rev(livesync.WORKDAY),
        },
        "options": {
            "factories": sorted({p["factory"] for p in out if p["factory"]}),
            "cats": sorted({p["mc_cat"] for p in out if p["mc_cat"]}),
            "types": sorted({p["type"] for p in out if p["type"]}),
            "mcs": sorted({x["mc"] for p in out for x in p["machines"] if x["mc"]}),
            "gauges": sorted({str(p["guage"]) for p in out if p["guage"] != ""}),
        },
    }


# =============================================================================
# เขียนกลับ
# =============================================================================
def save(changes: dict, revs: dict = None, user: str = "") -> dict:
    """บันทึกสิ่งที่ผู้ใช้แก้ในหน้าเดียว → กระจายลงไฟล์เดิมให้ครบทุกที่

    changes = {
      "machines": [{"row": i, "total": 12}],                       # จำนวนเครื่อง (Master MC)
      "pools":    [{"key": "PHET|DOUBLE-30|22", "work_days": 5.5,  # วันทำงาน/ชั่วโมงของกอง
                    "work_hours": 24,
                    "week_days": {"202632": 12},                   # วันเฉพาะสัปดาห์ (None = ล้าง)
                    "week_hours": {"202632": 24}}],                # ชั่วโมงเฉพาะสัปดาห์ (None = ล้าง)
    }
    สัปดาห์เป็น YW (ปี*100+สัปดาห์) — ค่าของปีหนึ่งต้องไม่ไหลไปใช้กับปีอื่น
    แก้เฉพาะช่องที่ส่งมา — ไม่เขียนทั้งไฟล์ทับด้วยค่าที่หน้าเว็บถืออยู่ (กันของคนอื่นหาย)
    """
    revs = revs or {}
    done = []

    # ---- ด่านกรอกให้ครบ (user สั่ง 24 ก.ย. 2569) --------------------------------
    # ทำไม: เดิม "ยังไม่ได้กรอก" กับ "ตั้งใจใช้ค่าพื้นฐาน" หน้าตาเหมือนกันเป๊ะ ระบบเลยเงียบ ๆ
    # ใช้ 6 วัน 24 ชม. โดยไม่มีใครรู้ (เคยทำให้ 12 กองคิดกำลังผลิตเกินจริง 31%)
    # ตอนนี้ทุกกองต้องมีทั้งวันและชั่วโมง ไม่มีข้อยกเว้น — เว้นว่าง = บันทึกไม่ผ่าน
    bad = []
    for ch in (changes.get("pools") or []):
        k = _s(ch.get("key")) or "(ไม่ระบุ Cat)"
        for field, label, lo, hi in (("work_days", "วันทำงาน", 0, 7),
                                     ("work_hours", "ชั่วโมงทำงาน", 0, 24)):
            if field not in ch:
                continue
            v = _num(ch.get(field))
            if v is None:
                bad.append(f"{k}: ต้องกรอก{label} (เว้นว่างไม่ได้)")
            elif not lo <= float(v) <= hi:
                bad.append(f"{k}: {label}ต้องอยู่ระหว่าง {lo}-{hi}")
        # ชั่วโมงเฉพาะสัปดาห์: ส่ง None มา = ตั้งใจล้างค่าทิ้ง (กลับไปใช้ชั่วโมงของกอง) จึงผ่านได้
        for w, hv in (ch.get("week_hours") or {}).items():
            if hv is None or _s(hv) == "":
                continue
            v = _num(hv)
            if v is None or not 0 <= float(v) <= 24:
                bad.append(f"{k} W{w}: ชั่วโมงทำงานต้องเป็นตัวเลข 0-24")
        # วันเฉพาะสัปดาห์: None = ล้าง (กลับไปคิดจากค่ามาตรฐานของกอง + ปฏิทิน) ; เพดานตรวจหลังรู้การยุบ
        for w, dv in (ch.get("week_days") or {}).items():
            if dv is None or _s(dv) == "":
                continue
            v = _num(dv)
            if v is None or float(v) < 0:
                bad.append(f"{k} W{w}: วันทำงานต้องเป็นตัวเลข 0 ขึ้นไป")
        # จำนวนเครื่องรายสัปดาห์: None = ไม่ตั้ง (ใช้ Total ของทะเบียน)
        for yw, mv in (ch.get("mc_weeks") or {}).items():
            if mv is None or _s(mv) == "":
                continue
            v = _num(mv)
            if v is None or v < 0 or float(v) != int(v):
                bad.append(f"{k} W{_num(yw) and _num(yw) % 100}: จำนวนเครื่องต้องเป็นจำนวนเต็ม 0 ขึ้นไป")
    for ch in (changes.get("machines") or []):
        if "total" in ch and _num(ch.get("total")) is None:
            bad.append(f"แถวที่ {ch.get('row')}: ต้องกรอกจำนวนเครื่อง (ใส่ 0 ได้ถ้าไม่มีเครื่อง)")
    if bad:
        return {"ok": False, "error": " · ".join(bad[:5])
                + (f" (และอีก {len(bad) - 5} รายการ)" if len(bad) > 5 else "")}

    # ---- รวมค่าที่แก้เข้ากับของเดิม แล้วตรวจด่านยุบสัปดาห์ "ก่อน" เขียนไฟล์ใด ๆ
    # ต้องตรวจก่อนเขียนจำนวนเครื่อง ไม่งั้นด่านไม่ผ่านแต่จำนวนเครื่องถูกเขียนไปแล้วครึ่งทาง
    allp = changes.get("pools") or []
    # จำนวนเครื่องรายสัปดาห์ไปชีท MC_Total — แยกจากแผงวันทำงาน (ไม่ต้องอ่าน/ตรวจแผงวันถ้าแก้แค่นี้)
    srows = [ch for ch in allp if ch.get("mc_weeks")]
    prows = [ch for ch in allp
             if any(f in ch for f in ("work_days", "work_hours", "week_hours", "week_days"))]
    if srows:
        # ตรวจเลขรุ่นก่อนเขียนอะไรทั้งนั้น — ไม่งั้นจำนวนเครื่องเขียนไปแล้วค่อยมาเจอว่ามีคนแก้ชีทนี้แซง
        livesync.check_rev(livesync.masters_res(MC_FILE, SH_TOTAL), revs.get("total"))
        if _timeline() is None:
            return {"ok": False, "error": "อ่านปฏิทินไม่ได้ — บันทึกจำนวนเครื่องรายสัปดาห์ไม่ได้"}
    wd = defaults = hours = weeks = week_hours = None
    nd = nh = nw = nwd = 0
    if prows:
        wd = workday.get_workday()
        tl = _timeline()
        defaults = dict(wd.get("defaults") or {})
        hours = dict(wd.get("hours") or {})
        weeks = dict(wd.get("weeks") or {})
        week_hours = dict(wd.get("week_hours") or {})
        for ch in prows:
            k = _s(ch.get("key"))
            if not k:
                continue
            # ค่าว่างถูกกันไปตั้งแต่ด่านข้างบนแล้ว — ที่มาถึงนี่คือตัวเลขที่ใช้ได้เสมอ
            if "work_days" in ch:
                defaults[k] = _num(ch.get("work_days"))
                nd += 1
            if "work_hours" in ch:
                hours[k] = _num(ch.get("work_hours"))
                nh += 1
            for w, hv in (ch.get("week_hours") or {}).items():
                wk = _num(w)
                if wk is None:
                    continue
                # คีย์ต้องมีปีเสมอ (YW) — เลขเปล่าจากจอรุ่นเก่า → ครั้งที่ใกล้วันนี้สุด
                wkey = f"{k}|{workday._yw_of(wk, None, tl)}"
                v = None if hv is None or _s(hv) == "" else _num(hv)
                if v is None:
                    week_hours.pop(wkey, None)      # ล้างค่า = กลับไปใช้ชั่วโมงของกอง
                else:
                    week_hours[wkey] = v
                nw += 1
            for w, dv in (ch.get("week_days") or {}).items():
                wk = _num(w)
                if wk is None:
                    continue
                wkey = f"{k}|{workday._yw_of(wk, None, tl)}"
                v = None if dv is None or _s(dv) == "" else _num(dv)
                if v is None:
                    weeks.pop(wkey, None)           # ล้างค่า = คิดจากค่ามาตรฐานของกอง + ปฏิทิน
                else:
                    weeks[wkey] = v
                nwd += 1

        merges = wd.get("merges") or {}
        cal_yw = {int(k): v for k, v in (wd.get("cal_days_yw") or {}).items()}
        # วันเฉพาะสัปดาห์ต้องไม่เกินวันที่ปฏิทินเปิด (รวมสัปดาห์ที่ยุบเข้ามา) — เกินคือโรงไม่ได้เปิดถึง
        over = []
        for ch in prows:
            k = _s(ch.get("key"))
            for w, dv in (ch.get("week_days") or {}).items():
                v, wk = _num(dv), _num(w)
                if v is None or wk is None:
                    continue
                yw = workday._yw_of(wk, None, tl)
                cap = _day_cap(yw, merges, cal_yw)
                if cap is not None and float(v) > cap + 1e-9:
                    over.append(f"{k} {workday._wlabel(yw)}: {v:g} วัน เกินวันที่ปฏิทินเปิด {cap:g} วัน")
        if over:
            return {"ok": False, "error": "บันทึกไม่ได้ — " + " · ".join(over[:5])
                    + (f" (และอีก {len(over) - 5} รายการ)" if len(over) > 5 else "")}
        # สัปดาห์ที่รับรวมมาต้องระบุวัน + ชั่วโมงครบทุกกองที่มีเครื่อง (เทียบกับผลหลังแก้ครั้งนี้
        # → กรอกครบในครั้งเดียวก็ผ่านได้เลย ไม่ต้องบันทึกสองรอบ)
        gaps = _merge_gaps(week_hours, weeks, merges, tl, cal_yw)
        if gaps:
            return {"ok": False, "error": _gap_message(gaps)}
        # แผงวันทำงานเขียนเป็นไฟล์สุดท้าย — ต้องตรวจเลขรุ่นตรงนี้ ไม่งั้นจำนวนเครื่อง/MC_Total
        # ถูกเขียนไปแล้วค่อยมาเจอว่ามีคนแก้แผงแซง (บันทึกได้ครึ่งทาง)
        if nd or nh or nw or nwd:
            livesync.check_rev(livesync.WORKDAY, revs.get("workday"))

    # ---- จำนวนเครื่อง
    mrows = changes.get("machines") or []
    if mrows:
        m = masters.read_sheet(MC_FILE, SH_MASTER)
        (toi,) = _cols(m["columns"], "Total MC")
        if toi < 0:
            return {"ok": False, "error": "ชีท Master MC ไม่มีคอลัมน์ 'Total MC'"}
        rows = m["rows"]
        n = 0
        for ch in mrows:
            i = ch.get("row")
            if not isinstance(i, int) or not 0 <= i < len(rows):
                continue
            v = _num(ch.get("total"))
            rows[i][toi] = 0 if v is None else v
            n += 1
        if n:
            masters.write_sheet(MC_FILE, SH_MASTER, m["columns"], rows,
                                user=user, base_rev=revs.get("master") or m.get("rev"))
            done.append(f"จำนวนเครื่อง {n} รายการ")

    # ---- จำนวนเครื่องรายสัปดาห์ (MC_Total)
    if srows:
        info = _pool_info()
        edits = []
        for ch in srows:
            p = info.get(_s(ch.get("key")))
            if not p:
                continue
            ed = {}
            for yw, mv in ch["mc_weeks"].items():
                k = _num(yw)
                if k is None:
                    continue
                ed[int(k)] = None if mv is None or _s(mv) == "" else int(_num(mv))
            if ed:
                edits.append((p, ed))
        n = _save_weeks(edits, revs.get("total"), user)
        if n:
            done.append(f"จำนวนเครื่องรายสัปดาห์ {n} Cat")

    # ---- วันทำงาน / ชั่วโมง (แผงเดียวทั้งไฟล์ → อ่าน-แก้-เขียนทีเดียว)
    if prows and (nd or nh or nw or nwd):
        workday.save_workday(defaults, weeks, hours, wd.get("merges") or {},
                             week_hours=week_hours,
                             user=user, base_rev=revs.get("workday"))
        if nd:
            done.append(f"วันทำงาน {nd} Cat")
        if nh:
            done.append(f"ชั่วโมงทำงาน {nh} Cat")
        if nwd:
            done.append(f"วันเฉพาะสัปดาห์ {nwd} รายการ")
        if nw:
            done.append(f"ชั่วโมงเฉพาะสัปดาห์ {nw} รายการ")

    return {"ok": True, "done": done}


# =============================================================================
# Export / Import ทั้ง 4 ชีทของ MasterMC ที่หน้านี้รวมไว้ (ไฟล์เดียว หลายชีท)
#   วันทำงาน/ชั่วโมง (Calendar/Work Day) ไม่รวม — เป็นแผงที่มีด่านกรอกครบของมันเอง
# =============================================================================
# ชีท → (คอลัมน์ที่ต้องมี, คอลัมน์ที่ต้องเป็นตัวเลขถ้ากรอก, คอลัมน์ที่ห้ามว่าง)
IO_SHEETS = {
    SH_MASTER: (("Factory", "MC_CAT", "MC", "Guage", "Total MC"),
                ("Total MC",), ("MC_CAT", "Guage", "Total MC")),
    SH_TOTAL: (("Week", "Factory", "MC_CAT", "Guage", "Total MC"),
               ("Week", "Total MC", "Year", "Week To", "Year To"), ("Week", "MC_CAT", "Guage", "Total MC")),
    SH_LOCK: (("Week", "MC_CAT", "Guage", "MC Lock"),
              ("Week", "MC Lock", "Year", "Week To", "Year To"), ("Week", "MC_CAT", "Guage", "MC Lock")),
    SH_SPECIAL: (("Factory", "MC_CAT", "MC", "Guage", "POLY", "COTTON"),
                 ("POLY", "COTTON"), ("MC_CAT", "Guage")),
}


def export_book() -> bytes:
    """4 ชีทในไฟล์เดียว — ชื่อชีท/หัวคอลัมน์ตามของจริง จึง Import กลับได้ทันที"""
    import io
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for sh in IO_SHEETS:
        d = masters.read_sheet(MC_FILE, sh)
        ws = wb.create_sheet(sh)
        ws.append([str(c) for c in d["columns"]])
        for r in d["rows"]:
            ws.append([None if v == "" else v for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _check_sheet(sh: str, fcols: list, frows: list) -> dict:
    """ข้อมูลชีทเดียวจากไฟล์ → จัดตามคอลัมน์ของชีทจริง + ตรวจ (ยังไม่เขียน)"""
    need, nums, req = IO_SHEETS[sh]
    cur = masters.read_sheet(MC_FILE, sh)
    cols = list(cur["columns"])
    miss = [n for n in need if _cols(fcols, n)[0] < 0]
    if miss:
        return {"sheet": sh, "ok": False, "error": f"ไม่พบคอลัมน์ {', '.join(miss)}"}
    # จับคู่คอลัมน์ตามชื่อ (ไม่สนตัวพิมพ์/ช่องว่าง) — สลับลำดับในไฟล์ได้
    at = _cols(fcols, *cols)
    warn = []
    lost = [c for c, i in zip(cols, at) if i < 0 and _s(c)
            and any(_s(_get(r, cols.index(c))) for r in cur["rows"])]
    if lost:
        warn.append(f"ไฟล์ไม่มีคอลัมน์ {', '.join(lost)} — คอลัมน์นี้จะว่างทั้งคอลัมน์")
    extra = [_s(c) for c in fcols if _s(c) and _cols(cols, c)[0] < 0]
    if extra:
        warn.append(f"ข้ามคอลัมน์ที่ชีทไม่มี: {', '.join(extra)}")
    num_i = {c: cols.index(c) for c in cols if _key(c) in {_key(n) for n in nums}}
    req_i = {c: cols.index(c) for c in cols if _key(c) in {_key(n) for n in req}}
    # แถวที่ช่องหลักว่าง/ไม่ใช่ตัวเลข = "เตือน" ไม่บล็อก — ชีทจริงมีแถวหมายเหตุ, ค่า "-"
    # และแถวที่ค้างว่างอยู่แล้ว (pipeline ข้ามเอง) ถ้าบล็อก ไฟล์ที่ Export ไปเองจะ Import กลับไม่ได้
    rows, odd = [], []
    for n, fr in enumerate(frows, start=2):            # เลขแถวตาม Excel (แถว 1 = หัวคอลัมน์)
        r = ["" if i < 0 or _get(fr, i) is None else _get(fr, i) for i in at]
        if not any(_s(v) for v in r):
            continue                                   # แถวว่าง ข้าม
        lack = [c for c, i in req_i.items() if not _s(r[i])]
        wrong = [c for c, i in num_i.items() if _s(r[i]) and _num(r[i]) is None]
        if lack or wrong:
            odd.append(f"แถว {n}: " + " · ".join(
                ([f"ว่าง {', '.join(lack)}"] if lack else [])
                + ([f"{', '.join(wrong)} ไม่ใช่ตัวเลข"] if wrong else [])))
        for i in num_i.values():
            if _num(r[i]) is not None:
                r[i] = _num(r[i])
        rows.append(r)
    if odd:
        warn.append(f"{len(odd)} แถวข้อมูลไม่ครบ/ไม่ใช่ตัวเลข (นำเข้าตามไฟล์ ตรวจดูก่อนยืนยัน) — " + " · ".join(odd[:3])
                    + (f" (และอีก {len(odd) - 3} แถว)" if len(odd) > 3 else ""))
    old = sum(1 for r in cur["rows"] if any(_s(v) for v in r))
    # เลขรุ่น ณ ตอนตรวจ — หน้าเว็บส่งกลับมาตอนยืนยัน กันมีคนแก้ชีทนี้แซงระหว่างนั้น
    return {"sheet": sh, "ok": True, "columns": cols, "rows": rows,
            "old": old, "count": len(rows), "warn": warn,
            "rev": livesync.rev(livesync.masters_res(MC_FILE, sh))}


def import_check(content: bytes) -> dict:
    """อ่านไฟล์ที่อัปโหลด → ผลตรวจทีละชีท ; ทับเฉพาะชีทที่มีในไฟล์ (ชื่อชีทต้องตรง)"""
    import io
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("เปิดไฟล์ไม่ได้ — รองรับเฉพาะ .xlsx")
    names = {_key(n): n for n in wb.sheetnames}
    out = []
    for sh in IO_SHEETS:
        src = names.get(_key(sh))
        if not src:
            continue
        data = [list(r) for r in wb[src].iter_rows(values_only=True)]
        if not data:
            out.append({"sheet": sh, "ok": False, "error": "ชีทว่าง"})
            continue
        out.append(_check_sheet(sh, data[0], data[1:]))
    wb.close()
    if not out:
        raise ValueError("ไม่พบชีทที่ใช้ได้ — ชื่อชีทต้องเป็น " + " / ".join(IO_SHEETS))
    return {"ok": all(s["ok"] for s in out), "sheets": out}


def import_apply(content: bytes, revs: dict = None, user: str = "") -> dict:
    """ตรวจซ้ำแล้วเขียนทับ — ชีทไหนไม่ผ่านไม่เขียนสักชีท (ตรวจครบก่อนเขียนอะไรทั้งนั้น)
    revs = {ชื่อชีท: เลขรุ่นที่ได้จาก import_check} — ไม่ตรงของจริง = มีคนแก้แซง (Conflict)"""
    revs = revs or {}
    chk = import_check(content)
    if not chk["ok"]:
        return {"ok": False, "error": " · ".join(f"{s['sheet']}: {s['error']}"
                                                 for s in chk["sheets"] if not s["ok"])}
    # เช็คเลขรุ่นครบทุกชีทก่อน — ไม่งั้นชีทแรกเขียนไปแล้ว ชีทที่สองค่อยเจอว่ามีคนแก้แซง
    for s in chk["sheets"]:
        livesync.check_rev(livesync.masters_res(MC_FILE, s["sheet"]), revs.get(s["sheet"]))
    done = []
    for s in chk["sheets"]:
        masters.write_sheet(MC_FILE, s["sheet"], s["columns"], s["rows"],
                            user=user, base_rev=revs.get(s["sheet"]))
        done.append(f"Import {s['sheet']} {s['count']} แถว (เดิม {s['old']})")
    return {"ok": True, "done": done}
