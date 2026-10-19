"""
mc_ava_live.py — Machine Availability (เครื่องที่เหลือ) สด ๆ จาก booking สำหรับหน้า Dashboard

แหล่งข้อมูล: Booking/view_booking.xlsx (ไฟล์เดียวกับที่การ์ด Booking หน้า DATA ดึงมา)
วิธีคิด    : รัน AVA_MC.py เป็น subprocess → ได้ logic การนับเครื่องเดียวกับ pipeline
             ทุกบรรทัด (พูล/นอกพูล, MC Special POLY-COTTON, MC_Total override, setup)
             แล้วอ่านชีท SUMMARY_MC_REMAIN มาทำตาราง กลุ่มเครื่อง × สัปดาห์

ทำไมไม่ import AVA_MC เข้ามาเรียกตรง ๆ: AVA_MC เป็นสคริปต์ top-level ทั้งไฟล์ (ไม่มี main())
→ แค่ import ก็รันทั้งไฟล์ + เขียน Excel ทันที และ state ค้างใน process ของ server
ถ้าลอก logic มาเขียนใหม่ก็จะกลายเป็นสูตรนับเครื่อง 2 ชุดที่หลุดจากกันเมื่อไหร่ก็ได้

ไฟล์ผลลัพธ์: data_plan/_ava_live.xlsx — ใบเดียวเขียนทับตลอด ตั้งใจให้ชื่อไม่ match
glob "booking_final_ready*.xlsx" จึงไม่ถูก Planning / plan_view / หน้า History หยิบไปใช้
"""
import os
import sys
import subprocess
import threading
import time
from pathlib import Path

import config
import runner
import plan_view

REPO_DIR = config.REPO_DIR
BOOKING_DIR = REPO_DIR / "Booking"
LIVE_FILE = config.OUTPUT_DIR / "_ava_live.xlsx"
SUMMARY_SHEET = "SUMMARY_MC_REMAIN"
OVERFLOW_WEEK = "99"      # sentinel ของ AVA_MC/Planning = "งานล้น" ไม่ใช่สัปดาห์จริง

# AVA_MC โหลดปฏิทินจาก SharePoint + MasterMC + booking ทั้งไฟล์ → ปกติหลักนาที
RUN_TIMEOUT = 900
# บรรทัดท้าย ๆ ของ log ที่เก็บไว้ให้หน้าเว็บดูตอนรันพลาด (ทั้งไฟล์ยาวเป็นพัน)
LOG_TAIL = 40

_lock = threading.Lock()
_state = {
    "running": False,
    "started_at": None,
    "finished_at": None,
    "returncode": None,
    "error": None,
    "log_tail": [],
}
_state_lock = threading.Lock()

# แคชตาราง: (path, mtime) → ผลที่แปลงแล้ว (อ่าน Excel ซ้ำทุกครั้งที่ poll หน้าเว็บไม่ไหว)
_cache: dict = {}

# mtime ของ booking ที่สั่งคำนวณอัตโนมัติไปแล้ว — กันเปิดหน้าซ้ำแล้วรัน AVA_MC วนไม่จบ
_auto_done_mtime: float | None = None


# ---------------------------------------------------------------- booking meta
def booking_meta() -> dict | None:
    """ไฟล์ booking ที่ AVA_MC จะอ่าน (ใหม่สุดในโฟลเดอร์ Booking) → {name, mtime, size}"""
    if not BOOKING_DIR.exists():
        return None
    files = [f for f in BOOKING_DIR.iterdir()
             if f.is_file() and f.suffix.lower() in (".xlsx", ".xls")
             and not f.name.startswith("~$")]
    if not files:
        return None
    f = max(files, key=lambda p: p.stat().st_mtime)
    st = f.stat()
    return {"name": f.name, "mtime": st.st_mtime, "size": st.st_size}


# ---------------------------------------------------------------- run AVA_MC
def _run():
    started = time.time()
    with _state_lock:
        _state.update({"running": True, "started_at": started, "finished_at": None,
                       "returncode": None, "error": None, "log_tail": []})

    env = dict(os.environ)
    env["AVA_OUT_FILE"] = str(LIVE_FILE)
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUNBUFFERED"] = "1"

    rc, err, tail = 1, None, []
    try:
        try:
            proc = subprocess.run(
                [sys.executable, str(REPO_DIR / "AVA_MC.py")],
                cwd=str(REPO_DIR), env=env, timeout=RUN_TIMEOUT,
                stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                text=True, encoding="utf-8", errors="replace",
            )
            rc = proc.returncode
            tail = (proc.stdout or "").splitlines()[-LOG_TAIL:]
            if rc != 0:
                err = "AVA_MC จบด้วย returncode=%d — ดูรายละเอียดใน log ด้านล่าง" % rc
        except subprocess.TimeoutExpired:
            err = f"รันเกิน {RUN_TIMEOUT} วินาที — ยกเลิก"
        except Exception as e:                   # noqa: BLE001 — รายงานให้หน้าเว็บทุกกรณี
            err = f"รัน AVA_MC ไม่สำเร็จ: {e}"
    finally:
        # ต้องปลด lock ทุกเส้นทาง ไม่งั้นค้างจนกดคำนวณใหม่ไม่ได้อีกเลยจนกว่าจะรีสตาร์ต server
        with _state_lock:
            _state.update({"running": False, "finished_at": time.time(),
                           "returncode": rc, "error": err, "log_tail": tail})
        _lock.release()


def _booking_has_rows() -> bool:
    """ไฟล์ booking มีแถวข้อมูลจริงไหม (ไม่ใช่มีแต่หัวคอลัมน์)

    ต้องเช็คก่อนรัน: AVA_MC ทำงานกับ booking เปล่าไม่ได้ (merge บนตารางว่าง → ValueError
    'columns overlap but no suffix specified') แล้ว traceback ยาว ๆ ไม่ได้บอกอะไรผู้ใช้เลย
    """
    meta = booking_meta()
    if not meta:
        return False
    p = BOOKING_DIR / meta["name"]
    if p.suffix.lower() != ".xlsx":
        return True                       # .xls อ่านแบบเร็วไม่ได้ → ปล่อยให้ AVA_MC ตัดสิน
    try:
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            for i, _ in enumerate(ws.iter_rows(values_only=True)):
                if i >= 1:                # แถวที่ 2 = ข้อมูลแถวแรก
                    return True
            return False
        finally:
            wb.close()
    except Exception:
        return True                       # อ่านไม่ได้ = ไม่ฟันธง ปล่อยให้ AVA_MC ลองเอง


def start_refresh() -> dict:
    """สั่งคำนวณใหม่จาก booking ล่าสุด (ทำใน background) → {ok, message, reason}

    reason: empty = booking ไม่มีข้อมูล / busy = ติดงานอื่นอยู่ (ลองใหม่ได้) / None = เริ่มแล้ว
    """
    if not _booking_has_rows():
        return {"ok": False, "reason": "empty",
                "message": "ไฟล์ booking ยังไม่มีข้อมูล (มีแต่หัวคอลัมน์) — "
                           "กด \"ดึงเฉพาะ Booking\" ที่หน้า DATA ให้ได้ข้อมูลก่อน"}
    if runner.get_status().get("running"):
        # pipeline กำลังรันอยู่ = AVA_MC อีกตัวกำลังอ่านไฟล์ชุดเดียวกัน — รอให้จบก่อน
        return {"ok": False, "reason": "busy",
                "message": "pipeline กำลังรันอยู่ — รอให้เสร็จก่อนค่อยคำนวณ Machine Availability"}
    if not _lock.acquire(blocking=False):
        return {"ok": False, "reason": "busy", "message": "กำลังคำนวณอยู่แล้ว — รอสักครู่"}
    threading.Thread(target=_run, daemon=True).start()
    return {"ok": True, "reason": None, "message": "เริ่มคำนวณ Machine Availability จาก booking ล่าสุดแล้ว"}


def _maybe_auto_refresh(bk, mtime) -> bool:
    """booking ใหม่กว่าตัวเลขที่มีอยู่ → สั่งคำนวณให้เองตอนเปิดหน้า (คืน True ถ้าเริ่มรันแล้ว)

    กันรันวน: จำ mtime ของ booking ที่สั่งไปแล้ว — booking เวอร์ชันเดิมจะสั่งซ้ำไม่ได้
    แม้รอบก่อนจะล้มเหลว (ไฟล์เสีย/ข้อมูลไม่ครบ) ไม่งั้นทุกครั้งที่เปิดหน้าจะรัน AVA_MC ใหม่ไม่จบสิ้น
    ส่วนกรณี 'busy' (pipeline หรือรอบก่อนยังไม่จบ) ไม่จำ เพื่อให้ลองใหม่รอบหน้าได้
    """
    global _auto_done_mtime
    if not bk or bk["mtime"] <= mtime:
        return False                       # ตัวเลขที่มีใหม่กว่า booking แล้ว ไม่ต้องทำอะไร
    if _auto_done_mtime == bk["mtime"]:
        return False
    r = start_refresh()
    if r["ok"] or r.get("reason") == "empty":
        _auto_done_mtime = bk["mtime"]
    return bool(r["ok"])


def status() -> dict:
    with _state_lock:
        s = dict(_state)
    s["booking"] = booking_meta()
    s["computed_at"] = LIVE_FILE.stat().st_mtime if LIVE_FILE.exists() else None
    return s


# ---------------------------------------------------------------- summary → ตาราง
def _gkey(g) -> str:
    try:
        return str(int(float(g)))
    except (TypeError, ValueError):
        return str(g).strip()


def _wk_int(w) -> int:
    """เลขสัปดาห์ของคีย์ ('40' / '12@2027' → 40 / 12) ; อ่านไม่ออก = 9999"""
    try:
        return int(str(w).split("@", 1)[0])
    except (TypeError, ValueError):
        return 9999


def _int(v) -> int:
    n = plan_view._num(v)
    if n is None or n != n:       # n != n → NaN
        return 0
    return int(n)


def _split_pools(df) -> set:
    """(cat|gauge) ที่แยกเป็นหลายพูล — กติกาเดียวกับ plan_view._split_pool_set()
    แต่ดูจาก df ที่ส่งเข้ามา (ไฟล์ live) ไม่ใช่ไฟล์ผลลัพธ์ pipeline"""
    if "MC_POOL" not in df.columns:
        return set()
    by_cg: dict = {}
    for _, r in df.iterrows():
        if str(r["MC_CAT"]).strip().upper() != "SINGLE-32":
            continue
        pool = str(r.get("MC_POOL", "")).strip()
        if not pool:
            continue
        by_cg.setdefault(f"{str(r['MC_CAT']).strip()}|{_gkey(r['GUAGE'])}", set()).add(pool)
    return {cg for cg, pools in by_cg.items() if len(pools) > 1}


def _build_table(path: Path) -> dict:
    """อ่านชีท SUMMARY_MC_REMAIN → {weeks, rows} (แถว = กลุ่มเครื่อง, ช่อง = ต่อสัปดาห์)

    ⚠️ แถวในชีทมี 2 ชนิด และห้ามบวกรวมกัน:
      - pool row (FACTORY ไม่ว่าง): TOTAL_MC / MC_USE_CEIL เป็นยอด "ทั้งพูล" ซึ่งซ้ำกัน
        ทุกแถวของพูลเดียวกัน (พูลหนึ่งครอบได้หลาย MC_CAT+เกจ) → นับพูลละครั้ง
      - non-pool row (FACTORY ว่าง): ยอดต่อ MC_CAT+เกจ ตรง ๆ
    AVA_MC ตัดแถว non-pool ที่ซ้ำกับ pool row ออกให้แล้ว จึงไม่ทับกันเอง
    """
    import pandas as pd
    df = pd.read_excel(path, sheet_name=SUMMARY_SHEET)
    need = {"MC_CAT", "GUAGE", "WEEK", "TOTAL_MC", "MC_USE_CEIL", "TOTAL_MC_REMAIN"}
    if not need <= set(df.columns):
        return {"weeks": [], "rows": [], "error": f"ชีท {SUMMARY_SHEET} ไม่มีคอลัมน์ที่ต้องใช้"}

    has_fac = "FACTORY" in df.columns
    has_pool = "MC_POOL" in df.columns
    has_year = "YEAR" in df.columns
    split = _split_pools(df)

    rows: dict = {}          # key → {meta..., cells: {week: {...}}}
    weeks: set = set()
    years: dict = {}
    for _, r in df.iterrows():
        # คีย์สัปดาห์ติดปีเมื่อไกลเกินครึ่งปี (plan_view.week_key) — ตารางเครื่องยาวหลายปีแล้ว
        # ถ้าคีย์ด้วยเลขเปล่า W40 ปีนี้กับ W40 ปีหน้าจะถูกรวมเป็นช่องเดียว
        wk = plan_view.week_key(r["WEEK"], r.get("YEAR") if has_year else None)
        if wk is None or wk == OVERFLOW_WEEK:
            continue                      # W99 = sentinel "งานล้น" ไม่ใช่สัปดาห์จริง
        cat = str(r["MC_CAT"]).strip()
        gauge = _gkey(r["GUAGE"])
        pool = str(r.get("MC_POOL", "")).strip() if has_pool else ""
        fac = str(r["FACTORY"]).strip() if has_fac else ""
        if fac.lower() == "nan":
            fac = ""
        is_pool = bool(fac) and bool(pool)
        # พูล = 1 แถวตาราง (ยอดของพูลซ้ำอยู่ทุกแถว) ; นอกพูล = แยกตาม cat+เกจ
        key = pool if is_pool else f"{cat}|{gauge}"

        weeks.add(wk)
        if has_year:
            y = _int(r.get("YEAR"))
            if y and wk not in years:
                years[wk] = y

        row = rows.setdefault(key, {
            "key": key, "factory": fac, "pool": pool, "is_pool": is_pool,
            "cats": [], "cells": {},
        })
        cg = f"{cat} {gauge}G"
        if cg not in row["cats"]:
            row["cats"].append(cg)
        # เก็บ (cat, เกจ) ไว้จับคู่ Lock_MC — ต้อง normalize เกจด้วยฟังก์ชันเดียวกับฝั่ง Lock_MC
        # ('26.0' / '26G' → '26') ไม่งั้นคีย์ไม่ตรงแล้วหักเครื่องที่กันไว้ไม่ติดแบบเงียบ ๆ
        row.setdefault("_cg", set()).add((cat.upper(), plan_view._lock_gkey(gauge)))

        total, used = _int(r["TOTAL_MC"]), _int(r["MC_USE_CEIL"])
        remain = _int(r["TOTAL_MC_REMAIN"])
        cell = row["cells"].get(wk)
        if cell is None:
            row["cells"][wk] = {"total": total, "used": used, "remain": remain}
        elif is_pool:
            # แถวซ้ำของพูลเดียวกัน = ยอดชุดเดิม → ไม่บวก (max กันค่าเพี้ยนจาก MC_Total override)
            cell["total"] = max(cell["total"], total)
            cell["used"] = max(cell["used"], used)
            cell["remain"] = min(cell["remain"], remain)
        else:
            cell["total"] += total
            cell["used"] += used
            cell["remain"] += remain

    # หักเครื่องที่ Lock_MC กันไว้ (user กันเครื่องไว้ให้ item → แผนเอาไปใช้ไม่ได้)
    # ทำหลังรวมแถวเสร็จ เพราะ key เดียวมาจากหลายแถวได้ — หักตอน loop จะหักซ้ำเท่าจำนวนแถว
    lock = plan_view._lock_mc_by_week()
    if lock:
        # แถวล็อกที่ระบุคอลัมน์ MC หักเฉพาะพูลของกลุ่มเครื่องนั้น — แมปเครื่อง→พูล ต้องอ่านจาก
        # ไฟล์ใบเดียวกับที่ตารางนี้ใช้ (หน้านี้คำนวณเองได้ ไม่ใช่ใบของ pipeline เสมอไป)
        pool_of = plan_view.lock_pool_resolver(path)
        for row in rows.values():
            for wk, cell in row["cells"].items():
                w = str(wk)       # คีย์ของ Lock_MC เป็นคีย์สัปดาห์ชุดเดียวกัน (plan_view._lock_span)
                # พูลที่ครอบหลาย cat+เกจ → รวมล็อกของทุก cat+เกจ ในพูล (คนละชุดเครื่องกัน)
                n = sum(plan_view.lock_mc_for_pool(w, c, g, row["pool"], pool_of)
                        for c, g in row.get("_cg", set()))
                if n <= 0:
                    continue
                cell["lock"] = n
                cell["remain"] -= n
                cell["used"] += n       # คงสมการ total − used = remain

    # เรียงสัปดาห์ตามเวลาจริง: มีคอลัมน์ YEAR → (ปี, สัปดาห์)
    # ไม่มี YEAR (ไฟล์รุ่นเก่า) → สัปดาห์ต้นปีที่โผล่คู่กับสัปดาห์ปลายปีคือ "ปีถัดไป" ให้ไปอยู่ท้าย
    # (ไม่งั้นตารางขึ้นต้นด้วย W1-W4 ของปีหน้า แล้วตามด้วย W27 ของปีนี้)
    wrap = (not years
            and any(_wk_int(w) <= 13 for w in weeks)
            and any(_wk_int(w) >= 40 for w in weeks))

    def _wnum(w):
        n = _wk_int(w)
        y = years.get(w)
        if y:
            return (y, n)
        return (1 if (wrap and n <= 13) else 0, n)

    out_rows = []
    for row in sorted(rows.values(),
                      key=lambda r: (r["factory"] or "zz", r["cats"][0] if r["cats"] else "")):
        row.pop("_cg", None)
        row["cat"] = ", ".join(row["cats"])
        # ป้ายพูลให้อ่านง่าย: ตัด "FACTORY|" ข้างหน้าออก (โชว์โรงงานเป็นคอลัมน์แยกอยู่แล้ว)
        row["pool_label"] = row["pool"].split("|", 1)[-1] if row["pool"] else ""
        out_rows.append(row)

    return {"weeks": sorted(weeks, key=_wnum), "years": years, "rows": out_rows}


def _current_source():
    """(source, path) ของไฟล์ผลลัพธ์ที่หน้านี้ใช้อยู่ — ใบที่ "ใหม่กว่า" ระหว่างไฟล์ที่หน้านี้
    คำนวณเอง กับผลลัพธ์ AVA_MC ของ pipeline → pipeline รันรอบไหน (รวมรอบอัตโนมัติ)
    หน้านี้อัปเดตตามเองโดยไม่ต้องกดคำนวณ ; (None, None) = ยังไม่มีผลลัพธ์เลย"""
    cands = []
    if LIVE_FILE.exists():
        cands.append(("live", LIVE_FILE))
    # full=True: หน้านี้คือ booking เต็มทุกสัปดาห์ — ไม่ส่งจะได้ booking_final_replan
    # (ตัดเหลือแต่สัปดาห์ล็อก) → สัปดาห์ที่ปลดล็อกกลายเป็น 0% ทั้งที่ booking จองไว้เต็ม
    p = plan_view._latest_booking_path(full=True)
    if p is not None:
        cands.append(("pipeline", p))
    if not cands:
        return None, None
    return max(cands, key=lambda c: c[1].stat().st_mtime)


def table() -> dict:
    """ตาราง Machine Availability ล่าสุด — ใช้ไฟล์ live ถ้ามี ไม่มีก็ตกไปใช้ผลลัพธ์ pipeline ล่าสุด
    (จะได้มีอะไรให้ดูก่อนกดคำนวณครั้งแรก) พร้อม meta ให้หน้าเว็บบอกที่มา/ความสด"""
    src, path = _current_source()
    if path is None:
        return {"ok": False, "source": None, "weeks": [], "rows": [],
                "booking": booking_meta(),
                "reason": "ยังไม่เคยคำนวณ — กดปุ่ม \"คำนวณใหม่\" เพื่อสรุป Machine Availability จาก booking"}

    mtime = path.stat().st_mtime
    cached = _cache.get((str(path), mtime))
    if cached is None:
        try:
            cached = _build_table(path)
        except Exception as e:                    # noqa: BLE001
            return {"ok": False, "source": src, "weeks": [], "rows": [],
                    "booking": booking_meta(), "reason": f"อ่านไฟล์ผลลัพธ์ไม่ได้: {e}"}
        _cache.clear()                            # เก็บแค่ผลล่าสุดพอ
        _cache[(str(path), mtime)] = cached

    bk = booking_meta()
    stale = bool(bk and bk["mtime"] > mtime)
    out = dict(cached)
    out.update({
        "ok": True,
        "source": src,
        "file": path.name,
        "computed_at": mtime,
        "booking": bk,
        # booking ใหม่กว่าผลที่คำนวณไว้ → ตัวเลขบนหน้าจอเก่ากว่าข้อมูลจริง
        "stale": stale,
        # เก่าแล้ว → สั่งคำนวณให้เองเลย (ผู้ใช้ไม่ต้องกด) หน้าเว็บ poll สถานะแล้วโหลดซ้ำเมื่อเสร็จ
        "auto_started": _maybe_auto_refresh(bk, mtime) if stale else False,
    })
    if not out["rows"]:
        out["reason"] = ("ไฟล์ booking ยังไม่มีข้อมูล — ดึง Booking ใหม่แล้วกดคำนวณอีกครั้ง"
                         if (bk and bk["size"] < 20_000) else
                         "ไม่พบข้อมูลเครื่องในผลลัพธ์ที่คำนวณได้")
    return out


# ---------------------------------------------------------------- รายละเอียดรายช่อง
ITEM_CAP = 1500          # กันหน้าเว็บอืดถ้าสัปดาห์เดียวมีงานเยอะผิดปกติ


def _pool_base(pool_type: str) -> str:
    """POOL_TYPE ของแถว DETAIL → คีย์พูลระดับที่ heatmap ใช้
    ตัด suffix ':POLY' / ':COTTON' (พูลย่อยของเครื่องที่กันไว้) ออก เพื่อให้จับคู่กันได้"""
    s = str(pool_type or "").strip()
    for sfx in (":POLY", ":COTTON"):
        if s.upper().endswith(sfx):
            return s[: -len(sfx)]
    return s


def detail(key: str, week: str) -> dict:
    """รายละเอียดของช่อง (กลุ่มเครื่อง × สัปดาห์) ที่ผู้ใช้คลิกบน heatmap

    - summary : ตัวเลขชุดเดียวกับช่องบน heatmap (เครื่องทั้งหมด/ใช้ไป/เหลือ/%)
    - groups  : กลุ่มเครื่องย่อย (MC_GROUP) ในพูลนั้น — ใช้เครื่องกี่ตัว กี่งาน
    - items   : รายการงาน booking ที่ลงสัปดาห์นั้น (จากชีท DETAIL)
    """
    src, path = _current_source()
    if path is None:
        return {"ok": False, "reason": "ยังไม่มีผลลัพธ์ที่คำนวณไว้"}

    tb = table()
    row = next((r for r in tb.get("rows", []) if r["key"] == key), None)
    if row is None:
        return {"ok": False, "reason": "ไม่พบกลุ่มเครื่องนี้ในผลลัพธ์ล่าสุด"}
    cell = row["cells"].get(str(week))

    try:
        import pandas as pd
        xl = pd.ExcelFile(path)
        df = xl.parse("DETAIL")
    except Exception as e:                            # noqa: BLE001
        return {"ok": False, "reason": f"อ่านชีท DETAIL ไม่ได้: {e}"}

    # เครื่องที่ MC Special กันไว้ให้ POLY/COTTON ของพูลนี้ (ชีท MC_RESERVED ใบเดียวกัน)
    # ยอด total ของช่อง = "พูลปกติ" ที่หักส่วนนี้ออกแล้ว → การ์ดต้องโชว์ทั้งกองให้เห็นว่าเครื่องไม่ได้หาย
    rsv_total = {"POLY": 0, "COTTON": 0}
    if row["is_pool"] and "MC_RESERVED" in xl.sheet_names:
        try:
            rdf = xl.parse("MC_RESERVED")
            if {"MC_POOL", "POLY", "COTTON"} <= set(rdf.columns):
                for _, rr in rdf[rdf["MC_POOL"].astype(str).str.strip() == row["pool"]].iterrows():
                    rsv_total["POLY"] += _int(rr["POLY"])
                    rsv_total["COTTON"] += _int(rr["COTTON"])
        except Exception:                             # noqa: BLE001 — ไม่มีก็แค่ไม่โชว์ส่วนกันไว้
            pass

    if "WEEK" not in df.columns:
        return {"ok": False, "reason": "ชีท DETAIL ไม่มีคอลัมน์ WEEK"}
    # คีย์ช่องเป็นคีย์สัปดาห์ของหน้าเว็บ (ติดปีเมื่อไกล) → แปลงแถว DETAIL ด้วยกติกาเดียวกัน
    # เทียบเลขเปล่าตรง ๆ จะดึงงาน W40 ของทุกปีมาปนในช่องเดียว
    if "YEAR" in df.columns:
        wk_keys = [plan_view.week_key(w, y) for w, y in zip(df["WEEK"], df["YEAR"])]
    else:
        wk_keys = [plan_view.week_key(w) for w in df["WEEK"]]
    df = df[[k == str(week) for k in wk_keys]]

    # จับคู่แถว DETAIL กับกลุ่มที่คลิก: pool row → เทียบ POOL_TYPE (ตัด suffix พูลย่อย)
    #                                  non-pool → POOL_TYPE ว่าง + MC_CAT|เกจ ตรงกัน
    has_pool_col = "POOL_TYPE" in df.columns
    if row["is_pool"] and has_pool_col:
        sel = df[df["POOL_TYPE"].map(lambda v: _pool_base(v) == row["pool"])]
    else:
        cat_key, _, g_key = key.partition("|")
        sel = df[df.apply(
            lambda r: (not has_pool_col or not str(r.get("POOL_TYPE") or "").strip())
            and str(r.get("MC_CAT", "")).strip().upper() == cat_key.strip().upper()
            and _gkey(r.get("GUAGE")) == g_key,
            axis=1,
        )] if len(df) else df

    def _s(v) -> str:
        s = str(v).strip()
        return "" if s.lower() in ("nan", "nat", "none") else s

    def _num(v) -> float:
        return plan_view._f(v)

    def _cap_adj(r, kp, cap_raw) -> float:
        """CAP ที่ AVA_MC ใช้หารจริง (ฐาน 24 ชม. × ชั่วโมงของกลุ่ม/24, Item Special ชนะ)
        AVA_MC ทิ้ง _CAP_ADJ ก่อนเขียน DETAIL → ย้อนจาก MC_USE = KP / (CAP_ADJ × วันทำงาน)
        ได้เลขเดียวกับที่ใช้คิดเครื่องเป๊ะ ไม่ต้องลอกสูตรชั่วโมง/Item Special มาซ้ำที่นี่"""
        use = _num(r.get("MC_USE"))
        wd = _num(r.get("_effective_working_days")) or _num(r.get("WORKING_DAY"))
        return kp / (use * wd) if (use > 0 and wd > 0 and kp > 0) else cap_raw

    items, groups = [], {}
    for _, r in sel.iterrows():
        mg = _s(r.get("MC_GROUP")).upper()
        mc = int(_num(r.get("MC_USE_CEIL")))
        kp = _num(r.get("KP_WEIGHT"))
        cap_raw = _num(r.get("CAP_KNIT"))
        cap = _cap_adj(r, kp, cap_raw)
        setup = bool(r.get("_is_new_setup")) if "_is_new_setup" in sel.columns else False
        # พูลย่อยของเครื่องที่กันไว้ (POLY/COTTON) — ตัวเลขชุดนี้ AVA_MC นับแยกจากพูลหลัก
        rsv = ""
        if has_pool_col:
            pt = str(r.get("POOL_TYPE") or "").strip().upper()
            rsv = "COTTON" if pt.endswith(":COTTON") else "POLY" if pt.endswith(":POLY") else ""
        items.append({
            "mc_group": mg,
            "gauge": _gkey(r.get("GUAGE")),
            "item": _s(r.get("ITEM_CODE")),
            "desc": _s(r.get("DESCRIPTION")),
            "so": _s(r.get("SO_NO")),
            "team": _s(r.get("TEAM_NAME")),
            "kp": round(kp, 1),
            "cap": round(cap, 1),
            "cap_raw": round(cap_raw, 1),
            "work_days": round(_num(r.get("_effective_working_days")) or _num(r.get("WORKING_DAY")), 2),
            "mc": mc,
            "mc_raw": round(_num(r.get("MC_USE")), 2),
            "setup": setup,
            "setup_days": int(_num(r.get("_setup_days"))) if "_setup_days" in sel.columns else 0,
            "material": _s(r.get("MATERIAL_CONTENT")),
            "yarn": _s(r.get("YARN_USED")),
            "reserved": rsv,
        })
        g = groups.setdefault(mg or "-", {"mc_group": mg or "-", "gauge": _gkey(r.get("GUAGE")),
                                          "mc": 0, "reserved_mc": 0, "items": 0, "kp": 0.0, "setup": 0})
        g["mc"] += mc
        g["reserved_mc"] += mc if rsv else 0      # ส่วนที่ลงเครื่องโควตา POLY/COTTON
        g["items"] += 1
        g["kp"] += kp
        g["setup"] += 1 if setup else 0

    items.sort(key=lambda x: (-x["mc"], -x["kp"]))
    for g in groups.values():
        g["kp"] = round(g["kp"], 1)

    total = cell["total"] if cell else 0
    used = cell["used"] if cell else 0
    mc_items = sum(i["mc"] for i in items)
    reserved_mc = sum(i["mc"] for i in items if i["reserved"])
    # เครื่องกันไว้แยกชนิด: กันไว้กี่ตัว / booking สัปดาห์นี้ใช้ไปกี่ตัว (โชว์เฉพาะชนิดที่มีจริง)
    reserved = [
        {"kind": k, "total": rsv_total[k],
         "used": sum(i["mc"] for i in items if i["reserved"] == k)}
        for k in ("POLY", "COTTON")
    ]
    reserved = [x for x in reserved if x["total"] or x["used"]]
    rsv_sum = sum(x["total"] for x in reserved)

    return {
        "ok": True,
        "key": key,
        "week": str(week),
        "year": (tb.get("years") or {}).get(str(week)),
        "factory": row["factory"],
        "cat": row["cat"],
        "pool_label": row["pool_label"],
        "summary": {
            "total": total,
            "used": used,
            "remain": cell["remain"] if cell else 0,
            "lock": (cell or {}).get("lock", 0),
            "pct": round(used / total * 100) if total else None,
            "items": len(items),
            "kp": round(sum(i["kp"] for i in items), 1),
            "mc_items": mc_items,
            # เครื่องในโควตา POLY/COTTON ถูกกันออกจากพูลหลักตั้งแต่ AVA_MC → ตัวเลขสองฝั่งไม่เท่ากันได้
            "reserved_mc": reserved_mc,
            "reserved": reserved,
            # ทั้งกอง = พูลปกติ + เครื่องที่กันไว้ (ตรงกับยอดที่หน้าเครื่องจักรแสดง/ให้กรอก)
            "pile_total": total + rsv_sum,
            "pile_used": used + reserved_mc,
            "groups_used": sum(1 for g in groups.values() if g["mc"] > 0),
            "groups_all": len(groups),
        },
        "groups": sorted(groups.values(), key=lambda g: (-g["mc"], g["mc_group"])),
        "items": items[:ITEM_CAP],
        "truncated": len(items) > ITEM_CAP,
    }
