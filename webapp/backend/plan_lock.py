"""
plan_lock.py — ล็อกแผน (Lock_Plan) : งาน booking ที่ห้าม Replan รื้อไปวางใหม่
=============================================================================

เก็บใน **ชีท "Lock_Plan" ของ MasterMC.xlsx** (ที่เดียวกับ Lock_MC / MC_Total / Job)
→ โผล่ในหน้า Master ของเว็บเอง (list_masters อ่าน sheetnames ทั้งหมด) แก้มือได้
   ทั้งใน Excel และในหน้าเว็บ ; กดล็อกบน Gantt ก็เขียนลงชีทเดียวกันนี้

1 แถว = 1 บล็อกงาน × สัปดาห์ ที่ถูกล็อก
    WEEK ITEM_CODE MC_GROUP GUAGE SO_NO QTY MC SOURCE SHEET LOCKED_BY LOCKED_AT REMARK

คีย์บล็อก = (WEEK, ITEM_CODE, MC_GROUP, GUAGE) + SOURCE ; SOURCE ชี้ว่าล็อกนี้หมายถึงแถวไหน
  - SOURCE=BOOKING → แถวในชีท DETAIL ของ booking_final (ของที่จองเครื่องไปแล้ว)
  - SOURCE=PLAN    → แถวในชีท PLAN ของไฟล์แผน (งาน SC ที่ Planning วางไว้) ; SHEET บอกว่า
                     กดล็อกจากชีทไหน (PLAN = production_plan_*.xlsx / REPLAN = replan_result.xlsx)
                     เพราะสองชีทนี้อ่านคนละไฟล์ Replan ต้องย้อนไปอ่านไฟล์ให้ถูกใบ
  - แถวเก่าที่ยังไม่มีคอลัมน์ SOURCE = BOOKING (ค่าเริ่มต้น เข้ากันได้กับล็อกที่บันทึกไว้ก่อนหน้า)

  - SO_NO เว้นว่าง = ล็อกทุก SO ของ item นั้นในช่องนั้น
  - SO_NO มีค่า    = ล็อกเฉพาะใบนั้น (คั่นด้วย , ได้หลายใบ — DETAIL เก็บหลาย SO ต่อแถว)
  - QTY / MC เป็นข้อมูลอ้างอิงตอนกดล็อก ; ฝั่ง SOURCE=PLAN ใช้เป็นค่าสำรองเมื่อหาแถวจริง
    ในไฟล์แผนไม่เจอ (แผน gen ใหม่แล้วงานย้ายสัปดาห์) — ดู Replan.py

ล็อก = "ตัวชี้" ไปที่แถวในไฟล์แผน/booking ไม่ได้เก็บตัวงานไว้เอง
ฝั่งที่ "ใช้" ล็อกจริงคือ Replan.py (load_plan_locks()) — ดูหัวไฟล์นั้นประกอบ
"""
import re
import shutil
from contextlib import contextmanager
from datetime import datetime

from openpyxl import load_workbook

import livesync
import masters

SHEET = "Lock_Plan"
# YEAR ต่อท้าย (ต้องตรงกับ LOCK_COLUMNS ใน program_lock.py) ; ว่าง = ครั้งที่ใกล้วันนี้สุด
COLUMNS = ["WEEK", "ITEM_CODE", "MC_GROUP", "GUAGE", "SO_NO", "QTY", "MC",
           "SOURCE", "SHEET", "LOCKED_BY", "LOCKED_AT", "REMARK", "YEAR"]

SRC_BOOKING = "BOOKING"
SRC_PLAN = "PLAN"

# ── ล็อกอัตโนมัติ (ตั๋วโปรแกรม + core item) + ทะเบียนช่วงสัปดาห์ที่ user ปลดแล้ว ──────
# program_lock.py (ก่อน Replan) เติมแถวล็อกให้เอง ประทับ LOCKED_BY = AUTO-PROGRAM (item ในชีท
# Program) หรือ AUTO-CORE (core item ในทะเบียน Target_Stock) ; user ปลดล็อกแถวพวกนี้ได้ตามปกติ
# แต่ต้อง "จำ" ไว้ด้วยว่าปลดสัปดาห์ไหนไปแล้ว ไม่งั้นรอบหน้าระบบจะล็อกกลับให้ทันที
# → เก็บเป็น "ช่วงสัปดาห์" ต่อ item ในชีท Program_Unlock / Core_Unlock (ที่เดียวกับที่ program_lock.py อ่าน)
AUTO_BY = "AUTO-PROGRAM"
AUTO_CORE_BY = "AUTO-CORE"
# LOCKED_BY → ชนิดที่ส่งให้หน้าเว็บเลือกข้อความบนบล็อก (ต้องตรงกับ KINDS ใน program_lock.py)
AUTO_KINDS = {AUTO_BY: "program", AUTO_CORE_BY: "core"}
UNLOCK_SHEET = "Program_Unlock"
CORE_UNLOCK_SHEET = "Core_Unlock"
# YEAR_FROM/YEAR_TO ต่อท้าย (ต้องตรงกับ program_lock.py) ; ว่าง = ครั้งที่ใกล้วันนี้สุด (แถวเก่า)
UNLOCK_COLUMNS = ["ITEM_CODE", "WEEK_FROM", "WEEK_TO", "MC_GROUP", "GUAGE",
                  "UNLOCKED_BY", "UNLOCKED_AT", "REMARK", "YEAR_FROM", "YEAR_TO"]


def _norm(v) -> str:
    return str(v if v is not None else "").strip()


def _source(v) -> str:
    """SOURCE → 'BOOKING' | 'PLAN' (ว่าง = BOOKING เพื่อความเข้ากันได้กับแถวเก่า)"""
    s = _norm(v).upper()
    return SRC_PLAN if s == SRC_PLAN else SRC_BOOKING


def _sheet(v) -> str:
    """SHEET → ชื่อชีทที่กดล็อก ('PLAN'/'REPLAN') ; ว่าง = PLAN (ชีทหลัก)"""
    s = _norm(v).upper().replace(" ", "_")
    return s or "PLAN"


def _gkey(g) -> str:
    """normalize เกจ: 22.0 / '22 ' → '22' (ต้องตรงกับ _gkey ใน plan_view.py)"""
    try:
        return str(int(float(g)))
    except (TypeError, ValueError):
        return _norm(g)


def _wk(v) -> str:
    """WEEK → string เลขสัปดาห์ ('29.0' → '29')"""
    try:
        return str(int(float(v)))
    except (TypeError, ValueError):
        return _norm(v)


def _wkey(week, year=None) -> str:
    """(WEEK[, YEAR]) → คีย์สัปดาห์ของหน้าเว็บ ('29' / '29@2028') — กติกาเดียวกับคอลัมน์ Gantt
    import ช้า: plan_view import โมดูลอื่นเยอะ (โหลดตอนเรียกจริงครั้งแรกพอ)"""
    import plan_view
    k = plan_view.week_key(week, year if _norm(year) else None)
    return k if k is not None else _norm(week)


def _split_key(key) -> tuple:
    """คีย์สัปดาห์ → (WEEK, YEAR) สำหรับเขียนลงชีท — เขียนปีจริงเสมอ (ไม่ปล่อยว่าง)
    ปีว่างในชีท = "ครั้งที่ใกล้วันนี้สุด" ซึ่งเลื่อนไปตามวันที่ ; ล็อกเขียนวันนี้ต้องชี้สัปดาห์เดิมตลอด"""
    import plan_view
    w, y = plan_view.key_week_year(key)
    if w is None:
        return _norm(key), ""
    return w, (y if y is not None else "")


def _yw_of_key(key):
    """คีย์สัปดาห์ → YW (ปี*100+สัปดาห์) ; อ่านไม่ออก = None"""
    w, y = _split_key(key)
    return y * 100 + w if isinstance(w, int) and isinstance(y, int) else None


_TL_CACHE: dict = {}


def _yw_shift(yw, n: int):
    """เลื่อน YW ไป n สัปดาห์บนปฏิทินจริง (ข้ามปี / ปี 52-53 ถูกต้อง)
    ปฏิทินอ่านไม่ได้ → คิดเลขเอง (ถือปีละ 52 สัปดาห์)"""
    if yw is None:
        return None
    try:
        import plan_view
        from Calendar import WeekTimeline
        weekly = plan_view._calendar_weekly()
        if weekly is None:
            raise ValueError("no calendar")
        key = id(weekly), len(weekly)
        tl = _TL_CACHE.get("tl")
        if tl is None or _TL_CACHE.get("k") != key:
            tl = WeekTimeline(weekly)
            _TL_CACHE.update(tl=tl, k=key)
        out = tl.shift(yw, n)
        if out is not None:
            return out
    except Exception:
        pass
    y, w = divmod(int(yw), 100)
    w += int(n)
    while w < 1:
        y, w = y - 1, w + 52
    while w > 52:
        y, w = y + 1, w - 52
    return y * 100 + w


def _so_set(v) -> set:
    """SO_NO ('250012, 250013') → {'250012','250013'} ; ว่าง = set() = ทุก SO

    ตัด '.0' ท้ายเลขทิ้ง — ฝั่ง booking อ่านด้วย pandas ได้ '720188.0' (คอลัมน์เป็น float
    เพราะมีช่องว่างปน) ถ้าเก็บคนละรูป ล็อกที่ระบุ SO จะจับคู่ไม่ติด (ดู _so_tokens ใน Replan.py)
    """
    out = set()
    for t in re.split(r"[,\s;]+", _norm(v)):
        t = t.strip().upper()
        if not t or t in ("NAN", "NONE", "NAT"):
            continue
        if re.fullmatch(r"\d+\.0+", t):
            t = t.split(".")[0]
        out.add(t)
    return out


def key_of(week, item, mc_group, gauge, source="", year=None) -> tuple:
    """คีย์บล็อก (ไม่รวม SO — SO เป็นตัวกรองย่อยอีกชั้น)

    สัปดาห์ในคีย์ = คีย์สัปดาห์ของหน้าเว็บ ('29' / '29@2028') — หน้าเว็บส่งคีย์คอลัมน์มาตรง ๆ
    ส่วนแถวในชีทส่ง WEEK + YEAR ; ทั้งสองทางได้คีย์ชุดเดียวกัน (ล็อก W35 ปีหน้ากับปีนี้ไม่ทับกัน)

    รวม SOURCE ด้วยเพราะช่องเดียวกันมีได้ทั้งบล็อก booking และบล็อกงานแผน (item/สัปดาห์/
    เครื่องเดียวกัน) — ถ้าคีย์ไม่แยก การล็อกฝั่งหนึ่งจะไปทับล็อกอีกฝั่งที่มีอยู่
    """
    return (_wkey(week, year), _norm(item).upper(), _norm(mc_group).upper(), _gkey(gauge),
            _source(source))


# ---------------------------------------------------------------- อ่าน / เขียน
def _read_raw() -> list:
    """แถวในชีท Lock_Plan → [dict ตาม COLUMNS] (ไม่มีชีท = [])"""
    try:
        grid = masters.read_sheet("MasterMC", SHEET)
    except (KeyError, FileNotFoundError):
        return []
    cols = [_norm(c).upper().replace(" ", "_") for c in grid.get("columns", [])]
    idx = {c: i for i, c in enumerate(cols)}
    out = []
    for row in grid.get("rows", []):
        def get(c):
            i = idx.get(c)
            return row[i] if (i is not None and i < len(row)) else ""
        if not _norm(get("ITEM_CODE")) and not _norm(get("WEEK")):
            continue                                  # แถวว่างท้ายชีท
        out.append({c: get(c) for c in COLUMNS})
    return out


# ---------------------------------------------------------------- ทะเบียนปลดล็อกตั๋วโปรแกรม
def _wk_int(v):
    """สัปดาห์ → int ; ว่าง/อ่านไม่ออก = None (None = ไม่จำกัดหัวหรือท้ายช่วง)"""
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _rng_yw(week, year):
    """ปลายช่วงในทะเบียนปลดล็อก → YW ; สัปดาห์ว่าง = None (ไม่จำกัดปลายนั้น)"""
    w = _wk_int(week)
    if w is None:
        return None
    y = _wk_int(year)
    return y * 100 + w if y else _yw_of_key(_wkey(w))


def _program_items() -> set:
    """ITEM_CODE ในชีท Program ของ MasterMC (อ่านสดทุกครั้ง — แก้ Master แล้วมีผลทันที)

    อ่านผ่าน masters เหมือนชีทอื่นในไฟล์นี้ ไม่ยืมจาก plan_view เพื่อให้ทุกชีทที่โมดูลนี้
    แตะ มาจากการ resolve ไฟล์ทางเดียวกันหมด
    """
    try:
        grid = masters.read_sheet("MasterMC", "Program")
    except (KeyError, FileNotFoundError):
        return set()
    cols = [_norm(c).upper().replace(" ", "_") for c in grid.get("columns", [])]
    if "ITEM_CODE" not in cols:
        return set()
    i = cols.index("ITEM_CODE")
    return {_norm(r[i]).upper() for r in grid.get("rows", [])
            if i < len(r) and _norm(r[i])}


def _core_items() -> set:
    """ITEM_CODE ในทะเบียน core (Target_Stock) — ดูแค่รหัสผ้า

    ไม่เช็คทีม/ลูกค้าตรงนี้ (เหตุผลเดียวกับชีท Program): ฝั่งที่ล็อกให้ (program_lock.py) เช็ค
    item+ทีม+ลูกค้าแล้ว ส่วนตรงนี้เป็นแค่ "ควรจดทะเบียนปลดล็อกไหม" — จดเกินไว้ไม่เสียหาย
    แต่จดขาดจะโดนล็อกกลับทันที ; อ่านทะเบียนไม่ได้ = เซ็ตว่าง (ปลด/ล็อกยังทำงานได้)
    """
    try:
        import core_registry
        return set(core_registry.items())
    except Exception:
        return set()


# ชนิดล็อกอัตโนมัติ → (ชีททะเบียนปลดล็อก, ตัวอ่านเซ็ต item ของชนิดนั้น, REMARK ตอนจด)
_UNLOCK_REGISTERS = [
    (UNLOCK_SHEET, _program_items, "user ปลดล็อกเอง — ห้ามล็อกตั๋วโปรแกรมกลับ"),
    (CORE_UNLOCK_SHEET, _core_items, "user ปลดล็อกเอง — ห้ามล็อก core item กลับ"),
]


def _unlock_updates(rm_keys, add_keys, user="") -> dict:
    """ปลด/ล็อกบล็อกของ item ที่มีล็อกอัตโนมัติ → {ชีททะเบียน: แถวใหม่ทั้งชีท}

    ปลด (rm_keys) = จดสัปดาห์นั้นเป็นช่วงที่ห้ามล็อกกลับ ; กดล็อกเอง (add_keys) = คว้าน
    สัปดาห์นั้นออกจากช่วงที่ปลดไว้ → รอบหน้าระบบล็อกให้ต่อ
    item เป็นทั้งตั๋วโปรแกรมและ core = จดทั้งสองชีท ; ไม่เกี่ยวชนิดไหนเลย = {} (ไม่เขียนชีททะเบียน)
    """
    keys = sorted(set(rm_keys) | set(add_keys))
    out = {}
    if not keys:
        return out
    for sheet, members_of, remark in _UNLOCK_REGISTERS:
        members = members_of()
        hit = [k for k in keys if k[1] in members]
        if not hit:
            continue
        ranges = _read_unlock(sheet)
        for k in hit:
            weeks = [_yw_of_key(k[0])]
            if k in rm_keys:
                ranges = _unlock_add(ranges, k[1], k[2], k[3], weeks, user, remark)
            else:
                ranges = _unlock_del(ranges, k[1], k[2], k[3], weeks)
        out[sheet] = _unlock_rows(ranges)
    return out


def _read_unlock(sheet=UNLOCK_SHEET) -> list:
    """ชีททะเบียนปลดล็อก (Program_Unlock / Core_Unlock) → [{item, mcg, gauge, from, to, by, at, remark}]"""
    try:
        grid = masters.read_sheet("MasterMC", sheet)
    except (KeyError, FileNotFoundError):
        return []
    cols = [_norm(c).upper().replace(" ", "_") for c in grid.get("columns", [])]
    idx = {c: i for i, c in enumerate(cols)}
    out = []
    for row in grid.get("rows", []):
        def get(c):
            i = idx.get(c)
            return row[i] if (i is not None and i < len(row)) else ""
        item = _norm(get("ITEM_CODE")).upper()
        if not item:
            continue
        out.append({
            "item": item,
            "mcg": _norm(get("MC_GROUP")).upper(),
            "gauge": _gkey(get("GUAGE")) if _norm(get("GUAGE")) else "",
            # ช่วงเก็บเป็น YW (ปี*100+สัปดาห์) — ช่วงคร่อมปีเทียบถูก ; ปีว่าง (แถวเก่า) = ครั้งที่ใกล้วันนี้สุด
            # (program_lock.py อ่านคอลัมน์ชุดเดียวกันด้วยกติกาเดียวกัน)
            "from": _rng_yw(get("WEEK_FROM"), get("YEAR_FROM")),
            "to": _rng_yw(get("WEEK_TO"), get("YEAR_TO")),
            "by": _norm(get("UNLOCKED_BY")), "at": _norm(get("UNLOCKED_AT")),
            "remark": _norm(get("REMARK")),
        })
    return out


def _rng_match(r, item, mcg, gauge) -> bool:
    """แถวทะเบียนนี้คุมบล็อกนี้ไหม — MC_GROUP/GUAGE ในทะเบียนเว้นว่าง = ทุกเครื่อง"""
    return (r["item"] == item
            and (not r["mcg"] or r["mcg"] == mcg)
            and (not r["gauge"] or r["gauge"] == gauge))


def _rng_has(r, w) -> bool:
    return (r["from"] is None or w >= r["from"]) and (r["to"] is None or w <= r["to"])


def _unlock_add(ranges: list, item, mcg, gauge, weeks, user="",
                remark="user ปลดล็อกเอง — ห้ามล็อกตั๋วโปรแกรมกลับ") -> list:
    """บันทึกว่า user ปลดล็อกสัปดาห์พวกนี้แล้ว → ต่อกับช่วงเดิมถ้าติดกัน ไม่งั้นเปิดแถวใหม่

    weeks = YW (ปี*100+สัปดาห์) ; "ติดกัน" ดูบนปฏิทินจริง (W52/2026 → W1/2027 ติดกัน)
    เขียน MC_GROUP/GUAGE ว่างเสมอ = ปลดทั้งสัปดาห์นั้นของ item (ไม่ผูกกับเครื่อง) —
    งานที่ปลดแล้วมีสิทธิ์ย้ายไปเครื่องอื่นในสัปดาห์เดิม ถ้าผูกเครื่องไว้จะถูกล็อกกลับทันที
    """
    out = [dict(r) for r in ranges]
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for w in sorted(w for w in weeks if w is not None):
        if any(_rng_match(r, item, mcg, gauge) and _rng_has(r, w) for r in out):
            continue                                   # อยู่ในช่วงที่ปลดไว้แล้ว
        touch = None
        for r in out:
            if r["item"] != item or r["mcg"] or r["gauge"]:
                continue
            if r["to"] is not None and w == _yw_shift(r["to"], 1):
                r["to"], touch = w, r
                break
            if r["from"] is not None and w == _yw_shift(r["from"], -1):
                r["from"], touch = w, r
                break
        if touch is None:
            out.append({"item": item, "mcg": "", "gauge": "", "from": w, "to": w,
                        "by": _norm(user), "at": now,
                        "remark": remark})
        else:
            touch["by"], touch["at"] = _norm(user) or touch["by"], now
    return out


def _unlock_del(ranges: list, item, mcg, gauge, weeks) -> list:
    """user กดล็อกสัปดาห์นั้นกลับเอง → คว้านสัปดาห์นั้นออกจากช่วงที่ปลดไว้
    (ช่วงถูกเจาะกลาง = แตกเป็นสองแถว ; ปลายเปิดยังคงปลายเปิดฝั่งเดิม)"""
    weeks = [w for w in weeks if w is not None]
    if not weeks:
        return ranges
    out = []
    for r in ranges:
        if not _rng_match(r, item, mcg, gauge):
            out.append(r)
            continue
        pieces = [r]
        for w in weeks:
            nxt = []
            for p in pieces:
                if not _rng_has(p, w):
                    nxt.append(p)
                    continue
                # ปลายใหม่ต้องเป็นสัปดาห์ที่มีจริงบนปฏิทิน (YW−1 ของ W1 = W0 ไม่มีจริง
                # → อ่านกลับแล้วกลายเป็น "ไม่จำกัดปลาย" ปลดล็อกเกินที่ user ตั้งใจ)
                prev_w, next_w = _yw_shift(w, -1), _yw_shift(w, 1)
                if p["from"] is None or p["from"] <= prev_w:
                    nxt.append({**p, "to": prev_w})
                if p["to"] is None or next_w <= p["to"]:
                    nxt.append({**p, "from": next_w})
            pieces = nxt
        out.extend(pieces)
    return out


def _unlock_rows(ranges: list) -> list:
    """[{...}] → แถวพร้อมเขียนลงชีท (เรียงตาม item แล้วสัปดาห์)"""
    ranges = sorted(ranges, key=lambda r: (r["item"], r["from"] if r["from"] is not None else -1))

    def _w(yw):
        return int(yw) % 100 if yw is not None else ""

    def _y(yw):
        return int(yw) // 100 if yw is not None else ""
    # ช่วงภายในเป็น YW → เขียนกลับเป็นสัปดาห์ + ปี (ปีว่างจะเลื่อนความหมายตามวันที่ ช่วงที่ปลดไว้จะเพี้ยน)
    return [{"ITEM_CODE": r["item"],
             "WEEK_FROM": _w(r["from"]), "WEEK_TO": _w(r["to"]),
             "MC_GROUP": r["mcg"], "GUAGE": r["gauge"],
             "UNLOCKED_BY": r["by"], "UNLOCKED_AT": r["at"], "REMARK": r["remark"],
             "YEAR_FROM": _y(r["from"]), "YEAR_TO": _y(r["to"])}
            for r in ranges]


def _backup_once_a_day(path) -> None:
    """สำรอง MasterMC วันละครั้ง (ก่อนแก้ครั้งแรกของวัน)

    ล็อก/ปลดล็อกเป็นการกดถี่ ๆ — ถ้า .bak ทุกครั้งเหมือน masters.write_sheet
    จะได้ไฟล์สำรองหลายสิบไฟล์ต่อวัน ; เก็บของ "ก่อนแก้ครั้งแรกของวัน" ไฟล์เดียวพอ
    """
    bak = path.with_suffix(path.suffix + f".lockplan.{datetime.now():%Y%m%d}.bak")
    if not bak.exists():
        shutil.copy2(path, bak)


def ensure_sheet() -> bool:
    """สร้างชีท Lock_Plan / Program_Unlock (หัวคอลัมน์อย่างเดียว) ถ้ายังไม่มี → True เมื่อสร้างใหม่

    ให้ชีทโผล่ในหน้า Master ตั้งแต่ยังไม่เคยกดล็อก user จะได้พิมพ์เพิ่มมือได้เอง
    (เดิมชีทเกิดตอนกดล็อกครั้งแรกเท่านั้น — เข้าหน้า Master ไปก็ไม่เห็นชีทนี้)
    Program_Unlock = ช่วงสัปดาห์ที่ห้ามล็อกตั๋วโปรแกรมกลับ — user แก้ช่วงเองในชีทได้
    """
    try:
        path = masters._resolve("MasterMC")
    except (KeyError, FileNotFoundError):
        return False
    wb = load_workbook(path, data_only=False)
    try:
        need = [(SHEET, COLUMNS), (UNLOCK_SHEET, UNLOCK_COLUMNS),
                (CORE_UNLOCK_SHEET, UNLOCK_COLUMNS)]
        missing = [(s, c) for s, c in need if s not in wb.sheetnames]
        if not missing:
            return False
        _backup_once_a_day(path)
        for s, c in missing:
            wb.create_sheet(s).append(c)
        wb.save(path)
        return True
    finally:
        wb.close()


def _write_raw(rows: list, unlocks: dict | None = None) -> None:
    """เขียนทับชีท Lock_Plan (+ ชีททะเบียนปลดล็อกที่ส่งมา {ชีท: แถว}) — ชีทอื่นคงเดิม

    ทุกชีทเขียนในการเปิด/เซฟรอบเดียว: ปลดล็อกอัตโนมัติ 1 ครั้ง = ลบแถวล็อก + จดทะเบียน
    ถ้าแยกเซฟแล้วพังกลางทาง จะได้สถานะครึ่ง ๆ (ล็อกหายแต่ไม่มีทะเบียน → รอบหน้าล็อกกลับ)

    ผู้เรียกต้องถือล็อกไฟล์อยู่แล้ว (ดู _edit()) — ตัวมันเองไม่ล็อกซ้ำ เพราะการล็อก/ปลดล็อก
    คืออ่านชีททั้งก้อนมาแก้แล้วเขียนกลับ ต้องกันคนอื่นตั้งแต่ตอนอ่าน
    """
    path = masters._resolve("MasterMC")
    _backup_once_a_day(path)

    wb = load_workbook(path, data_only=False)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.create_sheet(SHEET)
    if ws.max_row > 0:
        ws.delete_rows(1, ws.max_row)
    ws.append(COLUMNS)
    for r in rows:
        ws.append([masters._coerce(r.get(c, "")) for c in COLUMNS])

    for sheet, unlock in (unlocks or {}).items():
        wu = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
        if wu.max_row > 0:
            wu.delete_rows(1, wu.max_row)
        wu.append(UNLOCK_COLUMNS)
        for r in unlock:
            wu.append([masters._coerce(r.get(c, "")) for c in UNLOCK_COLUMNS])

    with livesync.atomic_save(path) as tmp:
        wb.save(tmp)
    wb.close()


# ---------------------------------------------------------------- API ของหน้าเว็บ
@contextmanager
def _edit(user: str = "", detail: str = ""):
    """ครอบทุกคำสั่งที่แก้ชีทล็อก — เข้าคิวกับคนอื่นที่แก้ MasterMC แล้วป่าวประกาศเมื่อสำเร็จ

    ต้องครอบตั้งแต่ _read_raw() ไม่ใช่แค่ตอนเขียน: ถ้าสองคนอ่านชีทเดียวกันพร้อมกันแล้ว
    ต่างคนต่างเขียนกลับ ล็อกของคนที่เขียนก่อนจะหายไปกับตาราง "ฉบับก่อนหน้า" ของคนที่สอง

    ไม่มีด่านเทียบเลขรุ่น (base_rev) ที่นี่โดยตั้งใจ — ล็อก/ปลดล็อกเป็นคำสั่งรายบล็อก
    (ล็อก A, ปลด B) ไม่ใช่การเขียนทับทั้งตาราง สองคนทำคนละบล็อกพร้อมกันจึงรวมกันได้เอง
    ไม่ต้องเด้ง "มีคนแก้แซง" ให้เสียจังหวะ
    """
    with livesync.file_lock(masters._resolve("MasterMC")):
        yield
        livesync.bump(livesync.PLAN_LOCK, user=user, detail=detail)


def _row_key(r) -> tuple:
    return key_of(r["WEEK"], r["ITEM_CODE"], r["MC_GROUP"], r["GUAGE"], r.get("SOURCE", ""),
                  year=r.get("YEAR"))


def _week_cells(key) -> dict:
    """คีย์สัปดาห์ → ช่อง WEEK/YEAR ของแถวในชีท"""
    w, y = _split_key(key)
    return {"WEEK": w, "YEAR": y}


def _sort_rows(rows: list) -> None:
    """เรียงแถวชีทตามเวลาจริง (ปี, สัปดาห์) แล้ว item — แถวเก่าที่ไม่มีปีใช้ปีของครั้งที่ใกล้วันนี้สุด"""
    def k(r):
        yw = _yw_of_key(_wkey(r.get("WEEK"), r.get("YEAR"))) or 0
        return yw, _norm(r.get("ITEM_CODE")), _source(r.get("SOURCE", ""))
    rows.sort(key=k)


def list_locks() -> list:
    """รายการล็อกทั้งหมด → ให้ Gantt เอาไปแสดงเครื่องหมายล็อกบนบล็อก

    key = 'WEEK|ITEM|MC_GROUP|GUAGE|SOURCE' (normalize แล้ว) — frontend เทียบตรง ๆ ได้
    เรียกตอนเปิดหน้าแผน จึงถือโอกาสสร้างชีทเปล่าให้ด้วย (ดู ensure_sheet)
    """
    try:
        ensure_sheet()
    except Exception:
        pass                                    # สร้างไม่ได้ (ไฟล์ถูกเปิดค้าง) ก็ยังอ่านต่อได้
    out = []
    for r in _read_raw():
        w, item, mcg, g, src = _row_key(r)
        out.append({
            "key": f"{w}|{item}|{mcg}|{g}|{src}",
            "week": w, "item": item, "mc_group": mcg, "gauge": g,
            # ชีทมีความหมายเฉพาะล็อกฝั่งแผน — ฝั่ง booking คืนค่าว่างไว้ ไม่ให้เข้าใจผิด
            "source": src, "sheet": _sheet(r.get("SHEET", "")) if src == SRC_PLAN else "",
            "so": _norm(r["SO_NO"]),
            "qty": r["QTY"] if r["QTY"] != "" else None,
            "mc": r["MC"] if r["MC"] != "" else None,
            "by": _norm(r["LOCKED_BY"]), "at": _norm(r["LOCKED_AT"]),
            "remark": _norm(r["REMARK"]),
            # ล็อกที่ระบบเติมให้เอง (ตั๋วโปรแกรม / core item) — หน้าเว็บใช้แยกข้อความบนบล็อก
            "auto": _norm(r["LOCKED_BY"]).upper() in AUTO_KINDS,
            "auto_kind": AUTO_KINDS.get(_norm(r["LOCKED_BY"]).upper(), ""),
        })
    return out


def add_lock(week, item, mc_group, gauge, so="", qty=None, mc=None,
             user="", remark="", source="", sheet="") -> dict:
    """ล็อกบล็อกหนึ่ง (ล็อกซ้ำคีย์เดิม = อัปเดตแถวเดิม ไม่เพิ่มแถวใหม่)"""
    w, it, mcg, g, src = key_of(week, item, mc_group, gauge, source)
    if not w or not it:
        raise ValueError("ต้องระบุสัปดาห์และ item ที่จะล็อก")

    with _edit(user, f"ล็อก {it} W{w}"):
        rows = [r for r in _read_raw() if _row_key(r) != (w, it, mcg, g, src)]
        # กดล็อกเอง = ยกเลิก "ช่วงที่เคยปลด" ของสัปดาห์นั้น → รอบหน้าระบบล็อกอัตโนมัติให้ต่อ
        unlocks = _unlock_updates(set(), {(w, it, mcg, g, src)})
        rows.append({
            **_week_cells(w),
            "ITEM_CODE": it, "MC_GROUP": mcg, "GUAGE": g,
            # เขียนรูปที่ normalize แล้ว ('720188' ไม่ใช่ '720188.0') — หน้าเว็บส่ง SO ที่มาจาก
            # pandas ซึ่งติดทศนิยมมาได้ ถ้าเก็บดิบ ๆ ชีทจะอ่านยากและจับคู่ booking ไม่ติด
            "SO_NO": ",".join(sorted(_so_set(so))),
            "QTY": qty if qty is not None else "",
            "MC": mc if mc is not None else "",
            "SOURCE": src,
            # ชีทมีความหมายเฉพาะล็อกฝั่งแผน (booking อ่านจาก booking_final ใบเดียวเสมอ)
            "SHEET": _sheet(sheet) if src == SRC_PLAN else "",
            "LOCKED_BY": _norm(user),
            "LOCKED_AT": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "REMARK": _norm(remark),
        })
        _sort_rows(rows)
        _write_raw(rows, unlocks)
    return {"ok": True, "locked": f"{w}|{it}|{mcg}|{g}|{src}", "count": len(rows),
            "rev": livesync.rev(livesync.PLAN_LOCK)}


def apply_locks(add=None, remove=None, user="") -> dict:
    """ล็อก/ปลดล็อกหลายบล็อกในคำสั่งเดียว — อ่าน/เขียนชีทรอบเดียว

    หน้าเว็บกด "Lock Plan" ของงานที่กระจายหลายสัปดาห์ แล้วติ๊กเลือกสัปดาห์ทีเดียวหลายอัน
    ถ้าวน add_lock()/remove_lock() ทีละใบ จะเปิด+เซฟ MasterMC.xlsx ทุกรอบ (ช้ามาก และถ้า
    พังกลางทางล็อกจะเข้าไปครึ่ง ๆ ไม่ตรงกับที่ user เห็นบนจอ)

    add / remove = [{week, item, mc_group, gauge, so, qty, mc, remark, source, sheet}]
    คีย์ที่อยู่ทั้งสองฝั่ง → ฝั่ง add ชนะ (ถือเป็น "อัปเดตแถวเดิม" เหมือน add_lock)
    """
    add, remove = list(add or []), list(remove or [])
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    def _k(d):
        return key_of(d.get("week"), d.get("item"), d.get("mc_group"), d.get("gauge"),
                      d.get("source", ""))

    new_rows, add_keys = [], set()
    for d in add:
        w, it, mcg, g, src = _k(d)
        if not w or not it:
            raise ValueError("ต้องระบุสัปดาห์และ item ที่จะล็อก")
        add_keys.add((w, it, mcg, g, src))
        new_rows.append({
            **_week_cells(w),
            "ITEM_CODE": it, "MC_GROUP": mcg, "GUAGE": g,
            "SO_NO": ",".join(sorted(_so_set(d.get("so", "")))),
            "QTY": d.get("qty") if d.get("qty") is not None else "",
            "MC": d.get("mc") if d.get("mc") is not None else "",
            "SOURCE": src,
            "SHEET": _sheet(d.get("sheet", "")) if src == SRC_PLAN else "",
            "LOCKED_BY": _norm(user),
            "LOCKED_AT": now,
            "REMARK": _norm(d.get("remark", "")),
        })
    rm_keys = {_k(d) for d in remove} - add_keys

    with livesync.file_lock(masters._resolve("MasterMC")):
        rows = _read_raw()
        keep = [r for r in rows if _row_key(r) not in add_keys and _row_key(r) not in rm_keys]
        removed = sum(1 for r in rows if _row_key(r) in rm_keys)
        if not new_rows and not removed:
            return {"ok": True, "added": 0, "removed": 0, "count": len(rows),
                    "rev": livesync.rev(livesync.PLAN_LOCK)}

        # ล็อกอัตโนมัติ (ตั๋วโปรแกรม / core): ปลด = จดช่วงสัปดาห์ไว้ (ห้ามล็อกกลับ) / ล็อกเอง = คว้านออกจากช่วง
        unlocks = _unlock_updates(rm_keys, add_keys, user)

        keep.extend(new_rows)
        _sort_rows(keep)
        _write_raw(keep, unlocks)
        # bump อยู่ในล็อก (ไม่ใช้ _edit เพราะต้อง return กลางทางตอนไม่มีอะไรเปลี่ยน)
        ent = livesync.bump(livesync.PLAN_LOCK, user=user,
                            detail=f"ล็อก {len(new_rows)} ปลด {removed}")
    return {"ok": True, "added": len(new_rows), "removed": removed, "count": len(keep),
            "rev": ent["rev"]}


def move_lock(week, item, mc_group, gauge, to_week, to_mc_group=None, to_gauge=None,
              qty=None, mc=None, user="", source=SRC_PLAN) -> dict:
    """ย้ายล็อกตามงานที่ถูกลากไปสัปดาห์/เครื่องอื่น

    ทำเป็นคำสั่งเดียว (เขียนชีทครั้งเดียว) ไม่ใช่ remove แล้ว add เพราะถ้าพังกลางทาง
    ล็อกจะหายไปเฉย ๆ ทั้งที่งานยังอยู่ — และประหยัดการเปิด/เซฟ MasterMC ไปหนึ่งรอบ

    ไม่พบล็อกเดิม = ไม่ทำอะไร (งานที่ลากไม่ได้ถูกล็อกไว้) → moved 0
    ปลายทางมีล็อกอยู่แล้ว = ยุบเป็นแถวเดียว (งานถูกลากไปรวมกับก้อนที่ล็อกไว้อยู่แล้ว)
    """
    src = _source(source)
    old = key_of(week, item, mc_group, gauge, src)
    new = key_of(to_week, item,
                 to_mc_group if to_mc_group is not None else mc_group,
                 to_gauge if to_gauge is not None else gauge, src)

    with livesync.file_lock(masters._resolve("MasterMC")):
        rows = _read_raw()
        hit = next((r for r in rows if _row_key(r) == old), None)
        if hit is None:
            return {"ok": True, "moved": 0, "count": len(rows),
                    "rev": livesync.rev(livesync.PLAN_LOCK)}

        keep = [r for r in rows if _row_key(r) not in (old, new)]
        hit = dict(hit)
        hit.update(_week_cells(new[0]))          # คีย์ปลายทางอาจเป็นปีอื่น → เขียนปีด้วยเสมอ
        hit["MC_GROUP"], hit["GUAGE"] = new[2], new[3]
        if qty is not None:
            hit["QTY"] = qty
        if mc is not None:
            hit["MC"] = mc
        if _norm(user):
            hit["LOCKED_BY"] = _norm(user)
        hit["LOCKED_AT"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        keep.append(hit)
        _sort_rows(keep)
        _write_raw(keep)
        ent = livesync.bump(livesync.PLAN_LOCK, user=user,
                            detail=f"ย้ายล็อก {old[1]} → W{new[0]}")
    return {"ok": True, "moved": 1,
            "from": "|".join(old), "to": "|".join(new), "count": len(keep),
            "rev": ent["rev"]}


def remove_lock(week, item, mc_group, gauge, source="", user="") -> dict:
    """ปลดล็อกบล็อกหนึ่ง (ไม่มีอยู่ = ไม่ทำอะไร ถือว่าสำเร็จ)

    เป็นตั๋วโปรแกรม / core item → จดสัปดาห์นั้นลงชีท Program_Unlock / Core_Unlock ด้วย
    ไม่งั้นรอบวางแผนถัดไป program_lock.py จะล็อกกลับให้ทันที (user ปลดแล้ว = ปลดจนกว่าจะกดล็อกเอง)
    """
    k = key_of(week, item, mc_group, gauge, source)
    with livesync.file_lock(masters._resolve("MasterMC")):
        rows = _read_raw()
        keep = [r for r in rows if _row_key(r) != k]

        unlocks = _unlock_updates({k}, set(), user)
        if len(keep) != len(rows) or unlocks:
            _write_raw(keep, unlocks)
            livesync.bump(livesync.PLAN_LOCK, user=user, detail=f"ปลดล็อก {k[1]} W{k[0]}")
    return {"ok": True, "removed": len(rows) - len(keep), "count": len(keep),
            "rev": livesync.rev(livesync.PLAN_LOCK)}


# ---------------------------------------------------------------- ใช้ฝั่ง pipeline
def locked_keys(source=SRC_BOOKING) -> dict:
    """{(week, item, mc_group, gauge): set(SO ที่ล็อก)} — set ว่าง = ล็อกทุก SO

    เลือกได้ว่าเอาล็อกฝั่งไหน (BOOKING = แถวใน booking DETAIL / PLAN = แถวในไฟล์แผน)

    Replan.py ใช้แนวคิดเดียวกันตัดสินว่าแถวไหนห้ามยกออกมาวางใหม่
    (import ตรง ๆ ไม่ได้เพราะ Replan อยู่คนละโฟลเดอร์ — Replan มีสำเนา loader ของตัวเอง
     ที่อ่านชีทเดียวกันด้วย pandas ; ถ้าแก้คอลัมน์ที่นี่ ต้องแก้ที่นั่นด้วย)
    """
    want = _source(source)
    out: dict = {}
    for r in _read_raw():
        k = key_of(r["WEEK"], r["ITEM_CODE"], r["MC_GROUP"], r["GUAGE"], r.get("SOURCE", ""),
                   year=r.get("YEAR"))
        if not k[0] or not k[1] or k[4] != want:
            continue
        out.setdefault(k[:4], set()).update(_so_set(r["SO_NO"]))
    return out
